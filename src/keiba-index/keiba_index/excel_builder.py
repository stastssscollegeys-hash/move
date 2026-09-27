# -*- coding: utf-8 -*-
"""Excel生成エンジン (openpyxl)"""
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .factors import FACTORS, F_MAX, GUIDE_DATA, PENDING_FACTORS
from .scoring import score_horse, get_shirushi
from .colors import (COL, GRADE_COLOR, VENUE_BG, VENUE_DARK,
                     SHIRUSHI_COLOR, SHIRUSHI_ESTBG)


# ──────────────────── スタイルヘルパー ────────────────────
def _thin():
    s = Side(style="thin")
    return Border(left=s, right=s, top=s, bottom=s)


def cell(ws, r, c, val="", bg="FFFFFF", fg="000000",
         bold=False, size=9, align="center", wrap=False, italic=False):
    cl = ws.cell(row=r, column=c, value=val)
    cl.fill = PatternFill("solid", fgColor=bg)
    cl.font = Font(color=fg, bold=bold, size=size, name="Meiryo UI", italic=italic)
    cl.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    cl.border = _thin()
    return cl


def hdr(ws, r, c, val, bg=None, size=9, wrap=True):
    return cell(ws, r, c, val, bg=bg or COL["NAVY"], fg=COL["WHITE"],
                bold=True, size=size, wrap=wrap)


# ──────────────────── サマリーシート ────────────────────
def build_summary_sheet(ws, race_program: dict, date_sat: str, date_sun: str):
    ws.title = "_cover"
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:K1")
    c = ws.cell(row=1, column=1,
        value=f"JRA週末全開催  アスメシ式 独自指数 v3.0  {date_sat}・{date_sun}")
    c.fill = PatternFill("solid", fgColor=COL["NAVY"])
    c.font = Font(color="FFFFFF", bold=True, size=13, name="Meiryo UI")
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:K2")
    c2 = ws.cell(row=2, column=1,
        value="枠順確定: 木曜14:00 | 調教評価: 金〜土 | 馬体重・クッション値: 当日 | G1は完全スコア済み・他レースは枠順確定後入力")
    c2.fill = PatternFill("solid", fgColor="FFF176")
    c2.font = Font(color="333333", size=8, name="Meiryo UI")
    c2.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 16

    hdrs = ["競馬場", "日程", "R#", "レース名", "グレード", "芝/ダ", "距離(m)", "条件", "シート名", "入力状況"]
    ws.row_dimensions[3].height = 20
    for ci, h in enumerate(hdrs, 1):
        hdr(ws, 3, ci, h)

    r = 4
    for (venue, day), races in sorted(race_program.items(),
            key=lambda x: ("土日".index(x[0][1]), list(VENUE_BG.keys()).index(x[0][0]))):
        date_str = date_sat if day == "土" else date_sun
        sheet_name = f"{venue}_{day}"
        vbg = VENUE_BG[venue]

        for race in races:
            rnum = race["num"]
            rname = race["name"]
            grade = race["grade"]
            surf  = race["surface"]
            dist  = race["distance"]
            cond  = race["conditions"]
            has_horses = bool(race.get("horses"))

            status = "✅ 完全スコア済" if has_horses else "⏳ 枠順確定後入力"
            gbg = GRADE_COLOR.get(grade, "999999")

            cell(ws, r, 1, venue, bg=vbg, fg=VENUE_DARK[venue], bold=True)
            cell(ws, r, 2, date_str, bg=vbg)
            cell(ws, r, 3, f"R{rnum}", bg=vbg, bold=(grade in ["G1","G2","G3"]))
            cn = ws.cell(row=r, column=4, value=rname)
            cn.fill = PatternFill("solid", fgColor=gbg if grade in ["G1","G2","G3","OP"] else "F5F5F5")
            cn.font = Font(color="FFFFFF" if grade in ["G1","G2","G3","OP"] else "333333",
                           bold=(grade in ["G1","G2","G3"]), size=9, name="Meiryo UI")
            cn.alignment = Alignment(horizontal="left", vertical="center")
            cn.border = _thin()

            gbg2 = gbg if grade in ["G1","G2","G3","OP"] else "EEEEEE"
            gfg  = "FFFFFF" if grade in ["G1","G2","G3","OP"] else "555555"
            cell(ws, r, 5, grade, bg=gbg2, fg=gfg, bold=(grade in ["G1","G2","G3"]))
            cell(ws, r, 6, surf, bg="E8F5E9" if surf == "芝" else "FFF3E0")
            cell(ws, r, 7, dist)
            cell(ws, r, 8, cond, align="left")
            cell(ws, r, 9, sheet_name, fg="1565C0", bold=True)
            cell(ws, r, 10, status, bg=COL["GOOD"] if has_horses else COL["PENDING"])
            ws.row_dimensions[r].height = 16
            r += 1

    for ci, w in zip(range(1, 11), [8, 9, 5, 26, 8, 6, 8, 12, 12, 16]):
        ws.column_dimensions[get_column_letter(ci)].width = w


# ──────────────────── 会場シート ────────────────────
def build_venue_sheet(wb, venue: str, day: str, races: list, date_str: str):
    sname = f"{venue}_{day}"
    ws = wb.create_sheet(sname)
    ws.sheet_view.showGridLines = False

    vbg = VENUE_DARK[venue]
    n_base = 8
    n_f    = len(FACTORS)
    total_cols = n_base + n_f + 5

    ws.merge_cells(f"A1:{get_column_letter(total_cols)}1")
    c = ws.cell(row=1, column=1,
        value=f"【{venue}競馬場】 {date_str}  全{len(races)}レース 独自指数 v3.0 (150点満点)")
    c.fill = PatternFill("solid", fgColor=vbg)
    c.font = Font(color="FFFFFF", bold=True, size=12, name="Meiryo UI")
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    current_row = 2
    BASE_HDRS = ["シルシ", "順", "馬番", "馬名", "騎手", "調教師", "所属", "オッズ"]
    SCORE_HDRS = [f[0] for f in FACTORS]
    EXTRA_HDRS = ["フラグ", "確定計", "未確定", "推定計", "備考"]
    all_h = BASE_HDRS + SCORE_HDRS + EXTRA_HDRS

    for race in races:
        rnum    = race["num"]
        rname   = race["name"]
        grade   = race["grade"]
        surf    = race["surface"]
        dist    = race["distance"]
        cond    = race["conditions"]
        horses  = race.get("horses", [])
        is_big  = grade in ["G1","G2","G3","OP"]

        # レースヘッダー行
        gbg = GRADE_COLOR.get(grade, "4A4A4A")
        ws.merge_cells(f"A{current_row}:{get_column_letter(total_cols)}{current_row}")
        label = (f"R{rnum:02d}  {rname}  [{grade}]  {surf}{dist}m  {cond}"
                 f"{'   ★ 完全スコア済' if horses else '   ⏳ 出走馬確定待ち'}")
        c = ws.cell(row=current_row, column=1, value=label)
        c.fill = PatternFill("solid", fgColor=gbg if is_big else "455A64")
        c.font = Font(color="FFFFFF", bold=True, size=10 if is_big else 9, name="Meiryo UI")
        c.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[current_row].height = 22 if is_big else 18
        current_row += 1

        # 列ヘッダー行
        for ci, h in enumerate(all_h, 1):
            fi = ci - n_base - 1
            if ci <= n_base:
                bg_h = "2C3E50"
            elif 0 <= fi < n_f:
                bg_h = "1A4A7A" if FACTORS[fi][2] == "new" else "2E5F4E"
            else:
                bg_h = "4A3560"
            cl = ws.cell(row=current_row, column=ci, value=h)
            cl.fill = PatternFill("solid", fgColor=bg_h)
            cl.font = Font(color="FFFFFF", bold=True, size=8, name="Meiryo UI")
            cl.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cl.border = _thin()
        ws.row_dimensions[current_row].height = 32
        current_row += 1

        # 馬行
        if horses:
            scored = [(h, score_horse(h), get_shirushi(score_horse(h)["est"]))
                      for h in horses]
            scored.sort(key=lambda x: x[1]["est"], reverse=True)

            for rank, (h, sc, shirushi) in enumerate(scored, 1):
                bg_r = COL["ROW_ODD"] if rank % 2 == 0 else COL["ROW_EVEN"]
                scol = SHIRUSHI_COLOR.get(shirushi, COL["KESHI"])

                cell(ws, current_row, 1, shirushi, bg=scol, fg="FFFFFF", bold=True, size=12)
                cell(ws, current_row, 2, rank, bg=bg_r, bold=True)
                cell(ws, current_row, 3, "—", bg=COL["PENDING"])
                cell(ws, current_row, 4, h["名"], bg=bg_r, bold=True, align="left",
                     fg="1E3A5F", size=10)
                jfg = "CC0000" if h.get("騎") in ("C.ルメール","川田将雅","D.レーン") else "000000"
                cell(ws, current_row, 5, h.get("騎",""), bg=bg_r, fg=jfg)
                cell(ws, current_row, 6, h.get("師",""), bg=bg_r)
                sfg = "FF6B35" if h.get("所") == "栗東" else "1565C0"
                cell(ws, current_row, 7, h.get("所",""), bg=bg_r, fg=sfg, bold=True)
                odds = h.get("オッズ", 0)
                cell(ws, current_row, 8, odds,
                     bg="FFE0E0" if odds <= 5 else bg_r,
                     fg="CC0000" if odds <= 5 else "000000",
                     bold=odds <= 5)

                for fi, s in enumerate(h["sc"]):
                    ci = n_base + fi + 1
                    if s is None:
                        cv = ws.cell(row=current_row, column=ci, value="—")
                        cv.fill = PatternFill("solid", fgColor=COL["PENDING"])
                        cv.font = Font(color="AAAAAA", size=8, italic=True, name="Meiryo UI")
                        cv.alignment = Alignment(horizontal="center", vertical="center")
                        cv.border = _thin()
                    else:
                        ratio = s / FACTORS[fi][1]
                        sbg = "E8F5E9" if ratio >= 0.85 else "FFF9C4" if ratio >= 0.65 else "FFEBEE"
                        sfg2 = "2E7D32" if ratio >= 0.85 else "F57F17" if ratio >= 0.65 else "C62828"
                        cell(ws, current_row, ci, s, bg=sbg, fg=sfg2, bold=(ratio >= 0.85))

                flg_col = n_base + n_f + 1
                ftotal = sc["flg"]
                fbg = COL["GOOD"] if ftotal >= 0 else COL["ALERT"]
                ffg = "1B5E20" if ftotal >= 0 else "B71C1C"
                cell(ws, current_row, flg_col,   f"{ftotal:+d}" if ftotal else "0",
                     bg=fbg, fg=ffg, bold=True)
                cell(ws, current_row, flg_col+1, sc["conf"], bold=True,
                     bg="E3F2FD" if sc["conf"] >= 100 else bg_r)
                cell(ws, current_row, flg_col+2, f"+{sc['pend']}", bg=COL["PENDING"])
                cell(ws, current_row, flg_col+3, sc["est"],
                     bg=SHIRUSHI_ESTBG.get(shirushi,"F5F5F5"), fg=scol, bold=True, size=11)
                memo = h.get("memo","")
                cm = ws.cell(row=current_row, column=flg_col+4, value=memo)
                cm.fill = PatternFill("solid", fgColor=bg_r)
                cm.font = Font(color="CC0000" if "⚠" in memo else "444444",
                               size=8, name="Meiryo UI")
                cm.alignment = Alignment(horizontal="left", vertical="center")
                cm.border = _thin()
                ws.row_dimensions[current_row].height = 18
                current_row += 1
        else:
            # テンプレート行 (16頭)
            for post in range(1, 17):
                bg_r = COL["ROW_ODD"] if post % 2 == 0 else COL["ROW_EVEN"]
                cell(ws, current_row, 1, "?", bg="EEEEEE", fg="BBBBBB")
                cell(ws, current_row, 2, post, bg=bg_r)
                cell(ws, current_row, 3, post, bg=bg_r)
                cell(ws, current_row, 4, "← 馬名入力", bg=COL["PENDING"],
                     fg="999999", align="left", italic=True)
                for ci_ in range(5, 9):
                    cell(ws, current_row, ci_, "", bg=COL["PENDING"])
                for fi in range(n_f):
                    ci = n_base + fi + 1
                    cv = ws.cell(row=current_row, column=ci, value="")
                    cv.fill = PatternFill("solid",
                        fgColor=COL["PENDING"] if fi in PENDING_FACTORS else "F5F5F5")
                    cv.font = Font(size=9, name="Meiryo UI")
                    cv.alignment = Alignment(horizontal="center", vertical="center")
                    cv.border = _thin()
                flg_col = n_base + n_f + 1
                for ei in range(5):
                    ce = ws.cell(row=current_row, column=flg_col+ei,
                                 value="0" if ei == 0 else "")
                    ce.fill = PatternFill("solid",
                        fgColor="EEEEEE" if ei < 4 else "F5F5F5")
                    ce.border = _thin()
                    ce.font = Font(size=9, name="Meiryo UI", color="BBBBBB")
                    ce.alignment = Alignment(horizontal="center", vertical="center")
                ws.row_dimensions[current_row].height = 16
                current_row += 1

        current_row += 1  # レース間スペース

    # 列幅
    cw = {1:5, 2:4, 3:5, 4:16, 5:12, 6:12, 7:6, 8:7}
    for fi in range(n_f):
        cw[n_base + fi + 1] = 7
    cw.update({n_base+n_f+1:12, n_base+n_f+2:7,
               n_base+n_f+3:8,  n_base+n_f+4:8, n_base+n_f+5:20})
    for ci, w in cw.items():
        ws.column_dimensions[get_column_letter(ci)].width = w


# ──────────────────── 採点ガイドシート ────────────────────
def build_guide_sheet(wb):
    wsg = wb.create_sheet("採点ガイド")
    wsg.sheet_view.showGridLines = False

    wsg.merge_cells("A1:E1")
    c = wsg.cell(row=1, column=1,
        value="アスメシ式 独自指数 v3.0 — 17ファクター 150点満点 採点ガイド")
    c.fill = PatternFill("solid", fgColor=COL["NAVY"])
    c.font = Font(color="FFFFFF", bold=True, size=11, name="Meiryo UI")
    c.alignment = Alignment(horizontal="center", vertical="center")
    wsg.row_dimensions[1].height = 25

    for ci, h in enumerate(["ファクター","満点","種別","主な採点基準","データ取得タイミング"], 1):
        hdr(wsg, 2, ci, h)

    for ri, (name, mp, kind, criteria, timing) in enumerate(GUIDE_DATA, 3):
        bg = COL["NEW_F"] if kind == "NEW" else COL["OLD_F"]
        cell(wsg, ri, 1, name, bg=bg, bold=True, align="left")
        cell(wsg, ri, 2, f"{mp}点", bg=bg, bold=True)
        cell(wsg, ri, 3, kind, bg=bg,
             fg="1A4A7A" if kind == "NEW" else "2E5F4E", bold=True)
        cell(wsg, ri, 4, criteria, bg=bg, align="left")
        pending_timing = "当日" in timing or "木曜" in timing or "金" in timing
        cell(wsg, ri, 5, timing, bg=COL["PENDING"] if pending_timing else bg, align="left")
        wsg.row_dimensions[ri].height = 18

    ri = len(GUIDE_DATA) + 3
    cell(wsg, ri, 1, "合計", bold=True)
    cell(wsg, ri, 2, f"{F_MAX}点", bold=True, bg="FFF176")

    ri += 2
    wsg.merge_cells(f"A{ri}:E{ri}")
    cc = wsg.cell(row=ri, column=1, value="判定基準")
    cc.fill = PatternFill("solid", fgColor=COL["NAVY"])
    cc.font = Font(color="FFFFFF", bold=True, name="Meiryo UI")
    cc.alignment = Alignment(horizontal="center", vertical="center")

    for sr, cr, sbg in [("◎ 本命","115点以上",COL["HONMEI"]),
                         ("○ 対抗","100〜114点",COL["TAIKOU"]),
                         ("▲ 穴","88〜99点",COL["ANA"]),
                         ("△ 注意","75〜87点",COL["CHUUI"]),
                         ("— 消し","74点以下",COL["KESHI"])]:
        ri += 1
        cell(wsg, ri, 1, sr, bg=sbg, fg="FFFFFF", bold=True)
        cell(wsg, ri, 2, cr)
        wsg.row_dimensions[ri].height = 18

    for ci, w in zip(range(1, 6), [14, 7, 7, 42, 26]):
        wsg.column_dimensions[get_column_letter(ci)].width = w


# ──────────────────── 入力チェックシート ────────────────────
def build_checklist_sheet(wb, date_sat: str, date_sun: str):
    wsc = wb.create_sheet("入力チェック")
    wsc.sheet_view.showGridLines = False

    wsc.merge_cells("A1:D1")
    c = wsc.cell(row=1, column=1, value="週末レース 入力待ちデータ & チェックリスト")
    c.fill = PatternFill("solid", fgColor=COL["NAVY"])
    c.font = Font(color="FFFFFF", bold=True, size=11, name="Meiryo UI")
    c.alignment = Alignment(horizontal="center", vertical="center")
    wsc.row_dimensions[1].height = 25

    for ci, h in enumerate(["項目","発表タイミング","情報源","対象"], 1):
        hdr(wsc, 2, ci, h)

    items = [
        ("⑦ 枠順（全レース）",       "木曜 14:00",           "JRA公式 / netkeiba 出馬表",    "全会場・全レース"),
        ("⑨ 最終追い切り（全馬）",   "金曜〜土曜",            "各スポーツ紙 / netkeiba 調教",  "全馬"),
        ("⑩ 馬体重（全馬）",         f"{date_sat}・{date_sun} 朝", "JRA馬体重公示 / netkeiba", "全馬"),
        ("⑯ クッション値（全場）",   f"{date_sat}・{date_sun} 朝", "JRA公式サイト 馬場情報",   "全競馬場"),
        ("⑪ 確定単勝オッズ（全馬）", "前日〜当日直前",         "JRAオッズ / netkeiba",         "全馬（⑪乖離再評価）"),
        ("フラグ再確認（距離未経験等）","枠順確定後",           "各馬の過去戦績",               "G1・G3馬を優先"),
    ]
    for ri, (item, timing, source, target) in enumerate(items, 3):
        urgent = "木曜" in timing or "朝" in timing
        cell(wsc, ri, 1, item, bg=COL["PENDING"] if urgent else "F5F5F5",
             bold=urgent, align="left")
        cell(wsc, ri, 2, timing, bg=COL["PENDING"] if urgent else "F5F5F5", align="left")
        cell(wsc, ri, 3, source, align="left")
        cell(wsc, ri, 4, target, align="left")
        wsc.row_dimensions[ri].height = 18

    for ci, w in zip(range(1, 5), [26, 22, 28, 28]):
        wsc.column_dimensions[get_column_letter(ci)].width = w


# ──────────────────── メインビルド ────────────────────
def build(race_program: dict, output_path: str,
          date_sat: str, date_sun: str) -> openpyxl.Workbook:
    """
    指数表Excelを生成する。

    Parameters
    ----------
    race_program : dict
        {(venue, day): [race_dict, ...]}
    output_path : str
        保存先Excelパス
    date_sat / date_sun : str
        表示用日付文字列 (例: "4/18(土)")

    Returns
    -------
    openpyxl.Workbook
    """
    wb = openpyxl.Workbook()

    # サマリー
    build_summary_sheet(wb.active, race_program, date_sat, date_sun)

    # 会場シート (土→日 × 中山→阪神→福島)
    order = [("中山","土"), ("中山","日"), ("阪神","土"), ("阪神","日"),
             ("福島","土"), ("福島","日")]
    for venue, day in order:
        if (venue, day) not in race_program:
            continue
        races    = race_program[(venue, day)]
        date_str = f"2026年{date_sat.replace('(土','')}" if day == "土" else f"2026年{date_sun.replace('(日','')}"
        # 整形: "4/18(土)" → "2026年4月18日(土)"
        ds = date_sat if day == "土" else date_sun
        full_ds = ds.replace("4/18", "2026年4月18日").replace("4/19", "2026年4月19日")
        build_venue_sheet(wb, venue, day, races, full_ds)

    build_guide_sheet(wb)
    build_checklist_sheet(wb, date_sat, date_sun)

    wb.save(output_path)
    return wb
