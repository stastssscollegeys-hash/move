# -*- coding: utf-8 -*-
"""
アスメシ競馬予想 SNS投稿案 — 週末 期待値の高いレース特集
2026/08/01(土)・08/02(日)
自信度7以上20レースをv5買い目（46パターン機械評価）で構築した結果から、
E[回収率]が高い非重賞レースTOP5をSNS投稿用にまとめる。
重賞2レース（クイーンS/アイビスSD）は別途フル投稿案あり、本編では対象外。
"""
import sys, json
sys.stdout.reconfigure(encoding='utf-8')
from pathlib import Path
from collections import defaultdict
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from gen_weekend_confidence_report import calc_confidence
from gen_kaime_v5 import choose_pattern, BUDGET_BY_CONF

JSON_PATH = r"C:\Users\User\Desktop\競馬予想レポート\20260801\週末ビッグデータ_20260801-0802_records.json"
ODDS_PATH = str(Path(__file__).resolve().parent / "odds_20260802_g3.json")
OUT = Path(r"C:\Users\User\Desktop\競馬予想レポート\20260801\週末期待値TOP5_SNS投稿案.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

FONT = '游ゴシック'
TOP_N = 5

def set_font(run, size=10, bold=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

def h1(doc, text, color=(15,71,97)):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        set_font(r, 14); r.font.color.rgb = RGBColor(*color)

def h3(doc, text):
    p = doc.add_heading(text, level=3)
    for r in p.runs:
        set_font(r, 11)

def box(doc, lines, size=10):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("\n".join(lines))
    set_font(r, size)


def load():
    with open(JSON_PATH, encoding='utf-8') as f:
        data = json.load(f)
    with open(ODDS_PATH, encoding='utf-8') as f:
        real_odds = json.load(f)
    races = defaultdict(list)
    for r in data['records']:
        races[(r['date'], r['競馬場'], r['R'])].append(r)
    for k in races:
        races[k].sort(key=lambda x: x['AI予測順位'])
    conf = {k: calc_confidence(v) for k, v in races.items()}
    return races, conf, real_odds


def main():
    races, conf, real_odds = load()
    candidates = []
    for k in races:
        c, _ = conf[k]
        if c < 7:
            continue
        top = races[k][0]
        if top['グレード']:   # 重賞は除外（別投稿案あり）
            continue
        budget = BUDGET_BY_CONF[c]
        odds_key = f"{k[0]}_{k[1]}_{k[2]}"
        name, desc, res, ranking, p, q = choose_pattern(races[k], budget, real_odds.get(odds_key))
        if name is None:
            continue
        candidates.append((k, c, name, desc, res, p, q))
    candidates.sort(key=lambda x: -x[4]['E_rate'])
    top5 = candidates[:TOP_N]

    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.8); s.bottom_margin = Cm(1.8)
        s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)

    t = doc.add_heading("週末 期待値TOP5レース SNS投稿案", level=0)
    for r in t.runs: set_font(r, 16)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("2026/08/01(土)・08/02(日)　自信度7以上の非重賞レースからE[回収率]上位5レース厳選")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in sub.runs: set_font(r, 10)
    doc.add_paragraph()

    weekday = {'20260801':'土', '20260802':'日'}

    def race_block(k, c, name, desc, res, recs):
        top = recs[0]
        marks = [r for r in recs if r['AI印']]
        lines = [f"{k[1]}{k[2]}R {top['レース名']} {top['距離']} {top['頭数']}頭　自信度{c}/10"]
        lines.append("印: " + "　".join(f"{r['AI印']}{r['馬番']}{r['馬名']}" for r in marks))
        lines.append(f"買い目: {name}（{desc}）")
        for btype, idxs, amount, odds in res['bets']:
            sep = '→' if btype in ('馬単','3連単') else '-'
            nums = sep.join(str(recs[i]['馬番']) for i in idxs)
            lines.append(f"　{btype} {nums}　{amount:,}円（想定{odds:.1f}倍）")
        lines.append(f"評価: E[回収率]{res['E_rate']*100:.0f}% / P(300%達成){res['P_target']*100:.0f}%")
        return lines

    # ── X投稿①: 全体ダイジェスト ─────────────────────────
    h1(doc, "【X投稿①】期待値TOP5 ダイジェスト（コピー用）", color=(20,100,50))
    digest = [
        "🏇【週末 期待値TOP5レース】2026/08/01-02",
        "━━ 自信度7以上20レースをビッグデータ×46パターン機械評価で分析 ━━",
        "",
        "重賞以外で最もEV（期待値）が高いと判定された5レースを厳選しました👇",
        "",
    ]
    for i, (k, c, name, desc, res, p, q) in enumerate(top5, 1):
        recs = races[k]
        top = recs[0]
        marks = [r for r in recs if r['AI印']]
        mark_line = "　".join(f"{r['AI印']}{r['馬番']}{r['馬名']}" for r in marks[:3])
        digest.append(f"{i}位 {k[0][4:6]}/{k[0][6:]} {k[1]}{k[2]}R {top['レース名']}"
                       f"　E[回収率]{res['E_rate']*100:.0f}%")
        digest.append(f"　{mark_line}")
    digest += ["", "各レースの根拠・買い目は返信スレッドで解説します🎯", "",
               "#競馬予想 #AI予想 #JRA #期待値 #ビッグデータ競馬"]
    box(doc, ["【コピー用】"] + digest)

    # ── X投稿②〜⑥: 個別スレッド ─────────────────────────
    h1(doc, "【X投稿②〜⑥】個別スレッド（1レース1投稿・コピー用）", color=(20,100,50))
    for i, (k, c, name, desc, res, p, q) in enumerate(top5, 1):
        recs = races[k]
        top = recs[0]
        h3(doc, f"投稿{i+1}：{i}位 {k[1]}{k[2]}R {top['レース名']}")

        # 根拠テキスト生成
        h_rec = next(r for r in recs if r['AI印'] == '◎')
        hi = recs.index(h_rec)
        pop_note = ("過小評価=大きな妙味" if p[hi] > q[hi]*1.2 else
                    "ほぼ人気どおり" if p[hi] > q[hi]*0.8 else "人気先行に注意")
        reasons = [
            f"◎{h_rec['馬名']}: 総合指数{h_rec['総合指数']}（独自{h_rec['独自指数']}・ML{h_rec['ML能力%']}%）",
            f"モデル勝率{p[hi]*100:.0f}% vs 大衆推定{q[hi]*100:.0f}% → {pop_note}",
            f"脚質{h_rec.get('脚質','?')}・近5走平均{h_rec.get('近5走平均着',0)}着・"
            f"騎手勝率{h_rec.get('騎手勝率%',0)}%",
        ]

        lines = [
            "【コピー用】",
            f"🏇【{k[1]}{k[2]}R {top['レース名']}】{k[0][4:6]}/{k[0][6:]}（{weekday.get(k[0],'')}）",
            f"━━ {top['距離']} {top['頭数']}頭立て / 自信度{c}/10 期待値ランキング{i}位 ━━",
            "",
            "🎯 予想印",
        ]
        marks = [r for r in recs if r['AI印']]
        for r in marks:
            lines.append(f"{r['AI印']} {r['馬番']}番 {r['馬名']}　総合指数{r['総合指数']}")
        lines += ["", "📊 データ根拠"]
        lines += [f"・{x}" for x in reasons]
        lines += ["", f"🎲 買い目（{name}）"]
        for btype, idxs, amount, odds in res['bets']:
            sep = '→' if btype in ('馬単','3連単') else '-'
            nums = sep.join(str(recs[i2]['馬番']) for i2 in idxs)
            names = "・".join(recs[i2]['馬名'] for i2 in idxs)
            lines.append(f"{btype} {nums}（{names}）　{amount:,}円　想定{odds:.1f}倍")
        lines += [
            "",
            f"E[回収率]{res['E_rate']*100:.0f}% / P(回収率300%達成){res['P_target']*100:.0f}%",
            "",
            f"#競馬予想 #AI予想 #{k[1]}競馬場 #期待値",
        ]
        box(doc, lines)

    doc.add_paragraph()

    # ── Threads投稿 ──────────────────────────────────────
    h1(doc, "【Threads投稿】期待値TOP5ダイジェスト", color=(80,40,120))
    box(doc, ["【コピー用】"] + digest)

    doc.add_paragraph()

    # ── Note記事 ──────────────────────────────────────────
    h1(doc, "【Note記事全文】週末 期待値TOP5レース徹底解説", color=(180,80,0))
    note_lines = [
        "週末ビッグデータ予想 期待値TOP5レース特集 — アスメシ競馬予想", "",
        "こんにちは、アスメシ競馬予想です🍱", "",
        "8/1（土）・8/2（日）の全72レース929頭をビッグデータで分析し、",
        "自信度7以上のレースをさらに46パターンの買い目ライブラリで機械評価。",
        "重賞を除く中から、期待回収率が最も高かった5レースを徹底解説します。", "",
        "「E[回収率]」はモデル勝率と大衆の想定オッズから計算した期待値、",
        "「P(300%達成)」は回収率300%を超える確率です。", "",
        "---", "",
    ]
    for i, (k, c, name, desc, res, p, q) in enumerate(top5, 1):
        recs = races[k]
        top = recs[0]
        marks = [r for r in recs if r['AI印']]
        note_lines.append(f"第{i}位　{k[0][4:6]}/{k[0][6:]}（{weekday.get(k[0],'')}）"
                           f"{k[1]}{k[2]}R {top['レース名']} {top['距離']} {top['頭数']}頭")
        note_lines.append(f"自信度{c}/10　E[回収率]{res['E_rate']*100:.0f}%　"
                           f"P(300%達成){res['P_target']*100:.0f}%")
        note_lines.append("")
        for r in marks:
            note_lines.append(f"　{r['AI印']} {r['馬番']}番 {r['馬名']}　総合指数{r['総合指数']}"
                               f"（独自{r['独自指数']}・ML{r['ML能力%']}%・{r.get('脚質','?')}）")
        note_lines.append("")
        note_lines.append(f"買い目パターン: {name}（{desc}）")
        for btype, idxs, amount, odds in res['bets']:
            sep = '→' if btype in ('馬単','3連単') else '-'
            nums = sep.join(str(recs[i2]['馬番']) for i2 in idxs)
            names = "・".join(recs[i2]['馬名'] for i2 in idxs)
            note_lines.append(f"　{btype} {nums}（{names}）　{amount:,}円　想定{odds:.1f}倍")
        note_lines.append("")
        note_lines.append("---")
        note_lines.append("")
    note_lines += [
        "重賞2レース（クイーンステークス・アイビスサマーダッシュ）は全頭診断・最新情報を",
        "反映した別記事で詳しく解説しています。あわせてご覧ください。", "",
        "予想が参考になったらフォロー＆いいねお願いします！",
        "みなさんの週末が楽しいものになりますように🍜🎉", "",
        "#競馬予想 #AI予想 #JRA #ビッグデータ競馬 #期待値",
    ]
    box(doc, note_lines, size=9.5)

    doc.save(str(OUT))
    print(f"保存完了: {OUT}")
    print(f"\n期待値TOP5:")
    for i, (k, c, name, desc, res, p, q) in enumerate(top5, 1):
        top = races[k][0]
        print(f"  {i}位 {k[0][4:6]}/{k[0][6:]} {k[1]}{k[2]}R {top['レース名'][:16]:<16} "
              f"自信度{c} {name:<16} E{res['E_rate']*100:>5.0f}% P300={res['P_target']*100:>3.0f}%")


if __name__ == '__main__':
    main()
