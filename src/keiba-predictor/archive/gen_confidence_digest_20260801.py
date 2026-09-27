# -*- coding: utf-8 -*-
"""
アスメシ競馬予想 SNS投稿案 — 週末 自信度7以上 注目レース ダイジェスト
2026/08/01(土)・08/02(日)
17ファクター独自指数×LightGBM v44×累積DB9,259戦のビッグデータから
自信度7以上（10段階）と判定された20レースをSNS投稿用にまとめる。
重賞2レース（クイーンS/アイビスSD）は別途フル投稿案あり、本ダイジェストでは概要のみ。
"""
import sys, json
sys.stdout.reconfigure(encoding='utf-8')
from pathlib import Path
from collections import defaultdict
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from gen_weekend_confidence_report import calc_confidence

JSON_PATH = r"C:\Users\User\Desktop\競馬予想レポート\20260801\週末ビッグデータ_20260801-0802_records.json"
OUT = Path(r"C:\Users\User\Desktop\競馬予想レポート\20260801\週末自信度7以上_SNS投稿案.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

FONT = '游ゴシック'
CONF_MIN = 7

def set_font(run, size=10, bold=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

def h1(doc, text, color=(15,71,97)):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        set_font(r, 14); r.font.color.rgb = RGBColor(*color)

def h3(doc, text):
    p = doc.add_heading(text, level=3)
    for r in p.runs:
        set_font(r, 11)

def box(doc, lines, size=10):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("\n".join(lines))
    set_font(r, size)


def load_races():
    with open(JSON_PATH, encoding='utf-8') as f:
        data = json.load(f)
    races = defaultdict(list)
    for r in data['records']:
        races[(r['date'], r['競馬場'], r['R'])].append(r)
    for k in races:
        races[k].sort(key=lambda x: x['AI予測順位'])
    conf = {k: calc_confidence(v) for k, v in races.items()}
    return races, conf


def race_line(k, races, conf, with_stars=True):
    recs = races[k]
    top = recs[0]
    c, _ = conf[k]
    grade = f"[{top['グレード']}]" if top['グレード'] else ""
    stars = f" 自信度{c}/10{'★'*c}" if with_stars else ""
    header = f"■{k[1]}{k[2]}R {top['レース名']}{grade} {top['距離']} {top['頭数']}頭{stars}"
    marks = [r for r in recs if r['AI印']]
    mark_line = "　".join(f"{r['AI印']}{r['馬番']}{r['馬名']}" for r in marks)
    return header, mark_line


def build_day_post(day, day_label, races, conf, keys):
    lines = [
        f"🏇【本日の自信度TOP注目レース】2026/08/{day[6:]}（{day_label}）",
        "━━ 札幌・新潟・中京 全レース中、自信度7以上（10段階）だけ厳選 ━━",
        "",
        "17ファクター独自指数×LightGBM ML×累積DB9,259戦のビッグデータで、",
        "確定枠順まで反映した本日の最終予想印です👇",
        "",
    ]
    grouped = defaultdict(list)
    for k in keys:
        grouped[k[1]].append(k)
    for venue in sorted(grouped, key=lambda v: -max(conf[k][0] for k in grouped[v])):
        vkeys = sorted(grouped[venue], key=lambda k: -conf[k][0])
        lines.append(f"【{venue}競馬場】")
        for k in vkeys:
            header, mark_line = race_line(k, races, conf)
            lines.append(header)
            lines.append(f"　{mark_line}")
        lines.append("")
    lines += [
        "重賞2レース（クイーンS・アイビスSD）は最終追い切り評価・陣営コメントまで反映した",
        "全頭診断＋買い目を別記事で詳述しています🎯",
        "期待値が特に高い5レースは「期待値TOP5」特集でさらに深掘りしています📊",
        "",
        "#競馬予想 #AI予想 #JRA #ビッグデータ競馬",
    ]
    return lines


def build_race_thread(k, races, conf):
    """自信度9-10の特に自信のあるレースは個別スレッドとして厚めに。"""
    recs = races[k]
    top = recs[0]
    c, tags = conf[k]
    grade = f"[{top['グレード']}]" if top['グレード'] else ""
    marks = [r for r in recs if r['AI印']]
    lines = [
        f"🏇【{k[1]}{k[2]}R {top['レース名']}{grade}】自信度{c}/10{'★'*c}",
        f"━━ {top['距離']} {top['頭数']}頭立て / 2026/08/{k[0][6:]} ━━",
        "",
    ]
    for r in marks:
        lines.append(f"{r['AI印']} {r['馬番']}番 {r['馬名']}　総合指数{r['総合指数']}（独自{r['独自指数']}・ML{r['ML能力%']}%）")
    lines.append("")
    lines.append("【自信度の根拠】")
    for t in tags:
        lines.append(f"・{t}")
    lines.append("")
    lines.append(f"#競馬予想 #AI予想 #JRA #{k[1]}競馬場")
    return lines


def main():
    races, conf = load_races()
    all_keys = [k for k in races if conf[k][0] >= CONF_MIN]
    all_keys.sort(key=lambda k: (k[0], -conf[k][0], k[1], k[2]))

    keys_by_day = defaultdict(list)
    for k in all_keys:
        keys_by_day[k[0]].append(k)

    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.8); s.bottom_margin = Cm(1.8)
        s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)

    t = doc.add_heading("週末 自信度7以上 注目レース SNS投稿案", level=0)
    for r in t.runs: set_font(r, 16)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("2026/08/01(土)・08/02(日)　全72レース中20レース厳選　17ファクター独自指数×ML v44×累積DB9,259戦")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in sub.runs: set_font(r, 10)
    doc.add_paragraph()

    weekday = {'20260801': '土', '20260802': '日'}

    # ── X投稿: 日別ダイジェスト ─────────────────────────
    h1(doc, "【X投稿】日別ダイジェスト（コピー用）", color=(20,100,50))
    for day in sorted(keys_by_day):
        h3(doc, f"{day[4:6]}/{day[6:]}（{weekday.get(day,'')}）版")
        lines = build_day_post(day, weekday.get(day, ''), races, conf, keys_by_day[day])
        box(doc, ["【コピー用】"] + lines)

    doc.add_paragraph()

    # ── X投稿: 自信度9-10 個別スレッド ──────────────────
    h1(doc, "【X投稿】自信度9〜10 個別スレッド（特に自信のあるレース）", color=(20,100,50))
    top_keys = [k for k in all_keys if conf[k][0] >= 9]
    for i, k in enumerate(top_keys, 1):
        h3(doc, f"スレッド{i}/{len(top_keys)}")
        lines = build_race_thread(k, races, conf)
        box(doc, ["【コピー用】"] + lines)

    doc.add_paragraph()

    # ── Threads投稿 ──────────────────────────────────────
    h1(doc, "【Threads投稿】日別ダイジェスト", color=(80,40,120))
    for day in sorted(keys_by_day):
        h3(doc, f"{day[4:6]}/{day[6:]}（{weekday.get(day,'')}）")
        lines = build_day_post(day, weekday.get(day, ''), races, conf, keys_by_day[day])
        box(doc, ["【コピー用】"] + lines)

    doc.add_paragraph()

    # ── Note記事 ──────────────────────────────────────────
    h1(doc, "【Note記事全文】自信度7以上 全20レース詳細", color=(180,80,0))
    note_lines = [
        "週末ビッグデータ予想 自信度TOPレース特集 — アスメシ競馬予想", "",
        "こんにちは、アスメシ競馬予想です🍱", "",
        "8/1（土）・8/2（日）の札幌・新潟・中京、全72レース929頭を",
        "17ファクター独自指数×LightGBM ML v44×累積DB9,259戦のビッグデータで分析し、",
        "確定枠順まで反映した最終予想から、自信度10段階評価で7以上のレースだけを厳選しました。", "",
        "自信度は「上位馬との指数差」「独自指数とML能力の一致」「累積DBの実績裏付け」",
        "「頭数」などから機械的に算出しています。", "",
        "---", "",
    ]
    for day in sorted(keys_by_day):
        note_lines.append(f"{day[4:6]}/{day[6:]}（{weekday.get(day,'')}）")
        note_lines.append("")
        for k in keys_by_day[day]:
            recs = races[k]
            top = recs[0]
            c, tags = conf[k]
            grade = f"[{top['グレード']}]" if top['グレード'] else ""
            note_lines.append(f"{k[1]}{k[2]}R {top['レース名']}{grade} {top['距離']} {top['頭数']}頭　自信度{c}/10 {'★'*c}")
            marks = [r for r in recs if r['AI印']]
            for r in marks:
                note_lines.append(
                    f"　{r['AI印']} {r['馬番']}番 {r['馬名']}　総合指数{r['総合指数']}（独自{r['独自指数']}・ML{r['ML能力%']}%・{r.get('脚質','?')}）")
            note_lines.append(f"　根拠: {' / '.join(tags)}")
            note_lines.append("")
        note_lines.append("---")
        note_lines.append("")
    note_lines += [
        "重賞2レース（クイーンステークス・アイビスサマーダッシュ）は最終追い切り評価・",
        "陣営コメント・確定オッズ想定まで反映した全頭診断＋買い目を別記事で詳しく解説しています。",
        "期待値が特に高い5レースの「期待値TOP5」特集もあわせてご覧ください。", "",
        "予想が参考になったらフォロー＆いいねお願いします！",
        "みなさんの週末が楽しいものになりますように🍜🎉", "",
        "#競馬予想 #AI予想 #JRA #ビッグデータ競馬 #札幌競馬場 #新潟競馬場 #中京競馬場",
    ]
    box(doc, note_lines, size=9.5)

    doc.save(str(OUT))
    print(f"保存完了: {OUT}")
    print(f"対象: {len(all_keys)}レース（8/1: {len(keys_by_day.get('20260801',[]))} / 8/2: {len(keys_by_day.get('20260802',[]))}）")
    print(f"個別スレッド化（自信度9-10）: {len(top_keys)}レース")


if __name__ == '__main__':
    main()
