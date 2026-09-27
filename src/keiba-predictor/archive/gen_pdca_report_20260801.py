# -*- coding: utf-8 -*-
"""
アスメシ競馬予想 — 8/1(土) レース終了後PDCAレポート
日次PDCA(独自指数) + 累積DB更新 + v5買い目答え合わせ + 改善アクション(v5.2)を1本のWordに。
"""
import sys, json
sys.stdout.reconfigure(encoding='utf-8')
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT = Path(r"C:\Users\User\Desktop\競馬予想レポート\20260801\20260801_PDCAレポート_v5検証.docx")
EVAL_JSON = Path(r"C:\Users\User\Desktop\競馬予想レポート\20260801\v5買い目_答え合わせ_20260801.json")

FONT = '游ゴシック'

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(1.8); s.bottom_margin = Cm(1.8)
    s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)

def set_font(run, size=10, bold=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

def h1(text, color=(15,71,97)):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        set_font(r, 14); r.font.color.rgb = RGBColor(*color)

def body(text, size=10, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, size, bold=bold)

def box(lines, size=9.5):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("\n".join(lines))
    set_font(r, size)

with open(EVAL_JSON, encoding='utf-8') as f:
    ev = json.load(f)

t = doc.add_heading("8/1(土) レース終了後 PDCAレポート", level=0)
for r in t.runs: set_font(r, 16)
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph("日次PDCA + 累積DB更新 + v5買い目検証 + 改善アクション（v5.2）")
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
for r in sub.runs: set_font(r, 10)
doc.add_paragraph()

# ── P: 当日実績サマリー ──────────────────────────────
h1("■ Step1-2 日次PDCA・累積DB（実行済み）")
box([
    "【日次PDCA（結果ベース17ファクター独自指数）】",
    "・◎的中: 24/36R（66.7%）",
    "・AI1位→1着: 24/36R（66.7%）",
    "・AI3位以内→1着: 35/36R（97.2%）",
    "・保存: 20260801_独自指数_PDCA.xlsx（全馬F01-F17スコア完全版）",
    "",
    "【累積DB更新】",
    "・463頭を新規追加 → DB総計 9,722レコード（21日分・720R）",
    "・累積◎的中率: 500/720R（69.4%）",
    "・保存: 累積DB_独自指数.xlsx",
])

# ── D/C: v5買い目の答え合わせ ────────────────────────
h1("■ v5買い目 答え合わせ（事前予想の検証）", color=(150,30,30))
box([
    f"総投資 {ev['total_inv']:,}円 / 総回収 {ev['total_ret']:,}円 / "
    f"回収率 {ev['total_ret']/ev['total_inv']*100:.1f}%",
    "",
    "レース別:",
])
tbl = doc.add_table(rows=1, cols=7)
tbl.style = 'Table Grid'
hdr = tbl.rows[0].cells
for i, txt in enumerate(['レース','自信度','パターン','投資','回収','◎着順','結果(1-2-3着)']):
    hdr[i].text = txt
    for p_ in hdr[i].paragraphs:
        for r in p_.runs: set_font(r, 9, bold=True)
for rr in ev['races']:
    row = tbl.add_row().cells
    row[0].text = rr['race']
    row[1].text = str(rr['conf'])
    row[2].text = rr['pattern']
    row[3].text = f"{rr['inv']:,}円" if rr['inv'] else "—"
    row[4].text = f"{rr['ret']:,}円" if rr['inv'] else "—"
    row[5].text = f"{rr['honmei_finish']}着" if rr['honmei_finish'] else "?"
    row[6].text = " / ".join(f"{o}着{ub}番{nm}" for o, ub, nm in rr['top3'])
    for c in row:
        for p_ in c.paragraphs:
            for r in p_.runs: set_font(r, 8.5)
doc.add_paragraph()

# ── C: ファクター深掘り分析 ──────────────────────────
h1("■ ファクター深掘り分析（何が当たり、何が外れたか）", color=(150,30,30))
box([
    "【1】自信度10の2レースで◎が大敗（中京1R 8着 / 札幌1R 9着）",
    "・両方とも2歳/3歳未勝利戦＝過去走データが薄い若馬戦。",
    "・データ由来の自信度が過大評価になり、単勝一点・◎絡み集中の買い目が全滅。",
    "・一方で○は両レースとも1着、▲も2着/3着 → モデルの2〜5位評価は機能していた。",
    "",
    "【2】V系（妙味馬）パターンが2戦0勝（札幌4R・札幌6R）",
    "・札幌6Rは◎ウーマンズパワーが1.6倍で快勝したのに、妙味馬軸ワイドのため回収0円。",
    "・推定オッズの誤差が「妙味の幻影」を生んだ。実オッズなしでのV系選択は危険。",
    "",
    "【3】印上位3頭中2頭がtop3に入ったレースが8戦中6戦",
    "・◎1頭に依存する券種構成では取りこぼす。BOX/2頭軸系が構造的に優位な日だった。",
    "",
    "【4】勝ちパターンの検証: F6 3連複▲2頭軸（中京8R・回収率380%）",
    "・◎1.2倍1人気×▲20.2倍5人気の構図で1170円×26＝30,420円。",
    "・「人気◎×人気薄▲の2頭軸」はまさに設計思想どおりに機能。",
    "",
    "【5】見送り判定は正解（中京7R豊橋S: ◎6着・買っても全滅だった）",
    "",
    "【6】自信度別: 自信度8が127%でプラス、自信度10が0% → 自信度と回収が逆転。",
    "・原因は【1】の若馬戦キャリブレーション不足に帰着。",
])

# ── A: 改善アクション ────────────────────────────────
h1("■ 改善アクション（v5.2として実装済み）", color=(20,100,50))
box([
    "① 若馬戦ガード: レース名に「新馬」「未勝利」を含むレースは、",
    "　 1点集中型パターン（T1/T2/T4/U1/W1/E1/S1/S2）を選択禁止 + 予算を60%に圧縮。",
    "　 → BOX/流し/2頭軸系のみで戦う（○▲の好走を取りこぼさない）。",
    "",
    "② V系（妙味馬）ガード: 実オッズ（--odds-file）がある時のみV1〜V4を候補に入れる。",
    "　 推定オッズの誤差によるEV幻影を排除。",
    "",
    "③ gen_kaime_v5.py に実装済み。日曜（8/2）分の買い目は改修版で再生成した。",
    "",
    "【v5.2適用後の日曜買い目の変化】",
    "・新潟11R（未勝利・V2だった）→ 見送りに変化",
    "・札幌10Rポプラス（V3だった）→ 見送りに変化",
    "・新潟10R（V2だった）→ S5 3連単1-2着固定に切替",
    "・新潟1R/中京1R/札幌3R（未勝利）→ 予算60%圧縮（BOX/軸流し系は維持）",
    "・日曜トータル: 12レース中9参加・3見送り / 総投資72,000円",
])

# ── 継続監視項目 ──────────────────────────────────────
h1("■ 継続監視項目（次回PDCAでチェック）")
box([
    "・F6/F1（2頭軸3連複）とW4（ワイドBOX）の成績を継続蓄積 → 主力パターン化の判断",
    "・自信度計算の若馬戦キャップ（現状は買い目側ガードのみ。的中傾向を見て自信度側も調整）",
    "・重賞のオッズ推定精度（明日のクイーンS/アイビスSDで想定オッズと実オッズの乖離を記録）",
    "・「◎が人気馬(2倍以下)の時の券種選択」: 今日は3勝とも人気◎ → 単勝一点では妙味なし、",
    "　 2頭軸で人気薄を拾う構成が正解だった。パターン選択ロジックは既にこの方向で機能中。",
])

doc.save(str(OUT))
print(f"保存完了: {OUT}")
