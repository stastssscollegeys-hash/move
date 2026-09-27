# -*- coding: utf-8 -*-
"""
アスメシ競馬予想 SNS投稿案 — 8/2(日) 注目レース（自信度7以上）
土曜版と同構成: Xダイジェスト + 自信度9-10個別スレッド + 深掘りNote記事（v5.2買い目込み）
重賞2レース（クイーンS/アイビスSD）は概要+別記事誘導、平場は買い目まで詳述。
"""
import sys, json
sys.stdout.reconfigure(encoding='utf-8')
from pathlib import Path
from collections import defaultdict
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

SRC = Path(__file__).resolve().parent
sys.path.insert(0, str(SRC))
from gen_weekend_confidence_report import calc_confidence, honmei_reasoning, FACTOR_LABELS, FKEYS
from gen_kaime_v5 import choose_pattern, BUDGET_BY_CONF, is_young_race

TARGET_DATE = "20260802"
JSON_PATH = r"C:\Users\User\Desktop\競馬予想レポート\20260801\週末ビッグデータ_20260801-0802_records.json"
ODDS_PATH = str(SRC / "odds_20260802_g3.json")
OUT = Path(rf"C:\Users\User\Desktop\競馬予想レポート\{TARGET_DATE}\{TARGET_DATE}_日曜注目レースSNS投稿案.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

FONT = '游ゴシック'
CONF_MIN = 7

def set_font(run, size=10, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(1.8); s.bottom_margin = Cm(1.8)
    s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)

def h1(text, color=(15,71,97)):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        set_font(r, 14); r.font.color.rgb = RGBColor(*color)

def h3(text):
    p = doc.add_heading(text, level=3)
    for r in p.runs:
        set_font(r, 11)

def box(lines, size=10):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("\n".join(lines))
    set_font(r, size)

def rival_line(r, label):
    fs = sorted(((k, r.get(k, 0)) for k in FKEYS), key=lambda x: -x[1])[:2]
    strong = "・".join(f"{FACTOR_LABELS[k]}{v:.0f}点" for k, v in fs)
    parts = [f"{label}{r['馬番']}番{r['馬名']}: 総合{r['総合指数']}（独自{r['独自指数']}・ML{r['ML能力%']}%）",
             f"強み: {strong}"]
    if r.get('DB最高指数', 0) >= 72:
        parts.append(f"累積DB最高{r['DB最高指数']}点")
    return "　".join(parts)


# ── データ ──────────────────────────────────────────────
with open(JSON_PATH, encoding='utf-8') as f:
    data = json.load(f)
with open(ODDS_PATH, encoding='utf-8') as f:
    real_odds = json.load(f)
races = defaultdict(list)
for r in data['records']:
    races[(r['date'], r['競馬場'], r['R'])].append(r)
for k in races:
    races[k].sort(key=lambda x: x['AI予測順位'])
conf = {k: calc_confidence(v) for k, v in races.items()}
keys = sorted([k for k in races if k[0] == TARGET_DATE and conf[k][0] >= CONF_MIN],
              key=lambda k: (-conf[k][0], k[1], k[2]))

race_data = []
for k in keys:
    recs = races[k]
    c, tags = conf[k]
    budget = BUDGET_BY_CONF[c]
    odds_key = f"{k[0]}_{k[1]}_{k[2]}"
    name, desc, res, ranking, p, q = choose_pattern(recs, budget, real_odds.get(odds_key))
    race_data.append((k, c, tags, recs, recs[0], name, desc, res))

n_join = sum(1 for rd in race_data if rd[5] is not None)
total_amt = sum(rd[7]['total'] for rd in race_data if rd[7])

# ── タイトル ────────────────────────────────────────────
t = doc.add_heading("8/2(日) 注目レースSNS投稿案【馬番訂正版】", level=0)
for r in t.runs: set_font(r, 16)
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph(f"自信度7以上 全{len(keys)}レース（重賞2本含む）　確定枠順で再計算した訂正版・改良版買い目（v5.2）")
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
for r in sub.runs: set_font(r, 10)
doc.add_paragraph()

# ══════════════════════════════════════════════════════════
# X投稿: 日別ダイジェスト
# ══════════════════════════════════════════════════════════
h1("【X投稿0】お詫びと訂正（最初に投稿・コピー用）", color=(150,30,30))
box([
    "【コピー用】",
    "🏇【お詫びと訂正】",
    "",
    "昨日投稿した本日8/2(日)の注目レース予想について、",
    "一部レースの馬番に誤りがありました。",
    "平場レースの枠順確定前のデータを使用してしまったことが原因です。",
    "申し訳ありません🙇",
    "",
    "確定した枠順で全レースを再計算した訂正版を、",
    "この投稿に続けて再掲します。",
    "重賞2本（クイーンS・アイビスSD）の予想・馬番に誤りはありません。",
    "",
    "今後は枠順確定のタイミングを確認してから投稿するよう改善します。",
    "引き続きよろしくお願いします🙏",
    "",
    "#競馬予想 #AI予想 #訂正とお詫び",
])

h1("【X投稿①】日曜ダイジェスト（コピー用）", color=(20,100,50))
digest = [
    "🏇【本日の自信度TOP注目レース・訂正版】2026/08/02（日）",
    "━━ 札幌・新潟・中京 全レース中、自信度7以上（10段階）だけ厳選 ━━",
    "",
    "昨日は1着馬が印内75%・3連複11.7倍的中🎯",
    "昨日の反省（若いレースの買い方）を修正した改良版で本日に臨みます👇",
    "",
]
grouped = defaultdict(list)
for k, c, tags, recs, top, name, desc, res in race_data:
    grouped[k[1]].append((k, c, recs, top))
for venue in sorted(grouped, key=lambda v: -max(c for _, c, _, _ in grouped[v])):
    digest.append(f"【{venue}競馬場】")
    for k, c, recs, top in sorted(grouped[venue], key=lambda x: -x[1]):
        grade = f"[{top['グレード']}]" if top['グレード'] else ""
        marks = [r for r in recs if r['AI印']]
        digest.append(f"■{k[2]}R {top['レース名']}{grade} {top['距離']} 自信度{c}/10{'★'*c}")
        digest.append("　" + "　".join(f"{r['AI印']}{r['馬番']}{r['馬名']}" for r in marks))
    digest.append("")
digest += [
    "重賞2本（クイーンS・アイビスSD）は全頭診断＋買い目を別スレッドで詳述🏆",
    "",
    "#競馬予想 #AI予想 #JRA #ビッグデータ競馬 #クイーンステークス #アイビスサマーダッシュ",
]
box(["【コピー用】"] + digest)

# ══════════════════════════════════════════════════════════
# X投稿: 自信度9-10個別スレッド
# ══════════════════════════════════════════════════════════
h1("【X投稿②〜】自信度9〜10 個別スレッド（コピー用）", color=(20,100,50))
top_races = [rd for rd in race_data if rd[1] >= 9]
for i, (k, c, tags, recs, top, name, desc, res) in enumerate(top_races, 1):
    grade = f"[{top['グレード']}]" if top['グレード'] else ""
    marks = [r for r in recs if r['AI印']]
    h3(f"スレッド{i}/{len(top_races)}: {k[1]}{k[2]}R {top['レース名']}{grade}")
    lines = [
        "【コピー用】",
        f"🏇【{k[1]}{k[2]}R {top['レース名']}{grade}】自信度{c}/10{'★'*c}",
        f"━━ {top['距離']} {top['頭数']}頭立て / 2026/08/02（日） ━━",
        "",
    ]
    for r in marks:
        lines.append(f"{r['AI印']} {r['馬番']}番 {r['馬名']}　総合指数{r['総合指数']}")
    lines.append("")
    lines.append("【自信度の根拠】")
    for tg in tags:
        lines.append(f"・{tg}")
    if top['グレード']:
        lines.append("")
        lines.append("→ 全頭診断・詳細な買い目は重賞専用スレッドで🏆")
    lines += ["", f"#競馬予想 #AI予想 #{k[1]}競馬場"]
    box(lines)

# ══════════════════════════════════════════════════════════
# Note記事: 深掘り+買い目統合（土曜版と同形式）
# ══════════════════════════════════════════════════════════
h1("【Note記事全文】日曜深掘り+買い目統合版（コピー用）", color=(180,80,0))
note = []
note.append("8月2日(日) 自信度TOPレース徹底解説 — アスメシ競馬予想")
note.append("")
note.append("はじめに")
note.append("")
note.append("※本記事は昨日公開版の訂正・再掲です。一部レースの馬番に誤りがあったため、")
note.append("　確定した枠順で全レースを再計算しています。申し訳ありませんでした。")
note.append("")
note.append("こんにちは、アスメシ競馬予想です🍱")
note.append("8月2日（日）札幌・新潟・中京の全レースをビッグデータで分析し、")
note.append(f"自信度7以上の{len(keys)}レース（重賞2本含む）を深掘り分析と買い目まで1本にまとめました。")
note.append("")
note.append("昨日8/1は「1着馬が印内」が8レース中6レース（75%）、中京8Rで3連複11.7倍を的中。")
note.append("一方でデータの少ない若いレースの買い方に課題が出たため、本日は")
note.append("「若いレースは1点勝負をやめて手広く構える」改良版の買い目で臨みます。")
note.append("")
note.append("---")
note.append("")
note.append("本日の自信度ランキング")
note.append("")
for k, c, tags, recs, top, name, desc, res in race_data:
    grade = f"[{top['グレード']}]" if top['グレード'] else ""
    note.append(f"自信度{c}/10{'★'*c}　{k[1]}{k[2]}R {top['レース名']}{grade} {top['距離']} {top['頭数']}頭")
note.append("")
note.append("---")
note.append("")

for k, c, tags, recs, top, name, desc, res in race_data:
    grade = f"[{top['グレード']}]" if top['グレード'] else ""
    marks = [r for r in recs if r['AI印']]
    h_rec = next(r for r in recs if r['AI印'] == '◎')
    o_rec = next((r for r in recs if r['AI印'] == '○'), None)
    s_rec = next((r for r in recs if r['AI印'] == '▲'), None)
    d_recs = [r for r in recs if r['AI印'] == '△']

    note.append(f"{k[1]}{k[2]}R　{top['レース名']}{grade}　{top['距離']}　{top['頭数']}頭立て　"
                 f"自信度{c}/10 {'★'*c}")
    note.append("")
    note.append("予想印")
    for r in marks:
        note.append(f"　{r['AI印']} {r['馬番']}番 {r['馬名']}　総合指数{r['総合指数']}"
                     f"（独自{r['独自指数']}・ML{r['ML能力%']}%）　脚質{r.get('脚質','?')}")
    note.append("")
    note.append("自信度の根拠")
    for tg in tags:
        note.append(f"・{tg}")
    note.append("")
    note.append(f"◎{h_rec['馬番']}番{h_rec['馬名']}を本命にする理由")
    for line in honmei_reasoning(h_rec, o_rec if o_rec else h_rec):
        note.append(f"　{line}")
    note.append("")
    if o_rec:
        note.append(rival_line(o_rec, "○対抗"))
    if s_rec:
        note.append(rival_line(s_rec, "▲穴"))
    for i, d in enumerate(d_recs, 1):
        note.append(rival_line(d, f"△相手{i}"))
    note.append("")
    if top['グレード']:
        note.append("※このレースは重賞専用記事で最終追い切り評価・陣営コメント・全頭診断まで")
        note.append("　詳しく解説しています。あわせてご覧ください。")
        note.append("")
    note.append("買い目")
    if name is None:
        reason = "若いレースで幅広く構えても期待値が確保できない" if is_young_race(top.get('レース名','')) \
                 else "期待値が確保できない"
        note.append(f"見送り（{reason}と機械判定）。オッズが想定と大きくズレた場合のみ再検討します。")
    else:
        budget = BUDGET_BY_CONF[c]
        note.append(f"選択パターン: {name}（{desc}）")
        note.append(f"予算配分 {res['total']:,}円")
        for btype, idxs, amount, odds in res['bets']:
            sep = '→' if btype in ('馬単', '3連単') else '-'
            nums = sep.join(str(recs[i2]['馬番']) for i2 in idxs)
            names = "・".join(recs[i2]['馬名'] for i2 in idxs)
            note.append(f"　{btype} {nums}（{names}）　{amount:,}円　想定{odds:.1f}倍")
        note.append(f"評価: 期待回収率{res['E_rate']*100:.0f}%　回収率300%達成確率{res['P_target']*100:.0f}%")
    note.append("")
    note.append("---")
    note.append("")

note.append("本日の参戦まとめ")
note.append("")
note.append(f"自信度7以上{len(keys)}レース中、{n_join}レースに参加（総投資{total_amt:,}円）、"
             f"{len(keys)-n_join}レースは期待値基準未達で見送りです。")
note.append("昨日の教訓を反映し、若いレースは幅広い券種構成に、妙味狙いの相手選びは")
note.append("実際のオッズを確認してから最終決定する運用に改めました。")
note.append("")
note.append("メインは重賞2本。クイーンSは週末929頭の分析1位フェスティバルヒル、")
note.append("アイビスサマーダッシュは千直最強の8枠を引いたアメリカンステージから勝負します。")
note.append("")
note.append("予想が参考になったらフォロー＆いいねお願いします！")
note.append("みなさんの昼飯代になりますように🍜🎉")
note.append("")
note.append("#競馬予想 #AI予想 #JRA #ビッグデータ競馬 #クイーンステークス #アイビスサマーダッシュ")

box(["【コピー用・そのままNoteに貼り付け可（マークダウン記号なし）】"] + note, size=9.5)

doc.save(str(OUT))
print(f"保存完了: {OUT}")
print(f"対象: {len(keys)}レース / 参加{n_join} / 総投資{total_amt:,}円 / 個別スレッド{len(top_races)}本")
