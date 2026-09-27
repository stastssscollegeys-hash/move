# -*- coding: utf-8 -*-
"""
アスメシ競馬予想 SNS投稿案 — 本日 自信度7以上 注目レース ダイジェスト
2026/08/15(土) 当日朝版
17ファクター独自指数×LightGBM v44×累積DB11,143レコードのビッグデータから
自信度7以上（10段階）と判定されたレースをSNS投稿用にまとめる。
※当日朝（10時台）投稿のため、発走済みの中京1R・新潟1Rは除外。
　明日の札幌記念（G2）・中京記念（G3）は別途詳細版あり。
Note記事のみ買い目付き（v5 EVアダプティブと同一ロジック）。
"""
import sys, json
sys.stdout.reconfigure(encoding='utf-8')
from pathlib import Path
from collections import defaultdict
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from gen_weekend_confidence_report import calc_confidence
from gen_kaime_v5 import choose_pattern, BUDGET_BY_CONF

JSON_PATH = r"C:\Users\User\Desktop\競馬予想レポート\20260815\週末ビッグデータ_20260815-0815_records.json"
OUT = Path(r"C:\Users\User\Desktop\競馬予想レポート\20260815\週末自信度7以上_SNS投稿案.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

FONT = '游ゴシック'
CONF_MIN = 7
# 発走済みレース（当日朝10時台生成のため除外）
SKIP_KEYS = {('20260815', '中京', 1), ('20260815', '新潟', 1)}
# まもなく発走の注意書き対象
SOON_KEYS = {('20260815', '札幌', 2)}

def set_font(run, size=10, bold=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

def h1(doc, text, color=(15,71,97)):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        set_font(r, 14); r.font.color.rgb = RGBColor(*color)

def h3(doc, text):
    p = doc.add_heading(text, level=3)
    for r in p.runs:
        set_font(r, 11)

def box(doc, lines, size=10):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("\n".join(lines))
    set_font(r, size)


def load_races():
    with open(JSON_PATH, encoding='utf-8') as f:
        data = json.load(f)
    races = defaultdict(list)
    for r in data['records']:
        races[(r['date'], r['競馬場'], r['R'])].append(r)
    for k in races:
        races[k].sort(key=lambda x: x['AI予測順位'])
    conf = {k: calc_confidence(v) for k, v in races.items()}
    return races, conf


def soon_tag(k):
    return " ⚠まもなく発走" if k in SOON_KEYS else ""


def race_line(k, races, conf, with_stars=True):
    recs = races[k]
    top = recs[0]
    c, _ = conf[k]
    grade = f"[{top['グレード']}]" if top['グレード'] and top['グレード'] not in top['レース名'] else ""
    stars = f" 自信度{c}/10{'★'*c}" if with_stars else ""
    header = f"■{k[1]}{k[2]}R {top['レース名']}{grade} {top['距離']} {top['頭数']}頭{stars}{soon_tag(k)}"
    marks = [r for r in recs if r['AI印']]
    mark_line = "　".join(f"{r['AI印']}{r['馬番']}{r['馬名']}" for r in marks)
    return header, mark_line


def build_day_post(day, day_label, races, conf, keys):
    lines = [
        f"🏇【本日8/{day[6:]}（{day_label}）の自信度TOP注目レース】",
        "━━ 札幌・新潟・中京 全レース中、自信度7以上（10段階）だけ厳選 ━━",
        "",
        "17ファクター独自指数×LightGBM ML×累積DB11,143レコードのビッグデータで、",
        "確定枠順まで反映した本日の最終予想印です👇",
        "",
    ]
    grouped = defaultdict(list)
    for k in keys:
        grouped[k[1]].append(k)
    for venue in sorted(grouped, key=lambda v: -max(conf[k][0] for k in grouped[v])):
        vkeys = sorted(grouped[venue], key=lambda k: -conf[k][0])
        lines.append(f"【{venue}競馬場】")
        for k in vkeys:
            header, mark_line = race_line(k, races, conf)
            lines.append(header)
            lines.append(f"　{mark_line}")
        lines.append("")
    lines += [
        "明日の札幌記念（G2）・中京記念（G3）は、枠順確定を反映した",
        "全頭診断＋買い目を別記事で詳述しています🎯",
        "",
        "#競馬予想 #AI予想 #JRA #ビッグデータ競馬",
    ]
    return lines


def build_race_thread(k, races, conf):
    """自信度9-10の特に自信のあるレースは個別スレッドとして厚めに。"""
    recs = races[k]
    top = recs[0]
    c, tags = conf[k]
    grade = f"[{top['グレード']}]" if top['グレード'] and top['グレード'] not in top['レース名'] else ""
    marks = [r for r in recs if r['AI印']]
    lines = [
        f"🏇【{k[1]}{k[2]}R {top['レース名']}{grade}】自信度{c}/10{'★'*c}{soon_tag(k)}",
        f"━━ {top['距離']} {top['頭数']}頭立て / 本日2026/08/{k[0][6:]} ━━",
        "",
    ]
    for r in marks:
        lines.append(f"{r['AI印']} {r['馬番']}番 {r['馬名']}　総合指数{r['総合指数']}（独自{r['独自指数']}・ML{r['ML能力%']}%）")
    lines.append("")
    lines.append("【自信度の根拠】")
    for t in tags:
        lines.append(f"・{t}")
    lines.append("")
    lines.append(f"#競馬予想 #AI予想 #JRA #{k[1]}競馬場")
    return lines


def build_threads_venue_posts(day, day_label, races, conf, keys):
    """Threads用: 競馬場ごとに分割した投稿（各500字以内・超える場合はさらに分割）"""
    grouped = defaultdict(list)
    for k in keys:
        grouped[k[1]].append(k)
    posts = []
    for venue in sorted(grouped, key=lambda v: -max(conf[k][0] for k in grouped[v])):
        vkeys = sorted(grouped[venue], key=lambda k: -conf[k][0])
        race_blocks = []
        for k in vkeys:
            recs = races[k]
            top = recs[0]
            c, _ = conf[k]
            grade = f"[{top['グレード']}]" if top['グレード'] and top['グレード'] not in top['レース名'] else ""
            marks = "　".join(f"{r['AI印']}{r['馬番']}{r['馬名']}" for r in recs if r['AI印'])
            race_blocks.append(f"■{k[2]}R {top['レース名']}{grade} 自信度{c}/10{soon_tag(k)}\n{marks}")

        def make_post(blocks, part=None):
            part_tag = f" その{part}" if part else ""
            lines = [
                f"🏇【本日8/{day[6:]}（{day_label}）{venue}の注目レース{part_tag}】",
                "━━ AI予想・自信度7以上だけ厳選 ━━",
                "",
            ]
            for b in blocks:
                lines.append(b)
                lines.append("")
            lines += ["買い目はNote記事で全公開🎯", "", "#競馬予想 #AI予想 #JRA"]
            return lines

        # まず1投稿で試し、500字超なら半分に分割
        one = make_post(race_blocks)
        if len("\n".join(one)) <= 500:
            posts.append((f"{day_label}・{venue}", one))
        else:
            mid = (len(race_blocks) + 1) // 2
            posts.append((f"{day_label}・{venue} その1", make_post(race_blocks[:mid], 1)))
            posts.append((f"{day_label}・{venue} その2", make_post(race_blocks[mid:], 2)))
    return posts


def build_kaime_lines(recs, c):
    """Note記事用: v5 EVアダプティブ買い目（gen_kaime_v5と同一ロジック・推定オッズ版）"""
    budget = BUDGET_BY_CONF[c]
    name, desc, res, ranking, p, q = choose_pattern(recs, budget)
    if name is None:
        best_line = ""
        if ranking:
            b = ranking[0]
            best_line = f"（最良パターンでも期待回収率{b[2]['E_rate']*100:.0f}%と基準115%未満）"
        return [f"　💰買い目: 見送り{best_line}。オッズが想定より付けば再判定"]
    lines = [f"　💰買い目（{name} / 予算{res['total']:,}円・自信度{c}連動）"]
    for btype, idxs, amount, odds in res['bets']:
        joiner = "→" if btype in ('馬単', '3連単') else "-"
        nums = joiner.join(str(recs[i]['馬番']) for i in idxs)
        lines.append(f"　　{btype} {nums}　{amount:,}円　想定{odds:.1f}倍")
    lines.append(f"　　→ 期待回収率{res['E_rate']*100:.0f}% / 300%超の確率{res['P_target']*100:.0f}% / ガミ率{res['gami_ratio']*100:.0f}%")
    return lines


def main():
    races, conf = load_races()
    all_keys = [k for k in races if conf[k][0] >= CONF_MIN and k not in SKIP_KEYS]
    skipped = [k for k in races if conf[k][0] >= CONF_MIN and k in SKIP_KEYS]
    all_keys.sort(key=lambda k: (k[0], -conf[k][0], k[1], k[2]))

    keys_by_day = defaultdict(list)
    for k in all_keys:
        keys_by_day[k[0]].append(k)

    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.8); s.bottom_margin = Cm(1.8)
        s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)

    t = doc.add_heading("本日 自信度7以上 注目レース SNS投稿案", level=0)
    for r in t.runs: set_font(r, 16)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("2026/08/15(土)　全36レース487頭から厳選　17ファクター独自指数×ML v44×累積DB11,143レコード")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in sub.runs: set_font(r, 10)
    note = doc.add_paragraph("※当日朝版: 発走済みの中京1R・新潟1R（いずれも自信度7以上）は投稿対象から除外済み。札幌2Rはまもなく発走のため投稿タイミングに注意")
    for r in note.runs: set_font(r, 9, color=(180, 0, 0))
    doc.add_paragraph()

    weekday = {'20260815': '土'}

    # ── X投稿: 日別ダイジェスト ─────────────────────────
    h1(doc, "【X投稿】本日ダイジェスト（コピー用）", color=(20,100,50))
    for day in sorted(keys_by_day):
        h3(doc, f"{day[4:6]}/{day[6:]}（{weekday.get(day,'')}）版")
        lines = build_day_post(day, weekday.get(day, ''), races, conf, keys_by_day[day])
        box(doc, ["【コピー用】"] + lines)

    doc.add_paragraph()

    # ── X投稿: 自信度9-10 個別スレッド ──────────────────
    h1(doc, "【X投稿】自信度9〜10 個別スレッド（特に自信のあるレース）", color=(20,100,50))
    top_keys = [k for k in all_keys if conf[k][0] >= 9]
    for i, k in enumerate(top_keys, 1):
        h3(doc, f"スレッド{i}/{len(top_keys)}")
        lines = build_race_thread(k, races, conf)
        box(doc, ["【コピー用】"] + lines)

    doc.add_paragraph()

    # ── Threads投稿（競馬場別・500字以内） ─────────────────
    h1(doc, "【Threads投稿】競馬場別ダイジェスト（各500字以内）", color=(80,40,120))
    for day in sorted(keys_by_day):
        for title, lines in build_threads_venue_posts(day, weekday.get(day, ''), races, conf, keys_by_day[day]):
            text = "\n".join(lines)
            n = len(text)
            status = "OK" if n <= 500 else "⚠超過!要短縮"
            h3(doc, f"{day[4:6]}/{day[6:]} {title}（{n}字/500字 {status}）")
            if n > 500:
                print(f"[WARN] Threads文字数超過: {day} {title} = {n}字")
            box(doc, ["【コピー用】"] + lines)

    doc.add_paragraph()

    # ── Note記事 ──────────────────────────────────────────
    h1(doc, "【Note記事全文】自信度7以上 全レース詳細（買い目付き）", color=(180,80,0))
    note_lines = [
        "本日のビッグデータ予想 自信度TOPレース特集 — アスメシ競馬予想", "",
        "こんにちは、アスメシ競馬予想です🍱", "",
        "本日8/15（土）の札幌・新潟・中京、全36レース487頭を",
        "17ファクター独自指数×LightGBM ML v44×累積DB11,143レコードのビッグデータで分析し、",
        "確定枠順まで反映した最終予想から、自信度10段階評価で7以上のレースだけを厳選しました。", "",
        "自信度は「上位馬との指数差」「独自指数とML能力の一致」「累積DBの実績裏付け」",
        "「頭数」などから機械的に算出しています。", "",
        "各レースには馬券の買い目も掲載します。買い目は14種のパターンを機械評価し、",
        "期待回収率115%以上×ガミ率35%以下を満たす中から「回収率300%超の確率」最大の",
        "パターンを自動選択したものです（基準を満たさないレースは正直に「見送り」）。",
        "予算は自信度連動（10=15,000円/9=12,000円/8=8,000円/7=5,000円）、",
        "オッズは朝時点の推定値です。", "",
        "---", "",
    ]
    for day in sorted(keys_by_day):
        note_lines.append(f"{day[4:6]}/{day[6:]}（{weekday.get(day,'')}）")
        note_lines.append("")
        for k in keys_by_day[day]:
            recs = races[k]
            top = recs[0]
            c, tags = conf[k]
            grade = f"[{top['グレード']}]" if top['グレード'] and top['グレード'] not in top['レース名'] else ""
            note_lines.append(f"{k[1]}{k[2]}R {top['レース名']}{grade} {top['距離']} {top['頭数']}頭　自信度{c}/10 {'★'*c}{soon_tag(k)}")
            marks = [r for r in recs if r['AI印']]
            for r in marks:
                note_lines.append(
                    f"　{r['AI印']} {r['馬番']}番 {r['馬名']}　総合指数{r['総合指数']}（独自{r['独自指数']}・ML{r['ML能力%']}%・{r.get('脚質','?')}）")
            note_lines.append(f"　根拠: {' / '.join(tags)}")
            note_lines.extend(build_kaime_lines(recs, c))
            note_lines.append("")
        note_lines.append("---")
        note_lines.append("")
    note_lines += [
        "明日8/16（日）は札幌記念（G2）と中京記念（G3）。枠順確定・確定騎手・買い目設計まで",
        "反映した全頭診断を別記事で詳しく解説しています。", "",
        "予想が参考になったらフォロー＆いいねお願いします！",
        "みなさんの週末が楽しいものになりますように🍜🎉", "",
        "#競馬予想 #AI予想 #JRA #ビッグデータ競馬 #札幌競馬場 #新潟競馬場 #中京競馬場",
    ]
    box(doc, note_lines, size=9.5)

    doc.save(str(OUT))
    print(f"保存完了: {OUT}")
    print(f"対象: {len(all_keys)}レース（発走済み除外: {len(skipped)}レース）")
    print(f"個別スレッド化（自信度9-10）: {len(top_keys)}レース")


if __name__ == '__main__':
    main()
