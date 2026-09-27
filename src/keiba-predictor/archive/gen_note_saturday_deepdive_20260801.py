# -*- coding: utf-8 -*-
"""
アスメシ競馬予想 Note記事案 — 8/1(土) 自信度7以上レース徹底解説
週末自信度レポート×v5買い目（46パターン機械評価）を1本のNote記事にがっちゃんこ。
各レースの深掘り分析（データ根拠・自信度の理由）＋詳細買い目（券種・金額・想定オッズ・
E[回収率]・P(300%達成)）まで全て統合した最終版Note記事案。
"""
import sys, json
sys.stdout.reconfigure(encoding='utf-8')
from pathlib import Path
from collections import defaultdict
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from gen_weekend_confidence_report import calc_confidence, honmei_reasoning, FACTOR_LABELS, FKEYS
from gen_kaime_v5 import choose_pattern, BUDGET_BY_CONF

TARGET_DATE = "20260801"
JSON_PATH = r"C:\Users\User\Desktop\競馬予想レポート\20260801\週末ビッグデータ_20260801-0802_records.json"
ODDS_PATH = str(Path(__file__).resolve().parent / "odds_20260802_g3.json")
OUT = Path(rf"C:\Users\User\Desktop\競馬予想レポート\{TARGET_DATE}\{TARGET_DATE}_土曜日深掘りNote記事案.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

FONT = '游ゴシック'
CONF_MIN = 7


def set_font(run, size=10, bold=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)


def h1(doc, text, color=(15, 71, 97)):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        set_font(r, 14); r.font.color.rgb = RGBColor(*color)


def h2(doc, text, color=(180, 80, 0)):
    p = doc.add_heading(text, level=2)
    for r in p.runs:
        set_font(r, 12); r.font.color.rgb = RGBColor(*color)


def box(doc, lines, size=10):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("\n".join(lines))
    set_font(r, size)


def load():
    with open(JSON_PATH, encoding='utf-8') as f:
        data = json.load(f)
    with open(ODDS_PATH, encoding='utf-8') as f:
        real_odds = json.load(f)
    races = defaultdict(list)
    for r in data['records']:
        races[(r['date'], r['競馬場'], r['R'])].append(r)
    for k in races:
        races[k].sort(key=lambda x: x['AI予測順位'])
    conf = {k: calc_confidence(v) for k, v in races.items()}
    return races, conf, real_odds


def rival_reasoning(r, label):
    """○▲△馬の一言根拠（honmei_reasoningの簡易版）。"""
    fs = sorted(((k, r.get(k, 0)) for k in FKEYS), key=lambda x: -x[1])[:2]
    strong = "・".join(f"{FACTOR_LABELS[k]}{v:.0f}点" for k, v in fs)
    parts = [f"{label}{r['馬番']}番{r['馬名']}: 総合指数{r['総合指数']}（独自{r['独自指数']}・ML{r['ML能力%']}%）"]
    parts.append(f"強み: {strong}")
    if r.get('DB最高指数', 0) >= 72:
        parts.append(f"累積DB最高指数{r['DB最高指数']}点の実績あり")
    if r.get('近5走複勝率%', 0) >= 55:
        parts.append(f"近5走複勝率{r['近5走複勝率%']:.0f}%と安定")
    elif r.get('近5走平均着', 0) >= 5.0:
        parts.append(f"近5走平均{r['近5走平均着']}着とやや低調も人気薄なら妙味")
    return "　".join(parts)


def main():
    races, conf, real_odds = load()
    keys = sorted([k for k in races if k[0] == TARGET_DATE and conf[k][0] >= CONF_MIN],
                  key=lambda k: (-conf[k][0], k[1], k[2]))

    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.8); s.bottom_margin = Cm(1.8)
        s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)

    t = doc.add_heading("8/1(土) 自信度TOPレース徹底解説 Note記事案", level=0)
    for r in t.runs: set_font(r, 16)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(f"自信度7以上 全{len(keys)}レース　深掘り分析＋v5買い目（46パターン機械評価）統合版")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in sub.runs: set_font(r, 10)
    doc.add_paragraph()

    h1(doc, "■ この記事の使い方（Word上のメモ・Noteには貼らない）")
    box(doc, [
        "以下「Note記事全文（コピー用）」を丸ごとNoteにコピペしてください。",
        "note記事のルールに従いマークダウン記号（# ## **）は使用していません。",
        "見出しはそのままテキストとして貼り付け、note側の見出し機能で装飾してください。",
    ])
    doc.add_paragraph()

    # ── データ集計 ──
    race_data = []
    for k in keys:
        recs = races[k]
        top = recs[0]
        c, tags = conf[k]
        budget = BUDGET_BY_CONF[c]
        odds_key = f"{k[0]}_{k[1]}_{k[2]}"
        name, desc, res, ranking, p, q = choose_pattern(recs, budget, real_odds.get(odds_key))
        race_data.append((k, c, tags, recs, top, name, desc, res, p, q))

    # ============================================================
    # Note記事全文の組み立て
    # ============================================================
    weekday_label = "土"
    note = []
    note.append(f"8月1日({weekday_label}) 自信度TOPレース徹底解説 — アスメシ競馬予想")
    note.append("")
    note.append("はじめに")
    note.append("")
    note.append("こんにちは、アスメシ競馬予想です🍱")
    note.append(f"8月1日（{weekday_label}）札幌・新潟・中京の全レースを、17ファクター独自指数×")
    note.append("LightGBM v44 ML能力×累積DB9,259戦のビッグデータで分析し、10段階の自信度評価を")
    note.append(f"付けました。その中から自信度7以上の{len(keys)}レースを、深掘り分析と詳細買い目まで")
    note.append("1本の記事にまとめてお届けします。")
    note.append("")
    note.append("自信度は「上位馬との指数差」「独自指数とML能力の一致」「累積DBの実績裏付け」")
    note.append("「頭数」から機械的に算出。買い目は46パターンの券種ライブラリを全レース分機械評価し、")
    note.append("E[回収率]115%以上×ガミ率35%以下という基準を満たすパターンの中から、")
    note.append("回収率300%達成確率が最大のものを自動選択しています。")
    note.append("")
    note.append("---")
    note.append("")
    note.append("本日の自信度ランキング")
    note.append("")
    for k, c, tags, recs, top, name, desc, res, p, q in race_data:
        grade = f"[{top['グレード']}]" if top['グレード'] else ""
        note.append(f"自信度{c}/10{'★'*c}　{k[1]}{k[2]}R {top['レース名']}{grade} "
                     f"{top['距離']} {top['頭数']}頭")
    note.append("")
    note.append("---")
    note.append("")

    for k, c, tags, recs, top, name, desc, res, p, q in race_data:
        grade = f"[{top['グレード']}]" if top['グレード'] else ""
        marks = [r for r in recs if r['AI印']]
        h_rec = next(r for r in recs if r['AI印'] == '◎')
        o_rec = next((r for r in recs if r['AI印'] == '○'), None)
        s_rec = next((r for r in recs if r['AI印'] == '▲'), None)
        d_recs = [r for r in recs if r['AI印'] == '△']
        hi = recs.index(h_rec)

        note.append(f"{k[1]}{k[2]}R　{top['レース名']}{grade}　{top['距離']}　{top['頭数']}頭立て　"
                     f"自信度{c}/10 {'★'*c}")
        note.append("")

        note.append("予想印")
        for r in marks:
            note.append(f"　{r['AI印']} {r['馬番']}番 {r['馬名']}　総合指数{r['総合指数']}"
                         f"（独自{r['独自指数']}・ML{r['ML能力%']}%）　脚質{r.get('脚質','?')}")
        note.append("")

        note.append("自信度の根拠")
        for tag in tags:
            note.append(f"・{tag}")
        note.append("")

        note.append(f"◎{h_rec['馬番']}番{h_rec['馬名']}を本命にする理由")
        for line in honmei_reasoning(h_rec, o_rec if o_rec else h_rec):
            note.append(f"　{line}")
        note.append("")

        if o_rec:
            note.append(rival_reasoning(o_rec, "○対抗"))
        if s_rec:
            note.append(rival_reasoning(s_rec, "▲穴"))
        for i, d in enumerate(d_recs, 1):
            note.append(rival_reasoning(d, f"△相手{i}"))
        note.append("")

        note.append("買い目")
        if name is None:
            note.append("EV基準未達のため見送り。人気が想定より大きくズレた場合のみ当日再判定します。")
        else:
            note.append(f"選択パターン: {name}（{desc}）")
            budget = BUDGET_BY_CONF[c]
            note.append(f"予算{budget:,}円（自信度{c}連動） 実配分{res['total']:,}円")
            for btype, idxs, amount, odds in res['bets']:
                sep = '→' if btype in ('馬単', '3連単') else '-'
                nums = sep.join(str(recs[i2]['馬番']) for i2 in idxs)
                names = "・".join(recs[i2]['馬名'] for i2 in idxs)
                note.append(f"　{btype} {nums}（{names}）　{amount:,}円　想定{odds:.1f}倍")
            note.append(f"評価: E[回収率]{res['E_rate']*100:.0f}%　"
                         f"P(回収率300%達成){res['P_target']*100:.0f}%　"
                         f"的中率{res['P_hit']*100:.0f}%　ガミ率{res['gami_ratio']*100:.0f}%")
        note.append("")
        note.append("---")
        note.append("")

    total_amt = sum(r[7]['total'] for r in race_data if r[7])
    n_join = sum(1 for r in race_data if r[5] is not None)
    note.append("本日のまとめ")
    note.append("")
    note.append(f"自信度7以上{len(keys)}レース中、EV基準を満たした{n_join}レースに参加、"
                 f"総投資額は{total_amt:,}円です。")
    note.append("見送りとなったレースはEV（期待値）が確保できないと機械判定されたレースであり、")
    note.append("「買わないことも回収率を守るための選択」というスタンスで運用しています。")
    note.append("")
    note.append("予想が参考になったらフォロー＆いいねお願いします！")
    note.append("みなさんの週末が楽しいものになりますように🍜🎉")
    note.append("")
    note.append("#競馬予想 #AI予想 #JRA #ビッグデータ競馬 #札幌競馬場 #新潟競馬場 #中京競馬場")

    h1(doc, "■ Note記事全文（コピー用）", color=(20, 100, 50))
    box(doc, ["【コピー用・そのままNoteに貼り付け可】"] + note, size=9.5)

    doc.save(str(OUT))
    print(f"保存完了: {OUT}")
    print(f"対象: {len(keys)}レース（うち買い目参加{n_join}）/ 総投資{total_amt:,}円")


if __name__ == '__main__':
    main()
