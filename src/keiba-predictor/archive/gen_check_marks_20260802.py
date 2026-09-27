# -*- coding: utf-8 -*-
"""
アスメシ競馬予想 — 8/2(日) 注目レース（自信度7以上）予想×結果 照合レポート+SNS投稿案
事前公開した印（訂正版）と実結果を全レース照合し、
照合表 + X/Threads結果報告投稿を1本のWordに出力する。
"""
import sys, json
sys.stdout.reconfigure(encoding='utf-8')
from pathlib import Path
from collections import defaultdict
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

SRC = Path(__file__).resolve().parent
sys.path.insert(0, str(SRC))
from gen_weekend_confidence_report import calc_confidence

TARGET_DATE = "20260802"
JSON_PATH = r"C:\Users\User\Desktop\競馬予想レポート\20260801\週末ビッグデータ_20260801-0802_records.json"
EVAL_PATH = rf"C:\Users\User\Desktop\競馬予想レポート\{TARGET_DATE}\v5買い目_答え合わせ_{TARGET_DATE}.json"
DB_PATH = Path.home() / 'Desktop' / '競馬予想レポート' / 'daily_pdca' / 'db' / 'race_results.json'
OUT = Path(rf"C:\Users\User\Desktop\競馬予想レポート\{TARGET_DATE}\{TARGET_DATE}_注目レース結果照合SNS投稿案.docx")
FONT = '游ゴシック'

# ── データ読込 ──────────────────────────────────────────
with open(JSON_PATH, encoding='utf-8') as f:
    data = json.load(f)
races = defaultdict(list)
for r in data['records']:
    races[(r['date'], r['競馬場'], r['R'])].append(r)
for k in races:
    races[k].sort(key=lambda x: x['AI予測順位'])
conf = {k: calc_confidence(v) for k, v in races.items()}
keys = sorted([k for k in races if k[0] == TARGET_DATE and conf[k][0] >= 7],
              key=lambda k: (-conf[k][0], k[1], k[2]))
# アイビスSD（再計算で自信度7未満だが公開・投票済みのため照合対象に含める）
isd_key = (TARGET_DATE, '新潟', 7)
if isd_key not in keys:
    keys.append(isd_key)

with open(DB_PATH, encoding='utf-8') as f:
    db = json.load(f)
results = defaultdict(dict)
for e in db:
    if e.get('date') != TARGET_DATE:
        continue
    try:
        ub = int(float(e.get('馬番', 0)))
    except (TypeError, ValueError):
        continue
    results[(e['競馬場'], int(e['R']))][ub] = e

with open(EVAL_PATH, encoding='utf-8') as f:
    ev = json.load(f)
bet_by_race = {r['race']: r for r in ev['races']}
bet_by_race['新潟7R'] = {'pattern': 'アイビスSD多層12点', 'inv': 10000, 'ret': 0}

# ── 照合 ────────────────────────────────────────────────
rows = []
n_win_in_marks = 0
n_two_top3 = 0
hon_stats = defaultdict(int)
for k in keys:
    recs = races[k]
    top = recs[0]
    c, _ = conf[k]
    # アイビスSDの公開印（当日最終版）
    if k == isd_key:
        marks = [('◎', 16, 'アメリカンステージ'), ('○', 9, 'カウンターセブン'),
                 ('▲', 6, 'ピューロマジック'), ('△', 10, 'エコロレジーナ'), ('△', 2, 'デュガ')]
        c = '-'
    else:
        marks = [(r['AI印'], int(r['馬番']), r['馬名']) for r in recs if r['AI印']]
    res_map = results.get((k[1], k[2]), {})
    finish = sorted(
        [(int(e['着順int']), ub, e['馬名'], e.get('単勝オッズ', 0), e.get('人気', 0))
         for ub, e in res_map.items() if e.get('着順int', 99) < 90])
    top3 = finish[:3]
    top3_nums = [ub for _, ub, *_ in top3]
    mark_of = {ub: m for m, ub, _ in marks}
    win_mark = mark_of.get(top3_nums[0]) if top3 else None
    in_top3 = [(m, ub, nm) for m, ub, nm in marks if ub in top3_nums]
    if win_mark:
        n_win_in_marks += 1
    if len(in_top3) >= 2:
        n_two_top3 += 1
    hon = next(((o, nm) for o, ub, nm, *_ in finish
                for m, mub, mnm in marks if m == '◎' and ub == mub), None)
    hon_num = next(ub for m, ub, _ in marks if m == '◎')
    hon_fin = next((o for o, ub, *_ in finish if ub == hon_num), None)
    if hon_fin == 1: hon_stats['勝利'] += 1
    elif hon_fin == 2: hon_stats['2着'] += 1
    elif hon_fin and hon_fin <= 3: hon_stats['3着'] += 1
    else: hon_stats['着外'] += 1
    bet = bet_by_race.get(f"{k[1]}{k[2]}R", {})
    rows.append({
        'k': k, 'conf': c, 'title': top['レース名'], 'grade': top['グレード'],
        'marks': marks, 'top3': top3, 'win_mark': win_mark, 'in_top3': in_top3,
        'hon_fin': hon_fin,
        'bet_pattern': bet.get('pattern', '—'), 'inv': bet.get('inv', 0), 'ret': bet.get('ret', 0),
    })

total = len(rows)
total_inv = sum(r['inv'] for r in rows)
total_ret = sum(r['ret'] for r in rows)

# ── Word出力 ────────────────────────────────────────────
doc = Document()
for s in doc.sections:
    s.top_margin = Cm(1.8); s.bottom_margin = Cm(1.8)
    s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)

def set_font(run, size=10, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

def h1(text, color=(15,71,97)):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        set_font(r, 14); r.font.color.rgb = RGBColor(*color)

def box(lines, size=9.5):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("\n".join(lines))
    set_font(r, size)

t = doc.add_heading("8/2(日) 注目レース 予想×結果 照合", level=0)
for r in t.runs: set_font(r, 16)
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph(f"自信度7以上10レース + アイビスSD = 全{total}レースの答え合わせ")
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
for r in sub.runs: set_font(r, 10)
doc.add_paragraph()

# 照合表
h1("■ 照合表（全レース）")
lines = [f"1着馬が印内: {n_win_in_marks}/{total}レース（{n_win_in_marks/total*100:.0f}%） / "
         f"印2頭以上が3着内: {n_two_top3}/{total}レース",
         f"◎の成績: 1着{hon_stats['勝利']}・2着{hon_stats['2着']}・3着{hon_stats['3着']}・着外{hon_stats['着外']}"
         f"（連対率{(hon_stats['勝利']+hon_stats['2着'])/total*100:.0f}%）",
         f"買い目: 投資{total_inv:,}円 / 回収{total_ret:,}円（{total_ret/total_inv*100:.0f}%）", ""]
for r in rows:
    k = r['k']
    grade = f"[{r['grade']}]" if r['grade'] else ""
    mark_str = "　".join(f"{m}{ub}{nm}" for m, ub, nm in r['marks'])
    top3_str = " / ".join(f"{o}着{ub}番{nm}" for o, ub, nm, *_ in r['top3'])
    hit_str = "・".join(f"{m}{nm}{next(o for o,ub2,_,*_ in r['top3'] if ub2==ub)}着"
                        for m, ub, nm in r['in_top3']) or "印は3着内なし"
    bet_str = (f"{r['bet_pattern']} {r['inv']:,}円→{r['ret']:,}円"
               if r['inv'] else "見送り")
    lines += [
        f"◆{k[1]}{k[2]}R {r['title']}{grade}（自信度{r['conf']}）",
        f"　予想: {mark_str}",
        f"　結果: {top3_str}",
        f"　照合: 1着は{'印内(' + r['win_mark'] + ')' if r['win_mark'] else '印外'} / 3着内の印: {hit_str}",
        f"　買い目: {bet_str}",
        "",
    ]
box(lines, 9)

# SNS投稿
h1("【X投稿】注目レース結果報告（コピー用）", color=(20,100,50))
x_lines = [
    "【コピー用】",
    "🏇【8/2(日) 注目レース結果報告】",
    f"━━ 自信度7以上+重賞 全{total}レースの答え合わせ ━━",
    "",
    "本日の注目レースの結果を正直にご報告します📝",
    "",
    "─────────────────────",
    "📊 予想印の精度",
    "─────────────────────",
    f"・1着馬が印の中にいたレース: {n_win_in_marks}/{total}（{n_win_in_marks/total*100:.0f}%）",
    f"・印2頭以上が3着内: {n_two_top3}/{total}レース",
    f"・◎の連対率: {(hon_stats['勝利']+hon_stats['2着'])/total*100:.0f}%"
    f"（1着{hon_stats['勝利']}回・2着{hon_stats['2着']}回）",
    "",
    "─────────────────────",
    "✅ 的中3レース",
    "─────────────────────",
    "中京1R 3連複13.7倍 → レース回収率256%🎯",
    "札幌3R 3連複9.2倍 → 回収率245%🎯",
    "中京6R 単勝的中 → 回収率140%",
    "※昨日の反省で「若いレースは幅広く」に変えた買い方が2本的中に直結",
    "",
    "─────────────────────",
    "💰 本日の収支（正直報告）",
    "─────────────────────",
    f"投資{total_inv:,}円 / 回収{total_ret:,}円（{total_ret/total_inv*100:.0f}%）",
    "重賞2本は勝ち馬を印内に収めながら、◎の大敗と大波乱で馬券は届かず。",
    "馬体重の扱いと特殊コースの評価を今夜のうちに修正済みです💪",
    "",
    "明日、今週末の全データから選んだ「来週の狙い目馬ベスト5」を発表します🎯",
    "",
    "#競馬予想 #AI予想 #JRA #結果報告 #ビッグデータ競馬",
]
box(x_lines)

h1("【Threads投稿】結果報告（コピー用）", color=(80,40,120))
th = [
    "🏇【8/2(日) 注目レース結果報告】",
    "",
    f"自信度7以上+重賞、全{total}レースの答え合わせです📝",
    "",
    f"・1着馬が印内: {n_win_in_marks}/{total}（{n_win_in_marks/total*100:.0f}%）",
    f"・◎連対率{(hon_stats['勝利']+hon_stats['2着'])/total*100:.0f}%（1着{hon_stats['勝利']}・2着{hon_stats['2着']}）",
    "・的中3レース: 3連複256%/245%・単勝140%",
    f"・収支: 投資{total_inv:,}円/回収{total_ret:,}円（{total_ret/total_inv*100:.0f}%）",
    "",
    "「若いレースは幅広く」の改良が3連複2本の的中に直結。",
    "重賞は勝ち馬を印内に収めるも◎大敗で悔しい結果に。",
    "馬体重ルールと特殊コース評価を修正して来週に臨みます💪",
    "",
    "明日は来週の狙い目馬ベスト5を発表🎯",
    "",
    "#競馬予想 #AI予想 #結果報告",
]
n_th = len("\n".join(th))
box([f"【コピー用】（{n_th}字/500字）"] + th)

doc.save(str(OUT))
print(f"保存完了: {OUT}")
print(f"1着が印内: {n_win_in_marks}/{total} / 印2頭以上top3: {n_two_top3}/{total} / "
      f"◎: 1着{hon_stats['勝利']} 2着{hon_stats['2着']} 3着{hon_stats['3着']} 着外{hon_stats['着外']}")
print(f"Threads文字数: {n_th}字")
