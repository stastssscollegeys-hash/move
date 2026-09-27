# -*- coding: utf-8 -*-
"""
gen_weekend_20260919.py — 2026/9/20 オールカマー(G2・中山芝2200) SNS投稿案（枠順確定版）
=====================================================================================
gen_weekend_20260912.py のコピー。変えたのはレース定義・入力ファイル・チェックリスト・内部参照だけ。
X6投稿＋Threads6投稿＋Note記事（1レース版12セクション）
【入力】research/score20_20260918.json・research/kaime_20260918.json・yoso_odds_20260918.json
"""
from __future__ import annotations
import json, sys
from pathlib import Path

sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, str(Path(__file__).resolve().parent))
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from kaime_format import fmt_manual
from kaime_design_v7 import hits

FONT = '游ゴシック'
BASE = Path.home() / 'Desktop' / '競馬予想レポート' / '20260920'
OUT = BASE / '20260920_オールカマー_SNS投稿案_枠順確定版.docx'
THREADS_LIMIT = 500
CIRC = '⓪①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱'
SEP = "─────────────────────"
TAGS = "#競馬予想 #AI予想 #JRA"

# 入力ファイルはコマンドラインで切り替えられる（9/12: 日曜分は別ファイルになるため）
#   例: --kaime research/kaime_20260913.json --odds odds_live_20260913.json --ens research/ensemble_20260912.json
def _arg(name, default):
    return sys.argv[sys.argv.index(name) + 1] if name in sys.argv else default


SCORE = json.load(open(BASE / _arg('--score', 'research/score20_20260918.json'), encoding='utf-8'))
KAIME = json.load(open(BASE / _arg('--kaime', 'research/kaime_20260918.json'), encoding='utf-8'))
YOSO = json.load(open(BASE / _arg('--odds', 'yoso_odds_20260918.json'), encoding='utf-8'))   # 前日オッズが出たレースは実オッズで上書き済み
YOSO_THU = json.load(open(BASE / 'yoso_odds_20260918.json', encoding='utf-8'))
ENS = {'オールカマー': []}   # 今週はYouTube予想の収集なし


def c(n):
    return CIRC[int(n)] if 0 <= int(n) <= 18 else str(n)


# ═══════════════════════════════════════════════════════════
# レース定義（事実はすべて research 配下・netkeiba戦績で確認済みのもの）
# ═══════════════════════════════════════════════════════════
RACES = [
 dict(
  key='オールカマー', name='オールカマー', grade='G2', date='2026/09/20（日）', when='日曜',
  place='中山競馬場', course='芝2200m（外回り）', field='13頭立て', start='15:45発走', sub='天皇賞・秋への前哨戦',
  target='200%前後',
  marks=[
   ('◎', 7, 'レガレイラ', '牝5・戸崎圭・56kg',
    ['独自の20ファクター指数155.7点でメンバー1位。2位とは6.4点の差があります',
     '昨年のこのレースの勝ち馬です。同じ中13週の休み明けで、1番人気に応えました',
     'エリザベス女王杯（G1）の勝ち馬。このメンバーでは実績が抜けています',
     '最終追い切りは最高評価のS。馬なりで好時計を出し、終いの伸びも良い内容でした',
     '不安は人気です。このレースの1番人気は過去10回で勝率40%・3着内率60%。4割は馬券圏外に消えています']),
   ('○', 12, 'エセルフリーダ', '牝5・武藤・55kg',
    ['独自の20ファクター指数149.3点で2位',
     '前走の中山牝馬S（G3）を6番人気で勝った中山の重賞ウイナーです',
     '最終追い切りは最高評価のS。◎と並んでメンバー最上位の評価でした',
     '斤量は55kg。牡馬の57〜58kg勢より2〜3kg軽く出られます',
     '不安は約半年（中27週）ぶりの実戦と8枠。このレースの過去10回で8枠の3着内率は14%です']),
   ('▲', 6, 'パンジャ', '牡5・丹内・57kg',
    ['独自の20ファクター指数148.8点で3位',
     '前走のサンシャインS（3勝クラス）を1番人気で勝っています。舞台は今回と同じ中山芝2200mです',
     '先行脚質で5枠。このレースの5枠は過去10回で3着内率28%と平均以上です',
     '不安は重賞が初挑戦であることと、中21週の休み明けです']),
   ('△', 9, 'コスモキュランダ', '牡5・横山武・57kg',
    ['昨年の有馬記念で12番人気ながら2着。中山の大舞台で結果を出しています',
     '前走の宝塚記念も8番人気で4着。G1で掲示板に入る力があります',
     'ただし最終追い切りの評価は低め（C）で、枠もこのレースで3着内率11%の6枠。相手までとします']),
   ('△', 13, 'ジューンテイク', '牡5・武豊・58kg',
    ['京都記念（G2）の勝ち馬。重賞の勝ち鞍はメンバー上位です',
     'ただし斤量58kgはメンバーで最も重く、枠も8枠。最終追い切りの評価も低め（C）でした']),
  ],
  ana=('🔥', 3, 'リビアングラス', '牡6・西村淳・57kg'),
  ana_detail=[
   'この馬を推す理由は3つあります。',
   '1つ目は状態です。最終追い切りはA評価。中間に何度も強めに追われ、最後は余裕のある動きでまとめています。',
   '2つ目は枠と脚質です。3枠3番を引きました。このレースの過去10回で3着以内に入った馬の脚質は逃げが34.8%と最も高く、前に行けるこの馬に合います。はっきり逃げそうな馬は多くありません。',
   '3つ目は重賞での実績です。日経新春杯（G2）では9番人気で3着に入っています。',
   'ただし予想では8番人気。このレースは過去10回で8〜9番人気が3着以内に一度も入っていません。人気がこれ以上落ちるなら評価も下げます。',
  ],
  danger=('コスモキュランダ', 9, '2番人気想定',
   ['有馬記念2着・宝塚記念4着と、G1での着順はメンバー屈指です。力は認めます。',
    'それでも軸にしない理由は2つです。',
    '1つ目は枠。6枠はこのレースの過去10回で3着内率11%。6〜8枠はそろって10%台にとどまり、内枠が明確に有利なレースです。',
    '2つ目は状態。最終追い切りの評価は低め（C）で、ポテンシャルほどの動きには見えないという見立てでした。',
    '2番人気なら、相手の評価までが妥当と判断しました。']),
  also_out=('ヴーレヴー', 4, '7番人気想定',
   '巴賞（オープン）を勝っていますが、その後はダートで16着・11着。芝の重賞で上位に入った実績がまだありません。'),
  hist=['過去10回で1番人気は勝率40%・3着内率60%',
        '4〜5番人気が10回中5勝（単勝回収率168%・364%）',
        '8〜9番人気は3着以内ゼロ（20頭中0頭）',
        '3連単の中央値は24,340円。10万円超は10回中2回',
        '3着以内の脚質は逃げ34.8%・差し30.4%・先行25.0%・追込6.5%',
        '枠は1枠46%に対して6枠・7枠11%、8枠14%'],
  outlook=[
   '秋の天皇賞へ向けた古馬の前哨戦です。',
   '過去10回で1〜3番人気が勝ったのは5回だけ。残り5回はすべて4〜5番人気が勝っています。一方で8番人気以下はほとんど馬券に絡まず、大穴よりも「中穴」が走るレースです。',
   '中山の外回り2200mは、スタート後に坂を上り、向こう正面から下って3コーナーから長く脚を使う形になりやすいコースです。3着以内の脚質は逃げ34.8%が最も高く、追込は6.5%。後ろから一気に差し切るのは難しいレースです。',
   '枠の差がはっきり出ています。1枠は3着内率46%に対し、6枠・7枠は11%、8枠は14%。今年は有力馬の⑫エセルフリーダと⑬ジューンテイクが8枠、⑨コスモキュランダが6枠に入りました。',
  ],
  caution=['○が6番人気の予想のため、検証中のルールでは買い目を見送ります。印は参考にしてください。',
           '当日の実オッズで◎と○がどちらも3番人気以内になれば、◎○▲の馬連BOXで買います。',
           '馬場の状態は当日の朝に最終判断します。'],
  type_line='買い方: 見送り（○が6番人気のため、検証中のルールに当てはまらない）',
  ana_short='最終追い切りA評価。3枠・前に行ける脚質で、日経新春杯（G2）3着の実績があります',
  x1_out=('キャントウェイト', 1, '5番人気想定',
          '1枠1番を引きましたが、近走は3勝クラスで9着・5着・3着・8着。重賞で上位に入った実績がなく、指数も8位にとどまりました'),
  danger_short='有馬記念2着・宝塚記念4着の実力馬ですが、6枠はこのレースの3着内率11%。最終追い切りの評価も低め（C）で、相手の評価までとしました',
  scen=[('◎と○で決着', [7, 12, 6]), ('◎と▲で決着', [7, 6, 9]), ('◎が飛ぶ', [12, 6, 9]), ('△が絡む', [7, 9, 13])],
 ),
]



# 買い方の説明は kaime_20260912.py がレースごとに選んだパターンから取る（9/11 レース別パターン版）
for _r in RACES:
    if _r['key'] in KAIME:      # 買い目ファイルに入っているレースだけ（日曜分など一部レースのみの時に対応）
        _r['type_line'] = KAIME[_r['key']]['type_line']
        if 'caution' in KAIME[_r['key']]:   # SNS提示用の買い目ファイル（kaime_20260919_sns.py）は注意書きも持つ
            _r['caution'] = KAIME[_r['key']]['caution']
SNS_ONLY = any('SNS提示用' in v.get('stats', '') for v in KAIME.values())


# ═══════════════════════════════════════════════════════════
# docx ヘルパー（gen_weekend_20260905.py と同じ体裁）
# ═══════════════════════════════════════════════════════════
def set_font(run, size=10, bold=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)


def h1(doc, t, col=(20, 60, 140)):
    p = doc.add_paragraph(); set_font(p.add_run(t), 13, True, col)


def h2(doc, t, col=(70, 70, 70)):
    p = doc.add_paragraph(); set_font(p.add_run(t), 11, True, col)


def box(doc, lines):
    tb = doc.add_table(rows=1, cols=1); tb.style = 'Table Grid'
    cell = tb.cell(0, 0); cell.text = ''
    for i, ln in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        set_font(p.add_run(ln), 9.5)
    doc.add_paragraph()


def note(doc, t, col=(0, 90, 160)):
    p = doc.add_paragraph(); set_font(p.add_run(t), 9, False, col)


def body(doc, t, size=10):
    for ln in t.split('\n'):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
        set_font(p.add_run(ln), size)


def table(doc, rows, widths=None):
    tb = doc.add_table(rows=0, cols=len(rows[0])); tb.style = 'Table Grid'
    for i, row in enumerate(rows):
        cells = tb.add_row().cells
        for j, v in enumerate(row):
            cells[j].text = ''
            set_font(cells[j].paragraphs[0].add_run(str(v)), 8.5, i == 0)
    doc.add_paragraph()


def post(doc, title, lines, limit=None):
    n = len("\n".join(lines))
    tag = f"（{n}字/{limit}字 {'OK' if n <= limit else '⚠超過'}）" if limit else f"（{n}字）"
    h2(doc, f"{title} {tag}")
    box(doc, lines)
    if limit and n > limit:
        print(f"[WARN] 超過 {title} {n}字")
    return n


# ═══════════════════════════════════════════════════════════
# 計算ヘルパー
# ═══════════════════════════════════════════════════════════
def sheet(r):
    return {x['uma']: x for x in SCORE[r['key']]}


def yoso(r):
    return {h['uma']: h for h in YOSO[r['key']]['horses'] if h.get('uma') and h.get('name') not in (None, '?')}


def kaime_lines(r, compact=False):
    k = KAIME[r['key']]
    if not k['bets']:
        return ["💰 買い目 このレースは見送り（買いません）"]
    order = ['3連複', '馬単', '馬連', 'ワイド']
    blocks = []
    for t in order:
        items = [b for b in k['bets'] if b['t'] == t]
        if not items:
            continue
        sep = '→' if t == '馬単' else '-'
        blocks.append((t, None, [(sep.join(c(x) for x in (b['combo'] if b['ordered'] else sorted(b['combo']))), b['amt']) for b in items]))
    # どの点が当たってもほぼ同じ額が戻る配分なので「本線」は置かない（9/11 レース別パターン版）
    lines = fmt_manual(blocks, roles="どの点が当たってもほぼ同じ額が戻るように配分")
    if compact:
        return [ln for ln in lines if not ln.startswith('役割')]
    return lines


def scen_lines(r):
    k = KAIME[r['key']]
    if not k['bets']:
        return []
    out = []
    for label, order in r['scen']:
        sc = tuple(order)
        ret = sum(b['amt'] * b['est'] for b in k['bets'] if hits(b, sc))
        pct = ret / k['budget'] * 100
        mark = ' 🎉' if pct >= 300 else (' ✅' if pct >= 100 else '')
        pre = '' if YOSO[r['key']].get('uren') else '推定'
        out.append(f"{label}（1着{c(order[0])}・2着{c(order[1])}・3着{c(order[2])}）→ {pre}{int(ret):,}円・回収率{pct:.0f}%{mark}")
    return out


def hd(r):
    return [f"🏇【{r['name']} {r['grade']}】{r['date']}",
            f"━━ {r['place']} {r['course']} / {r['sub']} / {r['field']} ━━"]


def tg(r):
    return f"{TAGS} #{r['name']}"


def odds_txt(r, uma):
    y = yoso(r).get(uma, {})
    lab = '前日' if (YOSO[r['key']].get('uren') or '前日' in YOSO[r['key']].get('source', '')) else '予想'
    return f"{lab}{y.get('pop')}番人気・{y.get('odds')}倍" if y.get('odds') else "人気未定"


def odds_note(r):
    src = YOSO[r['key']]['source']
    return f"※{src}オッズで組んでいます。当日朝の実オッズと馬場で最終調整します🐴" if (YOSO[r['key']].get('uren') or '前日' in src) \
        else "※枠順確定後の予想オッズで組んでいます。当日朝の実オッズと馬場で最終調整します🐴"


# ═══════════════════════════════════════════════════════════
# X / Threads
# ═══════════════════════════════════════════════════════════
def build_x(r):
    ps = []
    sh = sheet(r)
    hon = r['marks'][0]
    # ① 印一覧＋消し＋穴（冒頭150字で結論と数字フック）
    L = hd(r) + ["", f"{r['hist'][0]}。", f"本命は{c(hon[1])}{hon[2]}です。", "", SEP, "🎯 予想印と短評", SEP, ""]
    for mk, num, nm, who, cms in r['marks']:
        L += [f"{mk} {c(num)}{nm}（{who}）", f"　→ {cms[0]}", ""]
    a = r['ana']
    L += [f"{a[0]} {c(a[1])}{a[2]}（{a[3]}）", f"　→ {r['ana_short']}", ""]
    d = r['danger']
    o1 = r['x1_out']
    L += [SEP, "❌ 人気でも印を回さなかった馬", SEP, f"【{c(o1[1])}{o1[0]}（{o1[2]}）】", o1[3], ""]
    L += ["　各馬の詳しい根拠はスレッドで解説します👇", "", tg(r)]
    ps.append(("【X投稿①】予想印一覧＋消し＋穴", L))

    # ② 本命
    L = hd(r) + ["", SEP, f"◎ 本命 {c(hon[1])}{hon[2]}（{hon[3]}）", SEP, ""]
    L += [f"・{x}" for x in hon[4]]
    s = sh.get(hon[1])
    if s:
        L += ["", f"独自の20ファクター指数 {s['final170']}点（170点満点）"]
    L += ["", tg(r)]
    ps.append(("【X投稿②】本命◎の詳細根拠", L))

    # ③ 対抗・単穴
    L = hd(r)
    for m in r['marks'][1:3]:
        L += ["", SEP, f"{m[0]} {c(m[1])}{m[2]}（{m[3]}）", SEP, ""] + [f"・{x}" for x in m[4]]
    L += ["", tg(r)]
    ps.append(("【X投稿③】対抗○・単穴▲の詳細根拠", L))

    # ④ 穴
    L = hd(r) + ["", SEP, f"🔥 人気薄でも侮れない {c(a[1])}{a[2]}（{odds_txt(r, a[1])}）", SEP, ""]
    L += [f"　{x}" for x in r['ana_detail']] + ["", tg(r)]
    ps.append(("【X投稿④】穴馬の根拠", L))

    # ⑤ △＋軸にしない理由
    L = hd(r) + ["", SEP, "△ 相手に押さえる馬", SEP, ""]
    for mk, num, nm, who, cms in r['marks'][3:]:
        L += [f"△ {c(num)}{nm}（{who}）"] + [f"　・{x}" for x in cms] + [""]
    L += [SEP, f"❌ {c(d[1])}{d[0]}を軸にしない理由", SEP, ""] + [f"　{x}" for x in d[3]]
    if r['also_out']:
        o = r['also_out']
        L += ["", f"【{c(o[1])}{o[0]}（{o[2]}）も印なし】", f"　{o[3]}"]
    L += ["", SEP, "📊 過去データが示すこと", SEP, ""] + [f"・{x}" for x in r['hist']] + ["", tg(r)]
    ps.append(("【X投稿⑤】△の根拠＋人気馬を軸にしない理由＋過去データ", L))

    # ⑥ 買い目
    k = KAIME[r['key']]
    L = hd(r) + ["", SEP, f"🎯 このレースの買い方", SEP, "", r['type_line'], ""]
    L += [f"・{x}" for x in k['why']] + [""]
    L += kaime_lines(r)
    if k['bets']:
        lo, hi = k['ret_range']
        L += [f"当たった時の回収率: 1点的中で約{lo*100:.0f}〜{hi*100:.0f}%"]
        used = {x for b in k['bets'] for x in b['combo']}
        unused = [f"{mk}{c(num)}{nm}" for mk, num, nm, *_ in r['marks'][3:] if num not in used]
        if r['ana'][1] not in used:
            unused.append(f"🔥{c(r['ana'][1])}{r['ana'][2]}")
        if unused:
            L += [f"※{'・'.join(unused)}は印として注目していますが、この買い方では相手に入れていません"]
        L += ["", f"【当たった時の回収イメージ（{'前日オッズ' if YOSO[r['key']].get('uren') else '推定配当'}）】"] + [f"・{x}" for x in scen_lines(r)]
        L += ["", odds_note(r), "", tg(r)]
    else:
        L += ["", "※見送りも予想の一部です。当日朝の実オッズで条件に当てはまれば買い目を出します🐴", "", tg(r)]
    ps.append(("【X投稿⑥】買い方の考え方＋買い目", L))
    return ps


def build_threads(r):
    head = f"🏇【{r['name']} {r['grade']}】{r['date']}\n━━ {r['place']} {r['course']} / {r['field']} ━━"
    out = []
    t = [head, "", "🎯 予想印（1/6）", ""]
    for mk, num, nm, who, _ in r['marks']:
        t.append(f"{mk} {c(num)}{nm}")
    t += [f"{r['ana'][0]} {c(r['ana'][1])}{r['ana'][2]}", "", r['hist'][0], "", "根拠は次の投稿から👇", "", "あなたの本命はどの馬ですか？"]
    out.append(t)

    hon = r['marks'][0]
    out.append([head, "", f"◎本命 {c(hon[1])}{hon[2]}（2/6）", ""] + [f"・{x}" for x in hon[4][:4]] + ["", "この本命、どう見ますか？"])

    m1, m2 = r['marks'][1], r['marks'][2]
    out.append([head, "", f"○{c(m1[1])}{m1[2]}／▲{c(m2[1])}{m2[2]}（3/6）", "",
                f"○ {m1[4][0]}", f"　{m1[4][-1]}", "", f"▲ {m2[4][0]}", "", "この2頭、評価が割れそうです。"])

    a = r['ana']
    body_t = " ".join(r['ana_detail'][1:3])
    out.append([head, "", f"🔥穴 {c(a[1])}{a[2]}（4/6）", "", body_t[:260] + ("…" if len(body_t) > 260 else ""), "",
                "人気がないなら狙う価値ありと見ています。"])

    d = r['danger']
    out.append([head, "", f"❌{c(d[1])}{d[0]}を軸にしない理由（5/6）", "", r['danger_short'], "",
                "△ " + " ".join(f"{c(m[1])}{m[2]}" for m in r['marks'][3:]), "", "異論反論、歓迎です。"])

    k = KAIME[r['key']]
    t = [head, "", f"💰買い目（6/6）", "", r['type_line'], ""] + kaime_lines(r, compact=True)
    t += ["", "当日朝の実オッズで最終調整します🎯"]
    out.append(t)
    return [(f"【Threads {i+1}/6】", x) for i, x in enumerate(out)]


# ═══════════════════════════════════════════════════════════
# Note記事（3レース合同・12セクション・マークダウン記号なし）
# ═══════════════════════════════════════════════════════════
def build_note():
    if len(RACES) == 1:
        return build_note_single(RACES[0])
    P = []
    if len(RACES) == 3:
        P.append("【週末重賞3本 AI予想・枠順確定版】チャレンジC／セントライト記念／ローズS")
        P.append("")
        P.append("1. はじめに")
        P.append("こんにちは、アスメシ競馬予想です🍱")
        P.append("「明日の飯代」を懸けて、AIとデータで競馬に挑む予想アカウントです。")
        P.append("明日9/12（土）の阪神メイン・チャレンジカップ(G3)と、9/13（日）のセントライト記念(G2)・ローズステークス(G2)の、枠順確定版の予想をお届けします。")
        P.append("土日48レース630頭のデータに、独自の20ファクター指数（170点満点）、過去10年の同じ舞台の結果、最終追い切りの評価を重ねて分析しました。")
        P.append("")
        P.append("結論を先に言います。本命はチャレンジCが⑨マテンロウゲイル、セントライト記念が⑧ゴーイントゥスカイ、ローズSが⑦モンローウォークです。")
        P.append("セントライト記念は指数1位のバステールを本命にしませんでした。理由は最終追い切りです。")
        P.append("そして——買い目を出すのは「◎と○がどちらも3番人気以内」のレースだけにしました。過去のデータを期間ごとに分けて調べ、一番崩れにくかった条件です。今回はローズSを見送ります。")
        P.append("")
    else:
        # 2026-09-12: 日曜2重賞版など、レース数が3本でない時の冒頭（内容はRACESから生成する）
        P.append("【" + "／".join(r['name'] for r in RACES) + " AI予想】" + "・".join(f"本命{c(r['marks'][0][1])}{r['marks'][0][2]}" for r in RACES))
        P.append("")
        P.append("1. はじめに")
        P.append("こんにちは、アスメシ競馬予想です🍱")
        P.append("「明日の飯代」を懸けて、AIとデータで競馬に挑む予想アカウントです。")
        P.append(f"{RACES[0]['when']}{RACES[0]['date']}の" + "・".join(f"{r['place'][:2]}メイン・{r['name']}({r['grade']})" for r in RACES) + "の予想をお届けします。")
        P.append("独自の20ファクター指数（170点満点）に、同じ舞台の過去の結果、最終追い切りの評価、出走馬全頭の直近の全戦績を重ねて分析しました。")
        P.append("")
        P.append("結論を先に言います。本命は" + "、".join(f"{r['name']}が{c(r['marks'][0][1])}{r['marks'][0][2]}" for r in RACES) + "です。")
        buy = [r for r in RACES if KAIME[r['key']]['bets']]
        skip = [r for r in RACES if not KAIME[r['key']]['bets']]
        P.append("買い目を出すのは「◎と○がどちらも3番人気以内」のレースだけにしています。過去のデータを期間ごとに分けて調べ、一番崩れにくかった条件です。"
                 + ("今回は" + "・".join(r['name'] for r in skip) + "が条件に当てはまらないため見送ります。" if skip else "")
                 + ("買うのは" + "・".join(r['name'] for r in buy) + "です。" if buy else "印だけ参考にしてください。"))
        P.append("")
    P.append("2. レース基本情報")
    for r in RACES:
        P.append(f"{r['name']}({r['grade']})　{r['date']} {r['start']}　{r['place']} {r['course']}　{r['field']}　{r['sub']}")
    P.append("")
    P.append("3. コース特性と今回の狙い")
    for r in RACES:
        P.append(f"【{r['name']}】")
        P += r['outlook']
        P.append("過去データのポイント")
        P += [f"・{x}" for x in r['hist']]
        P.append("")
    P.append("4. 枠順確定後の予想オッズ一覧（全頭）")
    for r in RACES:
        P.append(f"【{r['name']}】（{YOSO[r['key']]['source']}オッズ。当日の実オッズとは異なります）")
        mk = {m[1]: m[0] for m in r['marks']}; mk[r['ana'][1]] = '🔥'
        for u, y in sorted(yoso(r).items(), key=lambda kv: kv[1]['pop'] or 99):
            P.append(f"{mk.get(u, '　')} {c(u)} {y['name']}（{y['jockey']}）{y['odds']}倍・{y['pop']}番人気")
        P.append("")
    P.append("5. 最終追い切り評価（上位）")
    for r in RACES:
        top = [x for x in SCORE[r['key']] if x['oikiri'] in ('S', 'A')]
        top.sort(key=lambda x: ({'S': 0, 'A': 1}[x['oikiri']], -x['final170']))
        P.append(f"【{r['name']}】" + " ／ ".join(f"{x['oikiri']}評価 {c(x['uma'])}{x['name']}" for x in top))
    P.append("評価はS（最上位）・A（上位）の馬だけを載せています。2つ以上の見立てが揃った馬だけを印の判断に使いました。")
    P.append("")
    P.append("6. オッズ動向（木曜の予想オッズ→最新のオッズ）")
    for r in RACES:
        thu = {h['name']: h for h in YOSO_THU[r['key']]['horses']}
        moves = []
        for u, y in yoso(r).items():
            t = thu.get(y['name'])
            if t and t.get('pop') and y.get('pop') and t['pop'] != y['pop']:
                moves.append((abs(t['pop'] - y['pop']), f"{c(u)}{y['name']} {t['pop']}番人気→{y['pop']}番人気"))
        moves.sort(key=lambda x: -x[0])
        P.append(f"【{r['name']}】" + (" ／ ".join(m[1] for m in moves[:5]) if moves else "大きな変動なし"))
    P.append("枠順が出たことで人気が動いています。人気が下がったことだけを理由に評価は下げません。")
    P.append("")
    P.append("7. 有力馬の詳細分析")
    for r in RACES:
        P.append(f"【{r['name']}】")
        for mk, num, nm, who, cms in r['marks'][:3]:
            P.append(f"{mk} {c(num)}{nm}（{who}）")
            P.append("".join(x + "。" if not x.endswith("。") else x for x in cms))
        a = r['ana']
        P.append(f"🔥 {c(a[1])}{a[2]}（{a[3]}）")
        P.append("".join(r['ana_detail']))
        P.append("")
    P.append("8. 全頭短評と独自の20ファクター指数（170点満点）")
    for r in RACES:
        P.append(f"【{r['name']}】")
        mk = {m[1]: m[0] for m in r['marks']}; mk[r['ana'][1]] = '🔥'
        for x in SCORE[r['key']]:
            w = x['why']
            P.append(f"{mk.get(x['uma'], '　')} {c(x['uma'])} {x['name']} {x['final170']}点　前走 {w['sc12']}／{w['sc18']}／最高実績 {w['sc19']}")
        P.append("")
    P.append("9. 独自指数ランキング TOP5")
    for r in RACES:
        P.append(f"【{r['name']}】" + " ／ ".join(f"{i}位 {c(x['uma'])}{x['name']} {x['final170']}点" for i, x in enumerate(SCORE[r['key']][:5], 1)))
    P.append("")
    P.append("10. 最終予想印まとめ")
    for r in RACES:
        sh = sheet(r)
        P.append(f"【{r['name']}】")
        for mk, num, nm, who, cms in r['marks']:
            P.append(f"{mk} {c(num)}{nm}　{sh[num]['final170']}点　{cms[0]}")
        a = r['ana']
        P.append(f"🔥 {c(a[1])}{a[2]}　{sh[a[1]]['final170']}点　{r['ana_detail'][0]}")
    P.append("")
    P.append("11. 買い目設計")
    for r in RACES:
        P.append(f"【{r['name']}】{r['type_line']}")
        P += kaime_lines(r)
        if KAIME[r['key']]['bets']:
            P.append(f"当たった時の回収イメージ（{'前日オッズ' if YOSO[r['key']].get('uren') else '推定配当'}）")
            P += [f"・{x}" for x in scen_lines(r)]
        else:
            P += [f"・{x}" for x in KAIME[r['key']]['why'][1:]]
        P.append("")
    P.append("※枠順確定後の予想オッズで組んでいます。当日朝の実オッズと馬場を見て最終調整します。")
    P.append("")
    P.append("12. まとめ")
    if len(RACES) == 3:
        P.append("・チャレンジCは阪神芝2000mの過去8回で勝ち馬がすべて1〜3番人気。3歳の⑨マテンロウゲイルを軸にしました。")
        P.append("・セントライト記念は10番人気以下が10年で一度も3着に来ていない堅いレース。追い切り最上位の⑧ゴーイントゥスカイを本命にしました。")
        P.append("・ローズSは3連単10万円超が6回中4回の荒れるレース。無敗の逃げ馬⑦モンローウォークを軸にしました。")
    else:
        for r in RACES:
            m = r['marks'][0]
            P.append(f"・{r['name']}は{r['hist'][0]}。{c(m[1])}{m[2]}を本命にしました。{m[4][0]}。")
    P.append("・買い目は " + "／".join(f"{r['name']}＝{KAIME[r['key']]['label']}" for r in RACES) + "。")
    P.append("　◎と○がどちらも3番人気以内のレースだけ買う、というルールを先週から検証しています。当たり外れも含めて毎週正直に報告します。")
    P.append("前日の夜に最終予想、当日の朝に実オッズ版をお届けします。フォローしてお待ちください🐴")
    P.append("")
    P.append("#競馬予想 #AI予想 #JRA #チャレンジカップ #セントライト記念 #ローズステークス #競馬")
    return P


def build_note_single(r):
    """1レース版のNote記事（12セクション・マークダウン記号なし）。9/11夜 チャレンジC最新版用"""
    k = KAIME[r['key']]
    sh = sheet(r)
    hon = r['marks'][0]
    live = bool(YOSO[r['key']].get('uren'))
    P = [f"【{r['name']}({r['grade']}) AI予想・最新版】本命は{c(hon[1])}{hon[2]}", ""]
    P += ["1. はじめに", "こんにちは、アスメシ競馬予想です🍱", "「明日の飯代」を懸けて、AIとデータで競馬に挑む予想アカウントです。",
          f"{r['when']}{r['date']}の{r['place'][:2]}メイン・{r['name']}({r['grade']})の最新版の予想をお届けします。",
          "独自の20ファクター指数（170点満点）、同じ舞台の過去の結果、最終追い切りの評価を重ねて分析しました。"
          + ("最新の前日オッズと、出走馬の直近の全戦績で計算し直しています。" if live else ""), "",
          f"結論を先に言います。本命は{c(hon[1])}{hon[2]}です。"
          f"買い方は{k['label'] if k['bets'] else '見送り'}。" + (("本命が1番人気に推されるため、相手を指数上位の中穴に絞りました。" if SNS_ONLY else "◎と○がどちらも3番人気以内という、先週から検証している条件に当てはまりました。") if k['bets'] else ""), ""]
    P += ["2. レース基本情報", f"{r['name']}({r['grade']})　{r['date']} {r['start']}　{r['place']} {r['course']}　{r['field']}　{r['sub']}", ""]
    P += ["3. コース特性と今回の狙い"] + r['outlook'] + ["過去データのポイント"] + [f"・{x}" for x in r['hist']] + [""]
    P.append(f"4. オッズ一覧（全頭・{YOSO[r['key']]['source']}）")
    mk = {m[1]: m[0] for m in r['marks']}; mk[r['ana'][1]] = '🔥'
    for u, y in sorted(yoso(r).items(), key=lambda kv: kv[1]['pop'] or 99):
        P.append(f"{mk.get(u, '　')} {c(u)} {y['name']}（{y['jockey']}）{y['odds']}倍・{y['pop']}番人気")
    P.append("")
    top = sorted([x for x in SCORE[r['key']] if x['oikiri'] in ('S', 'A')], key=lambda x: ({'S': 0, 'A': 1}[x['oikiri']], -x['final170']))
    P += ["5. 最終追い切り評価（上位）", " ／ ".join(f"{x['oikiri']}評価 {c(x['uma'])}{x['name']}" for x in top),
          "評価はS（最上位）・A（上位）の馬だけを載せています。", ""]
    thu = {h['name']: h for h in YOSO_THU[r['key']]['horses']}
    moves = sorted([(abs(thu[y['name']]['pop'] - y['pop']), f"{c(u)}{y['name']} {thu[y['name']]['pop']}番人気→{y['pop']}番人気")
                    for u, y in yoso(r).items() if y['name'] in thu and thu[y['name']].get('pop') and y.get('pop') and thu[y['name']]['pop'] != y['pop']], key=lambda x: -x[0])
    P += ["6. オッズ動向", "日曜のレースの馬券は土曜に発売されます。現時点は枠順確定後の予想オッズのため、動きは当日朝の最終版でお伝えします。",
          "人気が下がったことだけを理由に評価は下げません。", ""]
    P.append("7. 有力馬の詳細分析")
    for m, num, nm, who, cms in r['marks'][:3]:
        P += [f"{m} {c(num)}{nm}（{who}）", "".join(x if x.endswith("。") else x + "。" for x in cms)]
    a = r['ana']
    P += [f"🔥 {c(a[1])}{a[2]}（{a[3]}）", "".join(r['ana_detail']), ""]
    P.append("8. 全頭短評と独自の20ファクター指数（170点満点）")
    for x in SCORE[r['key']]:
        w = x['why']
        P.append(f"{mk.get(x['uma'], '　')} {c(x['uma'])} {x['name']} {x['final170']}点　前走 {w['sc12']}／{w['sc18']}／最高実績 {w['sc19']}")
    P += ["", "9. 独自指数ランキング TOP5", " ／ ".join(f"{i}位 {c(x['uma'])}{x['name']} {x['final170']}点" for i, x in enumerate(SCORE[r['key']][:5], 1)), ""]
    P.append("10. 最終予想印まとめ")
    for m, num, nm, who, cms in r['marks']:
        P.append(f"{m} {c(num)}{nm}　{sh[num]['final170']}点　{cms[0]}")
    P += [f"🔥 {c(a[1])}{a[2]}　{sh[a[1]]['final170']}点　{r['ana_short']}", ""]
    P += ["11. 買い目設計", r['type_line']] + kaime_lines(r)
    if k['bets']:
        P += [f"当たった時の回収イメージ（{'前日オッズ' if live else '推定配当'}）"] + [f"・{x}" for x in scen_lines(r)]
    P += [f"・{x}" for x in k['why']]
    P += ["", odds_note(r).replace('🐴', ''), "", "注意点"] + [f"・{x}" for x in r['caution']] + [""]
    d = r['danger']
    P += ["12. まとめ", f"・本命は{c(hon[1])}{hon[2]}。{hon[4][1]}",
          f"・人気の{c(d[1])}{d[0]}は軸にしません。{r['danger_short']}",
          f"・穴は{c(a[1])}{a[2]}。{r['ana_short']}",
          ("・買い目は◎から指数2位・3位の中穴へ。1番人気の本命でも配当が残る形にしました。当たり外れも含めて毎週正直に報告します。" if SNS_ONLY else "・◎と○がどちらも3番人気以内のレースだけ買う、というルールを先週から検証しています。当たり外れも含めて毎週正直に報告します。"),
          "当日の朝に実オッズと馬場を見た最終版をお届けします。フォローしてお待ちください🐴", "",
          f"#競馬予想 #AI予想 #JRA #{r['name']} #競馬"]
    return P


# ═══════════════════════════════════════════════════════════
# チェックリスト・20因子シート・内部参照
# ═══════════════════════════════════════════════════════════
CHECKLIST = [
    ["項目", "状況", "備考"],
    ["データ収集（出馬表・予想オッズ）", "✅", "9/18: 枠順確定後に土日48レースの出馬表を --force で取り直し（失敗0）。オールカマーの馬番→馬名をnetkeiba出馬表と予想オッズで突き合わせ全13頭一致。日曜のオッズは未発売のためnetkeiba予想オッズ（9/18 15:15）"],
    ["出走馬の過去走キャッシュ取り直し", "✅", "日付の無い古い形式418頭を refresh_horse_past.py で取り直し → records 再生成"],
    ["20因子v3.3採点", "✅", "score_v33_20260919.py（9/12と同じ規則）。sc[10]体重・sc[11]EV・sc[17]輸送・sc[20]外部指数は当日／未取得のため中立"],
    ["race_history.py（型判定の自前集計）", "✅", "9/16実施（過去10回）。1〜3番人気の勝利5/10＝中間型。4〜5番人気が5勝、8〜9番人気は3着内0/20"],
    ["うましる追い切り＋sc[9]", "✅", "9/17公開分を全頭取得（S=レガレイラ・エセルフリーダ／A=リビアングラス／C=コスモキュランダ・ジューンテイク・ワイドエンペラー）。2ソース目（競馬予想のホネ）は評価未掲載のため、追い切りによる印の昇格は行っていない"],
    ["追い切り評価の保存", "✅", "oikiri_store.py --add で13頭登録（累計57頭）"],
    ["インフルエンサー収集→合算", "—", "今週は未実施（合算補正0）"],
    ["買い方（ルールR1）", "✅", "◎1番人気・○6番人気（予想オッズ）→ 実収支は見送り。" + ("SNSには提示用の買い目（馬連BOX3＋◎ワイド流し3・8,000円）を掲載。台帳には登録しない" if SNS_ONLY else "当日朝の実オッズで判定し直す")],
    ["必須チェック（100円単位・合計＝予算・当たれば100%以上・Σ(1/配当)≦1）", "✅", ("SNS提示用6点で全項目OK（Σ(1/推定配当)=0.66・推定4倍未満のワイドなし）" if SNS_ONLY else "見送りのため買い目なし")],
    ["9/16実測ガード（3連複なし／🔥穴を軸にしない／予算は自信度連動）", "✅", "3連複なし・🔥③は◎からのワイドの相手のみ・予算は自信度8＝8,000円"],
    ["X/Threads 各6投稿・冒頭ヘッダー・500字チェック", "✅", "本docxに掲載"],
    ["Note記事（冒頭フック＋12セクション・マークダウン記号なし）", "✅", "本docxに掲載"],
    ["予想印インフォグラフィック・noteヘッダー画像", "⏳", "未作成（ユーザー確認待ち）"],
    ["当日朝: 実オッズ・出走取消チェック・R1判定", "⏳", "当日朝【必須】 scratch_check.py --date 20260920 --marks オールカマー=7,12,6,9,13,3"],
]


def sheet_rows(r):
    mk = {m[1]: m[0] for m in r['marks']}; mk[r['ana'][1]] = '🔥'
    mk[r['danger'][1]] = mk.get(r['danger'][1], '') or '消'
    rows = [["印", "枠-番", "馬名", "人気（最新）", "sc[7]枠", "sc[9]調教", "sc[12]前走", "sc[18]コース", "sc[19]格", "sc[20]", "素点", "合算補正", "最終(170)"]]
    live = yoso(r)
    for x in SCORE[r['key']]:
        pop = live.get(x['uma'], {}).get('pop') or x['pop']
        if not (x['uma'] in mk or (pop or 99) <= 6):
            continue
        s = x['sc']
        rows.append([mk.get(x['uma'], ''), f"{x['waku']}-{x['uma']}", x['name'], pop, s['7'], f"{s['9']}({x['oikiri']})",
                     s['12'], f"{s['18']} {x['why']['sc18']}", f"{s['19']} {x['why']['sc19']}", s['20'], x['raw170'], f"{x['ens_adj']:+.1f}", x['final170']])
    return rows


def internal(doc):
    h1(doc, '■ 内部参照（投稿しない）', (150, 30, 30))
    h2(doc, '20因子の全頭順位（score_v33_20260919.py）')
    rows = [["順", "枠-番", "馬名", "脚質", "追切", "最終(170)", "総合指数", "予想人気"]]
    for i, x in enumerate(SCORE['オールカマー'], 1):
        rows.append([i, f"{x['waku']}-{x['uma']}", x['name'], x['style'], x['oikiri'], x['final170'], x['sogo'], f"{x['pop']}人({x['odds']})"])
    table(doc, rows)
    body(doc, "\n".join([
        "・印は20因子の順位どおり（1位◎・2位○・3位▲・4〜5位△）。穴は8番人気以内から追い切りA評価の③リビアングラス",
        "・枠順前リサーチの騎手と変更: ②ワイドエンペラー＝ルメール、⑤アスクドゥポルテ＝岩田康、⑩コスモブッドレア＝津村",
        "・ルールR1の成績（レース前の印385R中128R）: 回収率120%［95%区間86〜160%］上位1除外112%。下限が100%を割っており未証明",
    ]), 9.5)


# ═══════════════════════════════════════════════════════════
def main():
    # --race チャレンジC: 1レースだけの投稿案を出す（9/11夜 ユーザー指示「明日のチャレンジCを最新のもので」）
    global OUT
    if '--race' in sys.argv:
        key = sys.argv[sys.argv.index('--race') + 1]
        RACES[:] = [r for r in RACES if r['key'] == key]
        OUT = BASE / f"{RACES[0]['date'][:10].replace('/', '')}_{key}_SNS投稿案_{_arg('--suffix', '枠順確定版')}.docx"
    elif '--races' in sys.argv:   # 複数レースを1つのdocxにまとめる（日曜2重賞）
        i = sys.argv.index('--races') + 1
        keys = []
        while i < len(sys.argv) and not sys.argv[i].startswith('--'):
            keys.append(sys.argv[i]); i += 1
        RACES[:] = [r for r in RACES if r['key'] in keys]
        d = RACES[0]['date'][:10].replace('/', '')
        OUT = BASE.parent / d / f"{d}_日曜重賞2本_SNS投稿案.docx"
        OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.8); s.bottom_margin = Cm(1.8)
        s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)
    single = len(RACES) == 1
    t = doc.add_heading(f"{RACES[0]['date']} {RACES[0]['name']}({RACES[0]['grade']}) SNS投稿案（枠順確定版）" if single
                        else '2026/9/12-13 重賞3本 SNS投稿案（枠順確定版）', level=0)
    for x in t.runs:
        set_font(x, 16)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    total = sum(KAIME[r['key']]['budget'] for r in RACES)
    sub = doc.add_paragraph((f"X6投稿＋Threads6投稿＋Note記事　投資{total:,}円　2026-09-18作成" if single else
                             'チャレンジC(G3) / セントライト記念(G2) / ローズS(G2)　各レース X6投稿＋Threads6投稿＋Note記事（合同）'
                             f'　総投資{total:,}円　2026-09-11作成'))
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for x in sub.runs:
        set_font(x, 9.5)
    note(doc, ('【買い目】枠順確定後の予想オッズ版です。')
              + '当日朝に実オッズで再計算します（現在は集客フェーズのため実際には購入しない）。')
    note(doc, '【買い方（9/11 ルール検証版）】「◎と○がどちらも1〜3番人気のレースだけ、◎○▲の馬連BOXを買う。それ以外は見送る」（'
              + '／'.join(f"{r['name']}={KAIME[r['key']]['label']}" for r in RACES)
              + '）。金額は推定配当に反比例させ、どの点が当たってもほぼ同じ額が戻るようにしています。'
              '区分ごとに買い方を変える方式は、時系列の検証で固定の買い方に負けたため取りやめました。')
    if SNS_ONLY:
        note(doc, '【SNS提示用の買い目（9/18ユーザー指示）】投稿には買い目を載せるが、実際の収支管理ではルールR1どおり見送り。'
                  'rule_forward_ledger には登録しない。構成は馬連BOX3（◎○▲）＋◎ワイド流し3（⑫⑥③）、推定4倍未満のワイドは除外、3連複なし、🔥は相手のみ。', (180, 60, 0))
    note(doc, '【馬場】今の馬場状態は当日朝に判断します。前日までは過去の馬場状態別の成績だけを使っています。')
    note(doc, '【表記ルール】外部の予想家・サイト名は投稿文に一切記載していません（内訳は末尾の内部参照のみ）。')
    doc.add_paragraph()

    h2(doc, '★完全網羅チェックリスト実施状況（枠順確定版・9/18）')
    table(doc, CHECKLIST)
    h2(doc, '最終予想印 一覧')
    rows = [["レース", "買い方", "◎", "○", "▲", "△", "🔥穴", "予算"]]
    for r in RACES:
        m = r['marks']
        rows.append([r['name'], KAIME[r['key']]['label'], f"{c(m[0][1])}{m[0][2]}", f"{c(m[1][1])}{m[1][2]}", f"{c(m[2][1])}{m[2][2]}",
                     " ".join(f"{c(x[1])}{x[2]}" for x in m[3:]), f"{c(r['ana'][1])}{r['ana'][2]}", f"{KAIME[r['key']]['budget']:,}円"])
    table(doc, rows)
    doc.add_page_break()

    # 「◯番人気」の表記は最新のオッズ（前日オッズが出たレースは実オッズ）から付け直す
    for r in RACES:
        d = r['danger']; r['danger'] = (d[0], d[1], odds_txt(r, d[1]), d[3])
        o = r['x1_out']; r['x1_out'] = (o[0], o[1], odds_txt(r, o[1]), o[3])
        if r['also_out']:
            a = r['also_out']; r['also_out'] = (a[0], a[1], odds_txt(r, a[1]), a[3])

    n = 0; chars = []; ng_all = []
    for r in RACES:
        k = KAIME[r['key']]
        ng_all += [f"{r['key']}: {x}" for x in k['checks_ng']]
        h1(doc, f"■ {r['name']}（{r['grade']}）　{r['date']}　{r['place']} {r['course']}　{r['start']}")
        note(doc, f"　{r['field']} ／ {r['sub']} ／ {r['type_line']} ／ 予算{k['budget']:,}円")
        h2(doc, '最終予想印')
        sh = sheet(r)
        rows = [["印", "馬番", "馬名", "騎手・斤量", "20因子(170)", "予想人気", "決め手"]]
        for mk, num, nm, who, cms in r['marks']:
            rows.append([mk, c(num), nm, who, sh[num]['final170'], odds_txt(r, num), cms[0]])
        a = r['ana']
        rows.append([a[0], c(a[1]), a[2], a[3], sh[a[1]]['final170'], odds_txt(r, a[1]), r['ana_short']])
        table(doc, rows)
        h2(doc, 'コース特性と過去データ')
        body(doc, "\n".join(r['outlook']))
        body(doc, "\n".join(f"・{x}" for x in r['hist']))
        h2(doc, '20因子チェックシート（印馬＋予想6番人気以内）')
        table(doc, sheet_rows(r))
        note(doc, 'sc[7]枠・sc[3]脚質はこのレース自体（同じ競馬場の年だけ）の3着内率、sc[9]は最終追い切り評価、sc[12]前走・sc[19]格は全戦績、'
                  'sc[18]は同コース同距離の最高着順と同一レースのリピーター（初コースは減点しない）。合算補正はYouTube予想の重み付きシグナル。')
        h2(doc, '買い目（出力前チェック: ' + ('✅ 全項目OK' if not k['checks_ng'] else '❌ ' + ' / '.join(k['checks_ng'])) + '）')
        lo, hi = k['ret_range']
        box(doc, [r['type_line'], ""] + kaime_lines(r) + ["", f"当たった時の回収率: 1点的中で約{lo*100:.0f}〜{hi*100:.0f}%",
                  f"Σ(1/推定配当)={k['overround']:.2f}（1.00以下＝全点ガミなし）／◎を含む点{k['hon_share']*100:.0f}%／穴絡み{k['ana_share']*100:.0f}%",
                  "", f"【選んだ根拠】{k['stats']}"]
            + ["候補: " + " ／ ".join(f"{c['name']}（安定度{c['avg']*100:.0f}%・的中{c['hit']*100:.0f}%）" for c in k['candidates'])]
            + [f"⚠ {x}" for x in k.get('guard_clash', [])] + [f"※{x}" for x in k.get('notes', [])])
        h2(doc, '的中シナリオ（推定配当）')
        body(doc, "\n".join(f"{i}. {x}" for i, x in enumerate(scen_lines(r), 1)))
        doc.add_paragraph()
        h2(doc, '── X（Twitter）投稿 6本 ──', (20, 100, 50))
        for title, lines in build_x(r):
            chars.append(post(doc, title, lines)); n += 1
        h2(doc, '── Threads 投稿 6本（各500字以内）──', (150, 60, 120))
        for title, lines in build_threads(r):
            post(doc, title, lines, limit=THREADS_LIMIT); n += 1
        doc.add_page_break()

    h1(doc, '■ Note記事（全文無料）', (120, 70, 0))
    note(doc, 'noteに貼り付ける本文。見出しに記号は使っていません（ハッシュタグのみ末尾）。')
    body(doc, "\n".join(build_note()), 10)
    doc.add_page_break()
    internal(doc)

    doc.save(str(OUT))
    amts = [b['amt'] for r in RACES for b in KAIME[r['key']]['bets']]
    print(f'[保存] {OUT}')
    print(f'　投稿数: {n} / 総投資 {total:,}円 / 全金額100円単位: {all(a % 100 == 0 for a in amts)}')
    print(f'　X投稿の平均字数: {sum(chars)//len(chars)}字（最長{max(chars)}字）')
    print('　買い目チェック:', '✅ 全レースOK' if not ng_all else '❌ ' + ' / '.join(ng_all))


if __name__ == '__main__':
    main()


