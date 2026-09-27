# -*- coding: utf-8 -*-
"""
gen_note_20260920_combined.py — 9/20（日）note記事を1本にまとめる（オールカマー＋日曜の注目レース）
====================================================================================================
2026-09-19夜 ユーザー指示「ノート記事はオールカマーと日曜の注目レースを合わせたものを出したい。1つにまとめて」。
2つの投稿案docx（どちらも生成物）から本文を読み、1本のnote記事（12セクション・マークダウン記号なし）に組み直す。
朝に実オッズで2つのdocxを作り直したら、このスクリプトを再実行すれば記事も追従する。

入力: 20260920/20260920_オールカマー_SNS投稿案_最終版.docx ／ 20260920/週末自信度7以上_SNS投稿案.docx
出力: 20260920/20260920_note記事_オールカマー＋日曜注目レース.docx（と同名 .txt）
"""
from __future__ import annotations
import re, sys
from pathlib import Path

sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn

BASE = Path.home() / 'Desktop' / '競馬予想レポート' / '20260920'
AC = BASE / '20260920_オールカマー_SNS投稿案_最終版.docx'
FLAT = BASE / '週末自信度7以上_SNS投稿案.docx'
OUT = BASE / '20260920_note記事_オールカマー＋日曜注目レース.docx'
FONT = '游ゴシック'

# 馬柱の近走レース名が10文字で切れているもの
NAME_FIX = {'サンシャイ(3勝)': 'サンシャインS(3勝)', 'エリザベス(GI)': 'エリザベス女王杯(GI)', 'ダイヤモン(GIII)': 'ダイヤモンドS(GIII)',
            'メトロポリ(L)': 'メトロポリタンS(L)'}


def ac_sections() -> dict[str, list[str]]:
    d = Document(str(AC))
    ps = [p.text for p in d.paragraphs]
    i = next(k for k, x in enumerate(ps) if x.startswith('■ Note記事'))
    j = next(k for k, x in enumerate(ps) if x.startswith('■ 内部参照'))
    sec, cur = {}, None
    for ln in ps[i + 2:j]:
        m = re.match(r'^(\d+)\. (.+)$', ln)
        if m:
            cur = int(m.group(1)); sec[cur] = []
            continue
        if cur:
            for a, b in NAME_FIX.items():
                ln = ln.replace(a, b)
            sec[cur].append(ln)
    return sec


def flat_blocks():
    d = Document(str(FLAT))
    t = [p.text for p in d.paragraphs] + [c.text for tb in d.tables for r in tb.rows for c in r.cells]
    note = next(x for x in t if 'AI注目レース' in x and '完全版' in x and '---' in x)
    body = note.split('\n---\n')[1]
    blocks = [b for b in body.strip().split('\n\n') if b.strip()]
    out = []
    for b in blocks:
        lines = b.split('\n')
        res = [lines[0].replace('　自信度', '　AI自信度')]
        for ln in lines[1:]:
            s = ln.strip()
            if s.startswith(('【買い目の型】', '→ 期待回収率', '→ ')):
                continue
            if s.startswith('根拠:'):
                s = (s.replace('根拠:', '根拠：').replace('pt差', '点差').replace(' / ', '／')
                     .replace('独自指数・ML能力の両系統で1位が一致', '2つの評価軸でそろって1位')
                     .replace('累積DB最高指数', '過去の最高指数').replace('総合指数', '指数').replace('の実績裏付け', '').replace('／全体スコア低調', ''))
                res.append('　' + s); continue
            if s.startswith('判定:'):
                m_ = re.search(r'（Σ1/配当=[\d.]+）。?(.*)$', s); extra = m_.group(1).strip() if m_ else ''
                res.append('　買い方：◎から人気上位の相手へ流し。どれが当たってもほぼ同じ額が戻るように配分')
                if extra:
                    extra = re.sub(r'（推定[\d.]+倍）', '', extra)
                    res.append('　※' + extra.replace('1点的中が損になる', '1点だけ当たると損をする'))
                continue
            if s.startswith('⏸見送り'):
                res.append('　⏸見送り（2歳戦は自信度9以上のレースだけ買うルールのため）'); continue
            # 印の行: 「総合指数…（独自…・ML…%）・」を外して脚質と一言だけ残す
            s = re.sub(r'総合指数[\d.]+（独自[\d.]+・ML[\d.]+%）・?', '', s)
            s = re.sub(r'モデル(\d+)位 vs 推定人気(\d+)位の妙味', r'指数\1位なのに人気は\2位', s)
            s = re.sub(r'独自指数[\d.]+×近5走平均[\d.]+着の人気盲点', '指数のわりに人気がない', s)
            res.append('　' + s)
        out.append(res)
    return out, body


def main():
    s = ac_sections()
    fb, _ = flat_blocks()
    play = [b for b in fb if any('💰' in x for x in b)]
    skip = [b for b in fb if not any('💰' in x for x in b)]
    tot = 8000 + sum(int(re.search(r'合計\d+点・([\d,]+)円', x).group(1).replace(',', '')) for b in play for x in b if '💰' in x)

    N = ['【9/20（日）AI予想】オールカマーは◎レガレイラ＋日曜の注目3レース｜昨日の反省を反映した最終版', '']
    N += ['1. はじめに', 'こんにちは、アスメシ競馬予想です🍱', '「明日の飯代」を懸けて、AIとデータで競馬に挑む予想アカウントです。',
          '9月20日（日）のメイン・オールカマー(G2)と、日曜の中山・阪神からAIの自信度が高い注目レースを、1本の記事にまとめました。', '',
          '結論を先に言います。オールカマーの本命は⑦レガレイラ。日曜の注目レースは'
          + '・'.join(b[0].split(' ')[0] for b in play) + f'の{len(play)}レースです。',
          f'買い目はオールカマーと合わせて{len(play) + 1}レース・合計{tot:,}円。どのレースも「1点だけ当たって損をする」組み合わせは入れていません。', '']
    N += ['2. 昨日の反省と、今日からの買い方',
          '昨日（9/19）は本命◎が5レース中3勝しました。それなのに回収率は18%。△の馬が2着に来たレースが4つあり、◎○▲の3頭だけで組んだ買い目が全滅したからです。',
          'そこで今日から、相手の選び方を変えます。本命◎は独自の指数1位、相手は◎を除いた単勝人気の上位から選びます。過去452レースで比べても、この組み方が一番安定していました（的中率43.8%）。',
          '買い目は◎からの流し。どれが当たってもほぼ同じ額が戻るように金額を分けています。◎が断然人気のレースでは、人気馬とのワイドが1倍台になり1点だけ当たると損をするので、その相手は馬連に替えるか外します。', '']
    N += ['3. オールカマー レース基本情報'] + s[2] + ['']
    N += ['4. オールカマー コース特性と過去データ'] + s[3] + ['']
    N += ['5. オールカマー オッズ一覧（前日の単勝オッズ）'] + s[4][:] + ['']
    N += ['6. オールカマー 最終追い切り評価（上位）'] + s[5] + ['']
    N += ['7. オールカマー 有力馬の詳細分析'] + s[7] + ['']
    N += ['8. オールカマー 全頭短評と独自の20ファクター指数（170点満点）'] + s[8] + ['']
    kaime = [x for x in s[11] if not x.startswith(('役割', '・昨日', '・そこで', '※前日', '注意点'))]
    kaime = [x for x in kaime if not x.startswith(('・前日の単勝', '・馬場の状態'))]
    N += ['9. オールカマー 最終予想印と買い目'] + s[10] + [''] + kaime + ['']
    N += [f'10. 日曜の注目レース（AI自信度7以上から{len(play)}レース）']
    for b in play:
        N += b + ['']
    for b in skip:
        N += b + ['']
    N += ['11. 買い目まとめ（全レース）', '・オールカマー　◎⑦から ワイド⑬⑥⑫＋馬連⑨　4点 8,000円']
    for b in play:
        head = b[0].split('　')[0]
        amt = next(x for x in b if '💰' in x).strip().replace('💰 買い目 ', '')
        N.append(f'・{head}　{amt}')
    N += [f'合計 {tot:,}円', '前日の単勝オッズで組んでいます。当日朝の実オッズで金額を最終調整し、出走取消があれば組み直します。', '']
    N += ['12. まとめ', '・オールカマーの本命は⑦レガレイラ。昨年の勝ち馬で、最終追い切りも最高評価でした',
          '・ただしこのレースの1番人気は過去10回で4割が3着を外しています。◎が崩れると買い目は外れます',
          '・今日から相手は「単勝人気の上位」から選びます。昨日、◎が勝ったのに相手で落としたレースが続いたためです',
          '・当たり外れも含めて、レース後に全部報告します', '', 'フォローして結果報告もお待ちください🐴', '',
          '#競馬予想 #AI予想 #JRA #オールカマー #レガレイラ #中山競馬場 #阪神競馬場 #競馬']

    N = [x for i, x in enumerate(N) if not (x == '' and i > 0 and N[i - 1] == '')]   # 空行の連続を1行に
    # マークダウン記号が混ざっていないか（note記事ルール）
    bad = [x for x in N if x.lstrip().startswith('#') and not x.startswith('#競馬予想') or '**' in x]
    if bad:
        print('[WARN] マークダウン記号:', bad[:3])
    for jargon in ('ML', 'Σ', '混合型', 'PDCA', 'エンジン', '総合指数'):
        hit = [x for x in N if jargon in x]
        if hit:
            print(f'[WARN] 内部用語「{jargon}」', hit[0][:60])

    doc = Document()
    for sct in doc.sections:
        sct.top_margin = Cm(1.8); sct.bottom_margin = Cm(1.8); sct.left_margin = Cm(2.0); sct.right_margin = Cm(2.0)
    for ln in N:
        p = doc.add_paragraph(); r = p.add_run(ln)
        r.font.name = FONT; r.font.size = Pt(10.5); r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
        p.paragraph_format.space_after = Pt(2)
    doc.save(str(OUT))
    OUT.with_suffix('.txt').write_text('\n'.join(N), encoding='utf-8')
    print('[保存]', OUT, f'／ {len("".join(N)):,}字 ／ 買い目合計 {tot:,}円')


if __name__ == '__main__':
    main()
