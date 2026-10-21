# -*- coding: utf-8 -*-
"""
gen_flat_sns.py — 平場（重賞以外）の予想とSNS投稿案を、予想カード（predict_day.py）から作る（2026-10-03）
出力: {日付}/{日付}_平場_予想とSNS投稿案.docx（全レース: 印・確率・買い目と期待値・X1本・Threads1本）＋ まとめ投稿
使い方: python gen_flat_sns.py --dir 20261004
"""
from __future__ import annotations
import argparse, json, re, sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.stdout.reconfigure(encoding="utf-8")
from kaime_format import circ
import bet_engine as BE

BASE = Path.home() / "Desktop" / "競馬予想レポート"
import os as _os
if _os.environ.get("REPLAY"):                      # 再生テストは _replay/ の中で作る
    BASE = BASE / "_replay"
SEP = "─────"
BAN = ["netkeiba", "SPAIA", "うましる", "競馬ラボ", "重賞ナビ", "PDCA", "EV", "インフルエンサー", "両関門", "Harville", "lift", "合算", "指数"]
ORDER = {"◎": 0, "○": 1, "▲": 2, "△": 3, "🔥": 4}
ap = argparse.ArgumentParser()
ap.add_argument("--dir", required=True)
a = ap.parse_args()
D = BASE / a.dir
CARDS = [c for c in json.loads((D / "research" / f"cards_{a.dir}.json").read_text(encoding="utf-8"))["cards"] if not c["graded"]]
DATE = f"{a.dir[:4]}年{int(a.dir[4:6])}月{int(a.dir[6:])}日"
WD = "月火水木金土日"[__import__("datetime").date(int(a.dir[:4]), int(a.dir[4:6]), int(a.dir[6:])).weekday()]
TAGS = "#競馬予想 #AI予想 #JRA"
ODDS_NOTE = "前日の予想オッズで計算しています。当日の発走前に実際のオッズで組み直します。"


def title(c):
    n = re.sub(r"\s*出馬表.*$", "", c["name"]).strip()
    return n if n else c["grade"]


def head(c):
    return f"🏇【{c['venue']}{c['R']}R {title(c)}】{DATE}（{WD}）\n━━ {c['venue']} {c['dist']} / {c['n']}頭 ━━"


def nm(c, u):
    h = next(x for x in c["horses"] if x["uma"] == u)
    return f"{circ(u)}{h['name']}"


def marked(c):
    return sorted(((mk, int(u)) for u, mk in c["marks"].items()), key=lambda x: (ORDER[x[0]], x[1]))


def bets(c):
    """公開の買い目: 標準形V1 ＋（あれば）妙味の馬連。妙味の馬連が出たレースは「勝負レース」"""
    src = c["ref_bets"] or c["bets"]
    if not src:
        return None, "見送り（期待値に見合う組み合わせがありません）"
    std = [b for b in src if not b.get("value")]
    val = [b for b in src if b.get("value")]

    def ln(b):
        return f"【{b['type']}】{'-'.join(circ(x) for x in b['combo'])} {b['stake']:,}円"
    L = [f"💰 買い目 合計{len(src)}点・{sum(b['stake'] for b in src):,}円"] + [ln(b) for b in std]
    if val:
        L += ["", "🎯 妙味の馬連（人気より割安な組み合わせ）"] + [ln(b) for b in val]
    return ("勝負" if val else "標準"), L


def one_race(c):
    ms = marked(c)
    top = next(h for h in c["horses"] if h["uma"] == ms[0][1])
    kind, bl = bets(c)
    why = [r for r in top["reasons"][:2]]
    if top.get("chokyo") and not any(w.startswith("調教") for w in why):
        why.append(f"調教: {top['chokyo']}")
    trend = [t for t in c["trends"] if t.startswith("前日")][:1]
    X = [head(c), "", SEP, "🎯 予想印", SEP] + [f"{mk} {nm(c, u)}" for mk, u in ms] + ["", f"◎の見立て: 勝つ確率{top['p_model'] * 100:.0f}%（人気が示す確率{top['p_market'] * 100:.0f}%）"] + [f"・{x}" for x in why] + [f"・{t}" for t in trend] + [""]
    X += (bl if isinstance(bl, list) else [bl])
    X += ["", f"※{ODDS_NOTE}", TAGS]
    return kind, "\n".join(X)


def fit(text, limit=500):
    lines = text.split("\n")
    while len("\n".join(lines)) > limit:
        # 根拠の「・」行を後ろから落とす
        idx = [i for i, l in enumerate(lines) if l.startswith("・")]
        if not idx:
            break
        lines.pop(idx[-1])
    return "\n".join(lines)


PROBLEMS = []


def check(label, text, limit=None):
    hits = [b for b in BAN if b in text]
    if hits:
        PROBLEMS.append(f"{label}: 禁止語 {hits}")
    if limit and len(text) > limit:
        PROBLEMS.append(f"{label}: {len(text)}字")
    for ln in text.split("\n"):
        for amt in re.findall(r"([\d,]+)円", ln):
            v = int(amt.replace(",", ""))
            if 0 < v < 100000 and v % 100:
                PROBLEMS.append(f"{label}: 100円単位でない {amt}円")


from docx import Document
from docx.shared import Pt

doc = Document()
doc.styles["Normal"].font.size = Pt(10)


def box(t, text, limit=None):
    doc.add_heading(f"{t}　{len(text)}字" + (f"/{limit}字" if limit else ""), 3)
    tb = doc.add_table(rows=1, cols=1); tb.style = "Table Grid"; tb.cell(0, 0).text = text
    doc.add_paragraph(); check(t, text, limit)


results = [(c, *one_race(c)) for c in CARDS]
e2 = [(c, k) for c, k, _ in results if k == "勝負"]
summ = [f"🏇{DATE}（{WD}）平場 予想まとめ", "", "印は人気順ではなく「人気に対する割安さ」の順。🔥は人気薄の穴、🎯は人気より割安な馬連が出た勝負レースです。", "",
        SEP, "🎯 勝負レース（人気より割安な馬連が出たレース）", SEP]
for c, _ in e2:
    ms = marked(c)
    summ.append(f"・{c['venue']}{c['R']}R {title(c)}　◎{nm(c, ms[0][1])}")
summ += ["", f"それ以外の{len(results) - len(e2)}レースは、印と標準の買い目だけを載せています。", "各レースの印と根拠は、続くスレッドで👇", f"※{ODDS_NOTE}", TAGS]
doc.add_heading(f"平場 予想とSNS投稿案 — {DATE}（{WD}）", 0)
doc.add_paragraph("印・確率は、過去10年のデータ＋コースの傾向＋前日の馬場の特徴を使った新エンジンの値です。")
box("まとめ投稿（X・先頭に置く）", "\n".join(summ))
doc.add_heading("全レース一覧", 1)
tb = doc.add_table(rows=0, cols=4); tb.style = "Table Grid"
for row in [["レース", "印", "買い目", "区分"]] + [[f"{c['venue']}{c['R']}R {title(c)} {c['dist']}", " ".join(f"{mk}{circ(u)}" for mk, u in marked(c)),
                                              (f"{len(c['bets'] or c['ref_bets'])}点" if (c['bets'] or c['ref_bets']) else "見送り"), k or "見送り"] for c, k, _ in results]:
    cs = tb.add_row().cells
    for j, v in enumerate(row):
        cs[j].text = str(v)
for venue in ("京都", "東京"):
    doc.add_heading(f"■ {venue}", 1)
    for c, k, x in results:
        if c["venue"] != venue:
            continue
        doc.add_heading(f"{c['venue']}{c['R']}R {title(c)}（{c['grade']}）　{c['dist']} {c['n']}頭", 2)
        doc.add_paragraph("傾向: " + " ／ ".join(c["trends"][:3]))
        tbl = doc.add_table(rows=0, cols=6); tbl.style = "Table Grid"
        for row in [["印", "馬", "馬名", "性齢・斤量・騎手", "勝つ確率", "人気が示す確率"]] + [[mk, f"{h['waku']}枠{circ(h['uma'])}", h["name"], f"{h['sexage']} {h['kin']:g}kg {h['jockey']}", f"{h['p_model'] * 100:.1f}%", f"{h['p_market'] * 100:.1f}%（{h['pop']}人気）"]
                                                   for mk, u in marked(c) for h in c["horses"] if h["uma"] == u]:
            cs = tbl.add_row().cells
            for j, v in enumerate(row):
                cs[j].text = str(v)
        box(f"{c['venue']}{c['R']}R X投稿", x)
        box(f"{c['venue']}{c['R']}R Threads投稿", fit(x), 500)
out = D / f"{a.dir}_平場_予想とSNS投稿案.docx"
try:
    doc.save(out)
except PermissionError:
    out = D / f"{a.dir}_平場_予想とSNS投稿案_new.docx"; doc.save(out)
print(f"平場 {len(results)}レース（勝負レース {len(e2)}）")
print("検査: " + ("問題なし" if not PROBLEMS else f"{len(PROBLEMS)}件"))
for p in PROBLEMS[:10]:
    print("  ⚠", p)
print("保存:", out)
