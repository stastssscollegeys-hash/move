# -*- coding: utf-8 -*-
"""
gen_value_horses.py — 週末「期待値レース」＆「狙える馬」抽出 SNS投稿案（前日版）

入力: collect_weekend_bigdata.py が出力した records.json（総合指数・独自指数・ML能力%・累積DB照合）
処理:
  1. 期待値レース: 自信度7以上の全レースを v5 EVアダプティブ（53パターン機械評価）にかけ、
     E[回収率]・P(300%+)・ガミ率・的中率を算出 → E[回収率]順に一覧化（見送りも明示）
  2. 狙える馬: 全レース・全頭で モデル勝率p（総合指数順×実測テーブル）と
     大衆勝率q（近走着順・ML・騎手の大衆スコア＝推定人気）を計算し、
     p/q乖離（＝オッズ歪み＝期待値の源泉）が大きく、かつ独自指数・累積DBの裏付けがある馬を抽出
     ＋「人気盲点」（独自指数68+×近5走平均着4着以下×DB最高指数70+）
  3. 重賞は「週末重賞2本_確定版」の印が正のため、期待値レース一覧からは除外し、
     狙える馬は「重賞の妙味馬」として別枠で参考掲載

使い方:
  python gen_value_horses.py --date 20260822 [--odds-file odds_20260822.json]
出力:
  Desktop/競馬予想レポート/{date}/週末期待値TOP＆狙える馬_SNS投稿案.docx
"""
import sys, json, argparse
sys.stdout.reconfigure(encoding='utf-8')
from pathlib import Path
from collections import defaultdict
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from gen_weekend_confidence_report import calc_confidence
from gen_kaime_v5 import model_probs, public_probs, BUDGET_BY_CONF, EV_MIN, is_young_race
from v55_guard import choose_pattern_v55, apply_young_cap, WEEKLY_CONTROL_ON, WEEKLY_CONTROL_REASON
from kaime_format import fmt_bets, circ
from flat_race_rules import assign_marks, build_kaime, scenarios
FLAT = {}
ODDS_ALL = {}

FONT = '游ゴシック'
CONF_MIN = 7
TOP_RACES = 5
TOP_HORSES = 10
GRADE_NAMES = ('キーンランド', '新潟２歳', '新潟2歳')   # 重賞（確定版の印が正）
WEEKDAY = {'20260822': '土', '20260823': '日'}


def set_font(run, size=10, bold=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

def h1(doc, text, color=(15, 71, 97)):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        set_font(r, 14); r.font.color.rgb = RGBColor(*color)

def h3(doc, text):
    p = doc.add_heading(text, level=3)
    for r in p.runs:
        set_font(r, 11)

def box(doc, lines, size=10):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("\n".join(lines))
    set_font(r, size)

def table(doc, rows, size=8.5):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = 'Table Grid'
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ''
            run = cell.paragraphs[0].add_run(str(val))
            set_font(run, size, bold=(i == 0))
    doc.add_paragraph()


def is_grade(title):
    return any(n in (title or '') for n in GRADE_NAMES)


def load(json_path):
    with open(json_path, encoding='utf-8') as f:
        data = json.load(f)
    races = defaultdict(list)
    for r in data['records']:
        races[(r['date'], r['競馬場'], r['R'])].append(r)
    for k in races:
        races[k].sort(key=lambda x: x['AI予測順位'])
    conf = {k: calc_confidence(v) for k, v in races.items()}
    for k, v in races.items():   # 平場にも重賞ルール（印可変・4類型多層買い目）を適用
        odds = ODDS_ALL.get(f"{k[1]}{k[2]}") if ODDS_ALL else None
        marks, roles, info = assign_marks(v, q=public_probs(v, odds) if odds else None)
        for r in v:
            m = marks.get(r['馬番'], '')
            r['AI印'] = '' if m == '❌' else m
        FLAT[k] = (marks, roles, info, build_kaime(v, conf[k][0], marks, roles, info, odds_override=odds) if conf[k][0] >= CONF_MIN else None)
    return races, conf


def eval_races(races, conf, odds_all):
    """自信度7以上の各レースをv5で評価して一覧化。"""
    rows = []
    keys = [k for k in races if conf[k][0] >= CONF_MIN]
    keep, young_dropped = apply_young_cap(keys, races, conf)
    for k in keys:
        recs = races[k]
        c, tags = conf[k]
        top = recs[0]
        odds = odds_all.get(f"{k[1]}{k[2]}") if odds_all else None
        name, desc, res, ranking, p, q, info = choose_pattern_v55(recs, BUDGET_BY_CONF[c], odds)
        kb = FLAT[k][3]
        if kb is not None:   # 類型別多層テンプレ（v5評価器で採点）を採用。engine列は参考
            info = {'engine': info.get('engine'), 'changed': (kb['arch'] != 'V5'), 'reason': kb['why'], 'kb': kb}
            name, res = (None, None) if kb['skip'] else (kb['arch_name'], kb['res'])
        if k not in keep:
            info = {'engine': info.get('engine'), 'changed': True, 'reason': young_dropped[k]}
            name, desc, res = None, None, None
        best = ranking[0][2] if ranking else None
        rows.append({
            'info': info,
            'key': k, 'title': top['レース名'], 'grade': top.get('グレード', ''), 'dist': top['距離'],
            'n': top['頭数'], 'conf': c, 'tags': tags, 'is_grade': is_grade(top['レース名']),
            'pattern': name, 'desc': desc, 'res': res, 'best': best,
            'E': (res['E_rate'] if res else (best['E_rate'] if best else 0)),
            'P300': (res['P_target'] if res else (best['P_target'] if best else 0)),
            'hit': (res.get('P_hit', 0) if res else (best.get('P_hit', 0) if best else 0)),
            'gami': (res['gami_ratio'] if res else (best['gami_ratio'] if best else 0)),
            'recs': recs,
        })
    rows.sort(key=lambda x: (-(x['pattern'] is not None), -x['E'], -x['conf']))
    return rows


def value_horses(races, odds_all):
    """全レース全頭の p/q 乖離から狙える馬を抽出。"""
    cands = []
    for k, recs in races.items():
        if is_young_race(recs[0].get('レース名', '')):
            continue   # 新馬・未勝利は事前情報が薄く「妙味の幻影」が出やすい（v5.2教訓）→ 狙える馬の対象外
        p = model_probs(recs)
        odds = odds_all.get(f"{k[1]}{k[2]}") if odds_all else None
        q = public_probs(recs, odds)
        for i, r in enumerate(recs):
            ratio = p[i] / max(q[i], 1e-6)
            if r['総合指数'] < 55 or r['ML能力%'] < 8.0:
                continue   # 総合・ML水準が低い馬の乖離は「人気がないだけ」→ 除外
            blind = (r['独自指数'] >= 68 and (r.get('近5走平均着') or 0) >= 4.0 and (r.get('DB最高指数') or 0) >= 70)
            if r['AI予測順位'] <= 3 and ratio >= 1.25 and r['独自指数'] >= 62:
                kind = '妙味(乖離)'
            elif blind and r['AI予測順位'] <= 5 and ratio >= 1.0:   # 推定人気で既に上位なら盲点ではない
                kind = '人気盲点'
            else:
                continue
            score = ratio * (r['独自指数'] / 70.0) * (1.0 + 0.1 * (3 - min(r['AI予測順位'], 3)))
            if blind:
                score *= 1.1
            cands.append({
                'key': k, 'title': recs[0]['レース名'], 'dist': recs[0]['距離'], 'n': recs[0]['頭数'],
                'is_grade': is_grade(recs[0]['レース名']), 'r': r, 'p': p[i], 'q': q[i],
                'ratio': ratio, 'score': score, 'kind': kind, 'blind': blind,
                'crowd_rank': sorted(range(len(recs)), key=lambda j: -q[j]).index(i) + 1,
            })
    cands.sort(key=lambda x: -x['score'])
    return cands


def horse_reason(c):
    r = c['r']
    bits = [f"総合指数{r['総合指数']}（独自{r['独自指数']}・ML{r['ML能力%']}%）",
            f"モデル{r['AI予測順位']}位 vs 推定人気{c['crowd_rank']}位（乖離×{c['ratio']:.2f}）"]
    if r.get('DB最高指数'):
        bits.append(f"累積DB最高指数{r['DB最高指数']}")
    if (r.get('近5走平均着') or 0) >= 4.0:
        bits.append(f"近5走平均{r['近5走平均着']}着＝人気落ちの盲点")
    if r.get('脚質') and r['脚質'] != '?':
        bits.append(f"脚質{r['脚質']}")
    if (r.get('騎手勝率%') or 0) >= 12:
        bits.append(f"騎手勝率{r['騎手勝率%']}%")
    return " / ".join(bits)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--date', required=True)
    ap.add_argument('--odds-file', help='実オッズJSON {"札幌11": {"3": 1.5, ...}}（当日再実行用）')
    args = ap.parse_args()
    day = args.date
    wd = WEEKDAY.get(day, '')
    md = f"{int(day[4:6])}/{int(day[6:])}"
    base = Path.home() / 'Desktop' / '競馬予想レポート' / day
    json_path = base / f"週末ビッグデータ_{day}-{day[4:]}_records.json"
    out = base / "週末期待値TOP＆狙える馬_SNS投稿案.docx"
    odds_all = json.load(open(args.odds_file, encoding='utf-8')) if args.odds_file else {}
    ODDS_ALL.update(odds_all)
    odds_label = "実オッズ反映版" if odds_all else "前日推定オッズ版（当日朝に実オッズで再実行）"

    races, conf = load(json_path)
    n_races = len(races); n_horses = sum(len(v) for v in races.values())
    ev_rows = eval_races(races, conf, odds_all)
    ev_flat = [x for x in ev_rows if not x['is_grade']]
    ev_grade = [x for x in ev_rows if x['is_grade']]
    cands = value_horses(races, odds_all)
    top_flat = [c for c in cands if not c['is_grade']][:TOP_HORSES]
    top_grade = [c for c in cands if c['is_grade']]

    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.8); s.bottom_margin = Cm(1.8)
        s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)
    t = doc.add_heading(f"週末 期待値レースTOP＆狙える馬 SNS投稿案（{md}{wd}・前日版）", level=0)
    for r in t.runs: set_font(r, 16)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(
        f"全{n_races}レース{n_horses}頭 / 17ファクター独自指数×ML v44×累積DB12,120レコード / "
        f"v5 EVアダプティブ53パターン機械評価 / {odds_label}")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in sub.runs: set_font(r, 9.5)
    note = doc.add_paragraph(
        "※重賞（キーンランドC・新潟2歳S）は「20260823_週末重賞2本_確定版_SNS投稿案.docx」の印・買い目が正。"
        "本docxの重賞欄は機械抽出の参考値。※推定人気は近走着順・ML・騎手から機械推定したもので、"
        "実オッズとズレる馬がいる（妙味の幻影を避けるため当日朝に --odds-file で再実行すること）")
    for r in note.runs: set_font(r, 9, color=(180, 0, 0))

    # ── 1. 期待値レース一覧 ──
    h1(doc, f"■ 期待値レース一覧（自信度7以上 {len(ev_rows)}レース・E[回収率]順）", color=(20, 100, 50))
    rows = [["会場R", "レース", "頭数", "自信度", "採用（4類型多層・v5.5適用後）", "v5エンジン素の選択（参考）", "E[回収率]", "P(300%+)", "的中率", "ガミ率", "予算"]]
    for x in ev_rows:
        k = x['key']
        info = x['info']
        if x['pattern']:
            pat = x['pattern']
        elif info.get('reason', '').startswith('若馬戦上限'):
            pat = f"見送り（{info['reason']}）"
        else:
            pat = f"見送り（{info['reason'] or f'最良でも{x['E']*100:.0f}%<{EV_MIN*100:.0f}%'}）"
        rows.append([f"{k[1]}{k[2]}R", (x['title'] + ("【重賞・別記事】" if x['is_grade'] else "")), x['n'], f"{x['conf']}/10",
                     pat, (info.get('engine') or '見送り') + ("→変更" if info.get('changed') else ""),
                     f"{x['E']*100:.0f}%", f"{x['P300']*100:.0f}%", f"{x['hit']*100:.0f}%",
                     f"{x['gami']*100:.0f}%", f"{x['res']['total']:,}円" if x['res'] else "0円"])
    table(doc, rows)
    box(doc, [f"【v5.5実戦ガード】週次制御{'ON' if WEEKLY_CONTROL_ON else 'OFF'}: {WEEKLY_CONTROL_REASON}。",
              "若馬戦（新馬・未勝利）は1日2レース・自信度9以上のみ参加。「エンジン素の選択」はgen_kaime_v5の生出力（週末買い目_v5EVアダプティブ.docxと同じ）で、",
              "SNS公開用の買い目はv5.5適用後の列が正。当日朝は --odds-file で実オッズ再実行が必須（V系・単勝ゲートは実オッズでのみ発動）"], size=9)

    # ── 2. 狙える馬一覧 ──
    h1(doc, f"■ 狙える馬（全レース機械抽出・上位{len(top_flat)}頭＋重賞{len(top_grade)}頭）", color=(180, 80, 0))
    rows = [["順", "会場R", "レース", "馬番", "馬名", "印", "種別", "総合", "独自", "ML%", "モデル順", "推定人気", "乖離", "DB最高", "近5走平均着"]]
    for i, c in enumerate(top_flat + top_grade, 1):
        r = c['r']; k = c['key']
        rows.append([i if not c['is_grade'] else "重賞", f"{k[1]}{k[2]}R", c['title'][:10], r['馬番'], r['馬名'], r['AI印'], c['kind'],
                     r['総合指数'], r['独自指数'], r['ML能力%'], r['AI予測順位'], c['crowd_rank'], f"×{c['ratio']:.2f}",
                     r.get('DB最高指数', ''), r.get('近5走平均着', '')])
    table(doc, rows, size=8)
    box(doc, ["【抽出ロジック】",
              "・妙味(乖離): モデル上位3位以内 × モデル勝率p÷推定人気勝率q ≧1.25 × 独自指数62+（大衆評価とのズレ＝期待値の源泉）",
              "・人気盲点: 独自指数68+ × 近5走平均4着以下（近走不振で人気を落とす）× 累積DB最高指数70+（能力の裏付けあり）",
              "・スコア = 乖離 × 独自指数/70 × モデル順位ボーナス（盲点は×1.1）。重賞は別枠（確定版の印が正）"], size=9)

    # ── 3. X投稿 ──
    tense = "明日"
    H = f"🏇【{tense}{md}（{wd}）期待値で選ぶ注目レースTOP{TOP_RACES}】\n━━ 札幌・新潟・中京 全{n_races}レースをAI機械評価 / 自信度7以上のみ ━━"
    lines = [H, "",
             "独自の17ファクター指数×機械学習×累積1万2千レコードで全レースを採点し、",
             "53種類の馬券パターンを機械評価して「期待回収率」が高い順に並べました。",
             "基準（期待回収率115%以上×ガミ率35%以下）を満たさないレースは正直に見送りです👇", ""]
    for i, x in enumerate([e for e in ev_flat if e['pattern']][:TOP_RACES], 1):
        k = x['key']; marks = [r for r in x['recs'] if r['AI印']]
        lines.append(f"{i}. {k[1]}{k[2]}R {x['title']} {x['dist']} {x['n']}頭　自信度{x['conf']}/10")
        lines.append("　" + "　".join(f"{r['AI印']}{r['馬番']}{r['馬名']}" for r in marks))
        _ph = (x['info'].get('kb') or {}).get('phantom', False)
        lines.append(("　期待回収率は参考値（市場乖離大）" if _ph else f"　期待回収率{x['E']*100:.0f}%") + f" / 300%超の確率{x['P300']*100:.0f}% / {x['pattern']}・{len(x['res']['bets'])}点{x['res']['total']:,}円")
        lines.append("")
    lines += ["買い目の全点はNote記事で公開。オッズは前日推定なので、",
              f"{tense}朝の実オッズで最終調整します🎯", "",
              "#競馬予想 #AI予想 #JRA #期待値"]
    h1(doc, "【X投稿①】期待値レースTOP（コピー用）", color=(20, 100, 50))
    box(doc, lines)

    H2 = f"🏇【{tense}{md}（{wd}）AIが見つけた「狙える馬」{len(top_flat)}頭】\n━━ 全{n_horses}頭から「実力＞人気」のズレだけを機械抽出 ━━"
    lines = [H2, "",
             "AIの実力評価が、近走成績や騎手から推定した人気より明確に高い馬＝",
             "オッズに妙味が乗りやすい馬だけを抜き出しました。",
             "印の上位でなくても、相手・穴として組み込む価値がある馬たちです👇", ""]
    for i, c in enumerate(top_flat, 1):
        r = c['r']; k = c['key']
        lines.append(f"{i}. {k[1]}{k[2]}R {c['title']}　{r['AI印'] or '無印'} {r['馬番']}番 {r['馬名']}")
        lines.append(f"　{horse_reason(c)}")
        lines.append("")
    if top_grade:
        lines.append("【重賞の妙味馬（機械抽出・詳細は重賞予想記事で）】")
        for c in top_grade:
            r = c['r']; k = c['key']
            lines.append(f"・{c['title']}　{r['馬番']}番 {r['馬名']}（モデル{r['AI予測順位']}位 vs 推定人気{c['crowd_rank']}位）")
        lines.append("")
    lines += ["※推定人気ベースの抽出です。当日の実オッズで想定より売れていれば妙味は消えるので、",
              "　朝のオッズと照らして最終判断します🍱", "",
              "#競馬予想 #AI予想 #JRA #穴馬"]
    h1(doc, "【X投稿②】狙える馬（コピー用）", color=(20, 100, 50))
    box(doc, lines)

    # ── 4. Threads（500字以内） ──
    h1(doc, "【Threads投稿】（各500字以内）", color=(80, 40, 120))
    th1 = [f"🏇【{tense}{md}（{wd}）期待値レースTOP3】", "━━ 全レースをAI機械評価 ━━", ""]
    for i, x in enumerate([e for e in ev_flat if e['pattern']][:3], 1):
        k = x['key']; marks = [r for r in x['recs'] if r['AI印'] and r['AI印'] != '△']
        th1.append(f"{i}. {k[1]}{k[2]}R {x['title']} 自信度{x['conf']}/10")
        th1.append("　" + "　".join(f"{r['AI印']}{r['馬番']}{r['馬名']}" for r in marks))
        th1.append(f"　期待回収率{x['E']*100:.0f}%")
    th1 += ["", "買い目はNoteで全公開🎯", "あなたの注目レースはどれですか？👇", "", "#競馬予想 #AI予想"]
    th2 = [f"🏇【{tense}{md}（{wd}）AIが見つけた狙える馬5頭】", "━━ 実力＞人気のズレを機械抽出 ━━", ""]
    for i, c in enumerate(top_flat[:5], 1):
        r = c['r']; k = c['key']
        th2.append(f"{i}. {k[1]}{k[2]}R {r['馬番']}番 {r['馬名']}（AI{r['AI予測順位']}位/推定人気{c['crowd_rank']}位）")
    th2 += ["", "人気薄の相手に迷ったらこの中から🍱", "気になる馬はいますか？👇", "", "#競馬予想 #穴馬"]
    for title, tl in (("Threads 1/2 期待値レース", th1), ("Threads 2/2 狙える馬", th2)):
        text = "\n".join(tl); n = len(text)
        h3(doc, f"{title}（{n}字/500字 {'OK' if n <= 500 else '⚠超過!要短縮'}）")
        if n > 500:
            print(f"[WARN] Threads文字数超過: {title} = {n}字")
        box(doc, ["【コピー用】"] + tl)

    # ── 5. Note記事 ──
    h1(doc, "【Note記事全文】期待値レース全一覧＋狙える馬（買い目付き）", color=(180, 80, 0))
    nl = [f"{tense}の期待値レース全一覧と「狙える馬」— アスメシ競馬予想", "",
          "こんにちは、アスメシ競馬予想です🍱", "",
          f"{tense}{md}（{wd}）の札幌・新潟・中京、全{n_races}レース{n_horses}頭を",
          "17ファクター独自指数×機械学習×累積1万2千レコードのビッグデータで採点し、",
          "自信度7以上のレースすべてに53種類の馬券パターンを機械評価で当てはめました。",
          "結論を先に言います。期待回収率が最も高いのは",
          (f"{ev_flat[0]['key'][1]}{ev_flat[0]['key'][2]}R {ev_flat[0]['title']}（期待回収率{ev_flat[0]['E']*100:.0f}%）です。" if ev_flat and ev_flat[0]['pattern'] else "該当なし（全レース見送り）です。"),
          "", "---", "", "期待値レース一覧（期待回収率順）", ""]
    for x in ev_flat:
        k = x['key']
        nl.append(f"{k[1]}{k[2]}R {x['title']} {x['dist']} {x['n']}頭　自信度{x['conf']}/10")
        marks = [r for r in x['recs'] if r['AI印']]
        nl.append("　" + "　".join(f"{r['AI印']}{r['馬番']}{r['馬名']}" for r in marks))
        nl.append(f"　根拠: {' / '.join(x['tags'])}")
        if x['pattern']:
            _m, _roles, _info, _kb = FLAT[x['key']]
            nl.append(f"　【買い目の型】{x['pattern']}（{_kb['why']}／目標回収率{_kb['target']}）")
            nl += [f"　{l}" for l in fmt_bets(x['res']['bets'], x['recs'])]
            nl.append(f"　→ 期待回収率{x['E']*100:.0f}% / 300%超の確率{x['P300']*100:.0f}% / ガミ率{x['gami']*100:.0f}%")
            _sc = scenarios(x['recs'], _roles, x['res'], circ)
            if _sc:
                nl.append('　📌 的中シナリオ（前日推定オッズ）'); nl += [f"　　{t}" for t in _sc]
        else:
            nl.append(f"　💰買い目: 見送り（最良パターンでも期待回収率{x['E']*100:.0f}%と基準未満）")
        nl.append("")
    nl += ["---", "", "狙える馬（実力＞人気のズレを機械抽出）", ""]
    for i, c in enumerate(top_flat, 1):
        r = c['r']; k = c['key']
        nl.append(f"{i}. {k[1]}{k[2]}R {c['title']}　{r['AI印'] or '無印'} {r['馬番']}番 {r['馬名']}（{c['kind']}）")
        nl.append(f"　{horse_reason(c)}")
        nl.append("")
    if top_grade:
        nl += ["重賞の妙味馬（機械抽出・参考）", ""]
        for c in top_grade:
            r = c['r']; k = c['key']
            nl.append(f"・{c['title']}　{r['馬番']}番 {r['馬名']}　{horse_reason(c)}")
        nl.append("")
    nl += ["---", "",
           "推定人気は近走成績・機械学習・騎手から機械的に推定したものです。",
           f"{tense}朝の実オッズで想定より売れている馬は妙味が薄れるため、最終判断は朝のオッズと照らして行います。",
           "重賞（キーンランドカップ・新潟2歳ステークス）は全頭診断＋買い目を別記事で公開しています。", "",
           "予想が参考になったらフォロー＆いいねお願いします！",
           f"みなさんの{'土曜' if wd == '土' else '日曜'}が楽しいものになりますように🍜🎉", "",
           "#競馬予想 #AI予想 #JRA #期待値 #穴馬"]
    box(doc, nl, size=9.5)

    doc.save(str(out))
    print(f"保存完了: {out}")
    print(f"期待値レース: {len(ev_rows)}（うち参加{sum(1 for x in ev_rows if x['pattern'])}・見送り{sum(1 for x in ev_rows if not x['pattern'])}・重賞{len(ev_grade)}）")
    print(f"狙える馬: 平場{len(top_flat)}頭 / 重賞{len(top_grade)}頭 / 候補総数{len(cands)}")
    for c in top_flat:
        r = c['r']; k = c['key']
        print(f"  {k[1]}{k[2]}R {r['馬番']}番{r['馬名']} {c['kind']} 乖離×{c['ratio']:.2f} 独自{r['独自指数']} AI{r['AI予測順位']}位/人気{c['crowd_rank']}位")


if __name__ == '__main__':
    main()
