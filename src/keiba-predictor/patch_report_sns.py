# -*- coding: utf-8 -*-
"""
patch_report_sns.py — predict_and_report.py が出力する「YYYYMMDD_競馬予想レポート.docx」の
SNS該当箇所（🎯買い目のご参考／📱X投稿文①②③／🧵Threads／📝Note）を、現行ルールで置き換える後処理。

置換後の内容（2026-08-21ルール）:
  - 印: ◎○▲各1・△1〜5頭可変・🔥穴・❌危険な人気馬（flat_race_rules）
  - 買い目: 4類型の多層構成（v5評価器採点・降格・見送り）＋配当シナリオ、表記v2（kaime_format）
  - 全投稿に冒頭ヘッダー、前日時制（明日）、Noteは冒頭フック定型・マークダウン記号なし
ベースラインの predict_and_report.py 本体は改変しない。元docxは scratchpad にバックアップしてから上書きする。

使い方: python patch_report_sns.py --date 20260822
"""
import sys, re, shutil, argparse, datetime
sys.stdout.reconfigure(encoding='utf-8')
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from gen_digest_sns import (load, mark_items, marks_line, horse_phrase, hdr, kaime_block, skip_reason, is_grade,
                            WEEKDAY, GRADE_LINK, BABA_NOTE, DB_RECORDS, CONF_MIN)
GRADE_DOC = "20260823_週末重賞2本_確定版_SNS投稿案.docx"
from v55_guard import apply_young_cap
from kaime_format import fmt_bets, circ
from flat_race_rules import scenarios

FONT = '游ゴシック'
VENUE_MAP = {'札幌': '札幌', '新潟': '新潟', '中京': '中京', '函館': '函館', '東京': '東京', '阪神': '阪神', '京都': '京都', '小倉': '小倉', '福島': '福島', '中山': '中山'}


def set_font(run, size=10, bold=False, color=None):
    run.font.name = FONT; run.font.size = Pt(size); run.font.bold = bold
    if color: run.font.color.rgb = RGBColor(*color)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

def heading(doc, text, level):
    p = doc.add_heading(text, level=level)
    for r in p.runs: set_font(r, {2: 13, 3: 11.5, 4: 11}.get(level, 11))
    return p

def para(doc, lines, size=10, color=None):
    p = doc.add_paragraph()
    r = p.add_run("\n".join(lines) if isinstance(lines, list) else lines); set_font(r, size, color=color)
    return p


def main():
    ap = argparse.ArgumentParser(); ap.add_argument('--date', required=True); a = ap.parse_args()
    day = a.date; wd = WEEKDAY.get(day, ''); md = f"{int(day[4:6])}/{int(day[6:])}"
    base = Path.home() / 'Desktop' / '競馬予想レポート' / day
    path = base / f"{day}_競馬予想レポート.docx"
    bk_dir = Path(r"C:/Users/User/AppData/Local/Temp/claude/C--Users-User-dev-move/101fc63b-0785-48c5-a555-424f66e4848b/scratchpad/report_backup")
    bk_dir.mkdir(parents=True, exist_ok=True)
    bk = bk_dir / f"{path.stem}_before_sns_patch_{datetime.datetime.now():%H%M%S}.docx"
    shutil.copy2(path, bk)

    doc = Document(str(path))
    body = doc.element.body
    # ── 置換開始位置（🎯買い目のご参考）と、そこに列挙された推奨レースを取得 ──
    start = None; featured = []
    for i, p in enumerate(doc.paragraphs):
        if start is None and '買い目のご参考' in p.text and p.style.name.startswith('Heading'):
            start = i
        elif start is not None and p.style.name == 'Heading 4':
            m = re.match(r'(\S+)\s+R(\d+)', p.text)
            if m: featured.append((m.group(1), int(m.group(2))))
    if start is None:
        print("置換対象の見出しが見つかりません"); return
    start_el = doc.paragraphs[start]._p
    # 開始見出し以降の本文要素（sectPr以外）を削除
    children = list(body)
    idx = children.index(start_el)
    for el in children[idx:]:
        if el.tag.endswith('sectPr'): continue
        body.remove(el)

    races, conf, flat = load(day, {})
    keys7 = [k for k in races if conf[k][0] >= CONF_MIN]
    keep, young_dropped = apply_young_cap(keys7, races, conf)

    def status(k):
        c = conf[k][0]; kb = flat[k][3]
        if c < CONF_MIN: return f"自信度{c}/10（7未満のため買い目対象外）"
        if k not in keep: return f"見送り（{young_dropped[k]}）"
        if kb['skip']: return f"見送り（{skip_reason(kb)}）"
        return None

    fkeys = [k for v, r in featured for k in races if k[1] == v and k[2] == r]
    if not fkeys:
        fkeys = sorted(keys7, key=lambda k: -conf[k][0])[:3]

    # ── 🎯 買い目（新ルール） ──
    heading(doc, "🎯 買い目のご参考（上位推奨レース・現行ルール適用版）", 2)
    para(doc, ["印=◎○▲各1・△1〜5頭可変・🔥穴・❌危険な人気馬／買い目=4類型の多層構成（期待回収率115%以上×ガミ率35%以下で採点・基準未達は見送り）／"
               "予算は自信度連動（10=15,000円/9=12,000円/8=8,000円/7=5,000円、新馬・未勝利60%）／オッズは前日推定。",
               "全レースのSNS投稿案は「週末自信度7以上_SNS投稿案.docx」が正本です。"], size=9, color=(0, 90, 160))
    for k in fkeys:
        recs = races[k]; c, tags = conf[k]; marks, roles, info, kb = flat[k]
        top = recs[0]
        if is_grade(top['レース名']):
            heading(doc, f"{k[1]} R{k[2]:02d} {top['レース名']}【重賞】", 4)
            para(doc, [f"重賞の印・買い目は「{GRADE_DOC}」が正本（20因子採点・外部8サイト照合・手組み多層買い目）。本レポートの機械印は参考値のため掲載しません。"], size=9.5, color=(180, 0, 0))
            continue
        heading(doc, f"{k[1]} R{k[2]:02d} {top['レース名']} — ◎{circ(top['馬番'])}{top['馬名']}（自信度{c}/10）", 4)
        lines = [marks_line(recs, marks, with_danger=True), f"自信度の根拠: {' / '.join(tags)}"]
        st = status(k)
        if st: lines.append(f"⏸ {st}")
        else: lines += kaime_block(recs, roles, kb)
        para(doc, lines, size=9.5)

    # ── SNS対象レース: 推奨レースのうち重賞以外で参加確定の最上位 → なければ全体から自信度最上位の参加レース ──
    cand = [k for k in fkeys if not is_grade(races[k][0]['レース名']) and status(k) is None]
    if not cand:
        cand = sorted([k for k in keys7 if k in keep and not flat[k][3]['skip'] and not is_grade(races[k][0]['レース名'])],
                      key=lambda k: (-conf[k][0], -flat[k][3]['res']['E_rate']))
    if not cand:
        cand = [k for k in fkeys if not is_grade(races[k][0]['レース名'])] or fkeys
    k = cand[0]; recs = races[k]; c, tags = conf[k]; marks, roles, info, kb = flat[k]; top = recs[0]
    st = status(k)
    H = hdr(k, recs, c, day, wd)

    heading(doc, "📱 X投稿文①：予想（現行ルール版）", 2)
    L = H + ["", "🎯 予想印と根拠"]
    for m, num, nm in mark_items(recs, marks):
        r = next(x for x in recs if x['馬番'] == num)
        if m == '🔥穴': L.append(f"🔥穴 {circ(num)} {nm}：{info['ana_reason']}")
        elif m == '❌': L.append(f"❌ {circ(num)} {nm}：推定人気上位でもAI評価{info['danger']+1}位")
        else: L.append(f"{m} {circ(num)} {nm}：{horse_phrase(r, top)}")
    L += ["", ("買い目はスレッドで👇" if not st else f"このレースは{st}。印のみ参考にどうぞ"), "", "#競馬予想 #AI予想 #JRA"]
    para(doc, ["【コピー用】"] + L)

    heading(doc, "📱 X投稿文②：本命の根拠（現行ルール版）", 2)
    L = [f"🏇【{k[1]}{k[2]}R】本命◎{circ(top['馬番'])}{top['馬名']} — 推す理由", f"━━ {k[1]}競馬場 {top['距離']} / {day[:4]}/{day[4:6]}/{day[6:]}（{wd}） ━━", "",
         f"独自AI総合指数{top['総合指数']}で全頭トップ（独自指数{top['独自指数']}・機械学習{top['ML能力%']}%）💪", ""]
    L += [f"・{t}" for t in tags]
    if info['ana'] is not None:
        a = recs[info['ana']]; L += ["", f"🔥穴は{circ(a['馬番'])}{a['馬名']}：{info['ana_reason']}"]
    if info['danger'] is not None:
        dg = recs[info['danger']]; L += [f"❌危険な人気馬は{circ(dg['馬番'])}{dg['馬名']}：人気先行でAI評価{info['danger']+1}位"]
    L += ["", "#競馬 #本命 #AI予想"]
    para(doc, ["【コピー用】"] + L)

    heading(doc, "📱 X投稿文③：買い目（現行ルール版）", 2)
    L = [f"🏇【{k[1]}{k[2]}R {top['レース名']}】買い目", f"━━ {k[1]}競馬場 {top['距離']} / {top['頭数']}頭 ━━", ""]
    if st: L += [f"⏸ {st}", "期待値基準に届かないレースは買いません。印だけ参考にどうぞ🍱"]
    else: L += kaime_block(recs, roles, kb) + ["", "※明日朝の実オッズで最終調整します。"]
    L += ["", "#馬券 #競馬 #JRA"]
    para(doc, ["【コピー用】"] + L)

    heading(doc, "🧵 Threads投稿文（現行ルール版・500字以内）", 2)
    L = H + ["", marks_line(recs, marks), ""]
    if st: L += [f"⏸ {st}"]
    else:
        res = kb['res']
        L.append(f"💰 {kb['arch_name'].split('→')[0].split('・絞り込み')[0]}・合計{len(res['bets'])}点{res['total']:,}円")
        L += [l for l in fmt_bets(res['bets'], recs, with_header=False) if l.startswith("【")]
        L.append(f"期待回収率{res['E_rate']*100:.0f}%（前日推定）")
        sc = scenarios(recs, roles, res, circ)
        if sc: L.append(sc[0])
    L += ["", "全点と根拠はNoteで🎯", "この印、あなたはどう見ますか？👇", "", f"#競馬予想 #{k[1]}競馬場"]
    n = len("".join(L))
    para(doc, [f"【コピー用】（{n}字/500字 {'OK' if n <= 500 else '⚠超過'}）"] + L)

    heading(doc, "📝 Note記事全文（現行ルール版・コピペ用）", 2)
    N = [f"明日{md}（{wd}）{k[1]}{k[2]}R {top['レース名']} AI予想 — アスメシ競馬予想", "",
         "こんにちは、アスメシ競馬予想です🍱", "「明日の飯代」を懸けて、AIとデータで競馬に挑む予想アカウントです。",
         f"明日{md}（{wd}）{k[1]}競馬場 {k[2]}R {top['レース名']}（{top['距離']}・{top['頭数']}頭）の前日版予想をお届けします。",
         f"17ファクター独自指数×機械学習×累積{DB_RECORDS}レコードで全頭を採点しました。", "",
         f"結論を先に言います。本命は{circ(top['馬番'])}{top['馬名']}（総合指数{top['総合指数']}・全頭トップ）。"]
    if info['danger'] is not None:
        dg = recs[info['danger']]; N.append(f"人気上位が予想される{circ(dg['馬番'])}{dg['馬名']}はAI評価{info['danger']+1}位で、今回は❌危険な人気馬とします。")
    N.append(f"そして——{'このレースは' + st + '。買わない判断も含めて公開します。' if st else '買い目は' + kb['arch_name'].split('→')[0] + 'で組みます。'}")
    N += ["", "---", "", "予想印と全頭評価", ""]
    for m, num, nm in mark_items(recs, marks):
        r = next(x for x in recs if x['馬番'] == num)
        if m == '🔥穴': N.append(f"🔥穴 {circ(num)} {nm}　{horse_phrase(r, top)}／{info['ana_reason']}")
        elif m == '❌': N.append(f"❌危険な人気馬 {circ(num)} {nm}　推定人気上位でもAI評価{info['danger']+1}位（{horse_phrase(r, top)}）")
        else: N.append(f"{m} {circ(num)} {nm}　{horse_phrase(r, top)}")
    marked = {num for _, num, _ in mark_items(recs, marks)}
    for r in recs:
        if r['馬番'] not in marked:
            N.append(f"　 {circ(r['馬番'])} {r['馬名']}　総合指数{r['総合指数']}（独自{r['独自指数']}・ML{r['ML能力%']}%）")
    N += ["", f"自信度{c}/10 の根拠: {' / '.join(tags)}", ""] + BABA_NOTE.get(day, []) + ["", "---", "", "買い目", ""]
    if st: N += [f"⏸ {st}", "期待回収率115%以上×ガミ率35%以下の基準に届かないレースは買いません。"]
    else: N += kaime_block(recs, roles, kb)
    N += ["", "オッズは前日時点の推定値で、明日朝に実オッズで最終調整します。", "",
          "明日の自信度7以上レース全部の印と買い目は、ダイジェスト記事でまとめて公開しています。",
          GRADE_LINK[day][0], GRADE_LINK[day][1].replace("🎯", "。"), "",
          "予想が参考になったらフォロー＆いいねお願いします！",
          f"みなさんの{'土曜' if wd == '土' else '日曜'}が楽しいものになりますように🍜🎉", "",
          f"#競馬予想 #AI予想 #JRA #{k[1]}競馬場"]
    para(doc, N, size=9.5)

    doc.save(str(path))
    print(f"置換完了: {path}  （バックアップ: {bk.name}）")
    print(f"推奨レース: {[f'{v}{r}R' for v, r in featured]} → SNSは {k[1]}{k[2]}R {top['レース名']}（{'見送り' if st else kb['arch_name']}）")


if __name__ == '__main__':
    main()
