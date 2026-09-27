# -*- coding: utf-8 -*-
"""
gen_watch_sns.py — 来週重賞の「注目馬」SNS投稿案（X＋Threads・docx）を汎用生成する（2026-09-27）
================================================================================================
入力: {日付フォルダ}/research/next_week_watch.json（レース・型・注目馬・注意馬）
出力: X＝1レース3投稿（①レースの型＋注目馬一覧 ②注目馬詳細・前半 ③後半＋注意馬＋締め）
      Threads＝レース概要1投稿＋**1頭1投稿**（各500字以内・自動チェック）
      docx＝上記すべて（Threadsは見出しに文字数）
ルール（SKILL v4.0 §2 SNS）: 2行ヘッダー→リード→───見出し／印の直後に人気を書かない／サイト名・人物名・内部用語を書かない／
      Xは冒頭140字に言いたいことを置く／noteは作らない（ユーザー指示）
使い方: python gen_watch_sns.py --dir 20261004
"""
from __future__ import annotations
import argparse, json, sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
BASE = Path.home() / "Desktop" / "競馬予想レポート"
BAN = ["netkeiba", "SPAIA", "うましる", "競馬ラボ", "重賞ナビ", "アジフライ", "しろクロ", "うまログ", "ウマキング", "カリスマ", "PDCA", "EV", "sc[", "インフルエンサー", "指数"]

ap = argparse.ArgumentParser()
ap.add_argument("--dir", required=True)
a = ap.parse_args()
D = BASE / a.dir
J = json.loads((D / "research" / "next_week_watch.json").read_text(encoding="utf-8"))
DATE = J["post_date_label"]
TAGS = J["hashtags"]


def header(r):
    return f"🏇【{r['name']} {r['grade']}】{DATE}\n━━ {r['venue']} {r['course']} / 登録{r['entries']}頭 ━━"


def sec(title):
    return f"─────\n{title}\n─────"


def horse_block(h, with_jockey=True):
    head = h["horse"] + ("" if not h.get("age") or h["age"] == "—" else f"（{h['age']}") + \
        (f"・{h['kin']}" if h.get("kin") else "") + ("）" if h.get("age") and h["age"] != "—" else "")
    if with_jockey and h.get("jockey") and h["jockey"] != "—":
        head += f"　想定 {h['jockey']}"
    lines = [f"◆ {head}"] + ([h["lead"]] if h.get("lead") else []) + [f"・{x}" for x in h["detail"]] + [f"→ {h['stance']}"]
    return "\n".join(lines)


def x_posts(r):
    names = "／".join(h["horse"] for h in r["watch"])
    p1 = "\n".join([header(r), "", r["lead"], "", sec("レースの型（過去10回の自前集計）")] + [f"・{x}" for x in r["profile"]] +
                   ["", sec("注目馬"), names, "", "各馬の根拠はスレッドで👇", f"{TAGS} {r['tag']}"])
    half = (len(r["watch"]) + 1) // 2
    p2 = "\n".join([header(r), "【注目馬 詳細①】", ""] + [horse_block(h) + "\n" for h in r["watch"][:half]] + [f"{TAGS} {r['tag']}"])
    p3_parts = [header(r), "【注目馬 詳細②】", ""] + [horse_block(h) + "\n" for h in r["watch"][half:]]
    if r.get("caution"):
        p3_parts += [sec("人気になったら注意したい馬")] + [horse_block(h, with_jockey=False) + "\n" for h in r["caution"]]
    p3_parts += ["登録段階なので、枠順と当日の馬体重が出たら最終版を出します。", f"{TAGS} {r['tag']}"]
    p3 = "\n".join(p3_parts)
    return [p1, p2, p3]


def threads_posts(r):
    total = 1 + len(r["watch"]) + len(r.get("caution", []))
    out = []
    t0 = "\n".join([header(r), f"（1/{total}）レースの型", "", r["lead"], ""] + [f"・{x}" for x in r["profile"][:3]] +
                   ["", "注目馬は1頭ずつ次の投稿で👇", f"{TAGS} {r['tag']}"])
    out.append(t0)
    i = 2
    for h in r["watch"]:
        body = "\n".join([header(r), f"（{i}/{total}）{h['horse']}", "", h["lead"], ""] + [f"・{x}" for x in h["detail"]] + ["", f"→ {h['stance']}", f"{r['tag']}"])
        while len(body) > 500 and len(h["detail"]) > 2:           # 長ければ根拠を1行ずつ削る
            h = dict(h, detail=h["detail"][:-1])
            body = "\n".join([header(r), f"（{i}/{total}）{h['horse']}", "", h["lead"], ""] + [f"・{x}" for x in h["detail"]] + ["", f"→ {h['stance']}", f"{r['tag']}"])
        out.append(body)
        i += 1
    for h in r.get("caution", []):
        body = "\n".join([header(r), f"（{i}/{total}）注意したい人気馬：{h['horse']}", ""] + [f"・{x}" for x in h["detail"]] + ["", f"→ {h['stance']}", f"{r['tag']}"])
        out.append(body)
        i += 1
    return out


def check(text, label):
    hits = [b for b in BAN if b in text]
    if hits:
        print(f"  ⚠ {label}: 禁止語 {hits}")
    return hits


# ── 生成 ──────────────────────────────────────────────────────────
from docx import Document
from docx.shared import Pt
doc = Document()
doc.styles["Normal"].font.size = Pt(10)
doc.add_heading(f"来週重賞 注目馬 SNS投稿案（X＋Threads）— {DATE}", 0)
doc.add_paragraph("登録段階（枠順・オッズ未定）の注目馬。数字は蓄積データと過去10回の自前集計。noteは作らない（指示）。"
                  "Xは1レース3投稿、Threadsはレース概要1投稿＋1頭1投稿（500字以内）。")
nban = 0
for r in J["races"]:
    doc.add_heading(f"{r['name']}（{r['grade']}・{r['venue']} {r['course']}）", 1)
    doc.add_heading("X投稿（3本）", 2)
    for k, p in enumerate(x_posts(r), 1):
        first = p.split("\n")
        lead_chars = len("".join(first[:4]))
        doc.add_heading(f"X{k}　{len(p)}字（冒頭4行 {lead_chars}字）", 3)
        doc.add_paragraph(p)
        nban += len(check(p, f"{r['name']} X{k}"))
    doc.add_heading("Threads投稿（1頭1投稿）", 2)
    for k, p in enumerate(threads_posts(r), 1):
        ok = "OK" if len(p) <= 500 else "⚠超過"
        doc.add_heading(f"Threads {k}/{len(threads_posts(r))}　{len(p)}字/500字 {ok}", 3)
        doc.add_paragraph(p)
        nban += len(check(p, f"{r['name']} T{k}"))
    print(f"■ {r['name']}: X {[len(p) for p in x_posts(r)]}字 ／ Threads {[len(p) for p in threads_posts(r)]}字")
out = D / f"{a.dir}_来週重賞注目馬_SNS投稿案.docx"
try:
    doc.save(out)
except PermissionError:
    out = D / f"{a.dir}_来週重賞注目馬_SNS投稿案_new.docx"
    doc.save(out)
print(f"禁止語検出 {nban}件 ／ 保存: {out}")
