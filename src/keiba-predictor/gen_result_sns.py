"""
2026/04/26 事前予想の結果報告 SNS投稿案
「予想投稿を見てくれた人への結果報告」スタイル
"""
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.stdout.reconfigure(encoding='utf-8')

OUTPUT_DIR = Path(r"C:\Users\User\Desktop\競馬予想レポート\20260426")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_PATH = OUTPUT_DIR / "20260426_結果報告SNS投稿案_v3.docx"


def set_font(run, size_pt=10.5, bold=False, color=None, font_name="游ゴシック"):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:hint'), 'eastAsia')
    rFonts.set(qn('w:eastAsia'), font_name)
    rpr.insert(0, rFonts)


def add_label(doc, text, color):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_font(run, size_pt=10, bold=True, color=color)


def add_box(doc, content, bg_color, note=None):
    """投稿ボックス"""
    if note:
        pn = doc.add_paragraph()
        run = pn.add_run(f"  ※{note}")
        set_font(run, size_pt=8.5, color=RGBColor(0x88, 0x88, 0x88))

    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = 'Table Grid'
    cell = tbl.rows[0].cells[0]
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), bg_color)
    tcPr.append(shd)
    for line in content.strip().split('\n'):
        p = cell.add_paragraph()
        run = p.add_run(line)
        set_font(run, size_pt=10.5)
    doc.add_paragraph()


def add_section_bar(doc, text, color=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_font(run, size_pt=12, bold=True,
             color=color or RGBColor(0x0F, 0x47, 0x61))


# ══════════════════════════════════════════════════════
# 投稿テキスト（事前予想の結果報告スタイル）
# ══════════════════════════════════════════════════════

# ─────────────────────────────
# 【投稿順①】両レース速報（最速で出す1本）
# ─────────────────────────────
POST_SOKUHOH_X = """\
🍱【結果報告】本日のG2予想 2レース的中🔥

【フローラS】
◎ラフターラインズ → 🥇1着✅
○エンネ → 🥈2着✅
単勝・馬連・ワイド的中💰
回収率225%

【マイラーズC】
◎アドマイヤズーム → 🥇1着✅
▲ベラジオボンド → 🥉3着✅
単勝410円・ワイド900円的中💰
回収率303%

G2 2レース合計 投資2万→回収52,800円🔥
見ててくれた方ありがとうございます🙏

#アスメシ競馬 #的中 #競馬"""

POST_SOKUHOH_THREADS = """\
🍱【結果報告】2026/04/26 G2 2レース的中🔥

今日の事前予想、結果が出ました！

━【フローラS 東京芝2000m】━
◎5番 ラフターラインズ（1人気）→ 🥇1着✅
○13番 エンネ（3人気）→ 🥈2着✅
単勝220円・馬連970円・ワイド310円 的中💰
投資10,000円 → 回収22,500円（回収率225%）🎯

━【マイラーズC 京都芝1600m】━
◎9番 アドマイヤズーム（1人気）→ 🥇1着✅
▲7番 ベラジオボンド（6人気）→ 🥉3着✅
単勝410円・ワイド900円 的中💰
投資10,000円 → 回収30,300円（回収率303%）🎯

2レース合計 投資20,000円 → 回収52,800円（回収率264%）🔥

今日は2重賞連続で本命が1着でした🎉
予想見てくれた方、ありがとうございます！
来週もよろしくお願いします🙏

#アスメシ競馬 #AI競馬予想 #的中報告"""

# ─────────────────────────────
# 【投稿順②】フローラS 結果詳細
# ─────────────────────────────
POST_FLORA_X1 = """\
🌸【フローラS G2 結果報告】

今朝予想した通り…

◎5番 ラフターラインズ 🥇1着✅
○13番 エンネ 🥈2着✅
（3着 リアライズルミナス）

本命と対抗でワンツー決着🎯
D.レーン騎手が直線で一気に突き抜けました！

事前予想を信じてくれた方、
一緒に獲れましたか？😊

#フローラステークス #アスメシ競馬 #的中"""

POST_FLORA_X2 = """\
💰【フローラS 買い目結果】

✅ 単勝 5番（220円）→ 3,000円 → 6,600円
✅ ワイド 5-13（310円）→ 2,000円 → 6,200円
✅ 馬連 5-13（970円）→ 1,000円 → 9,700円

投資10,000円 → 回収22,500円
回収率 225%🎯

消し判断の2人気ラベルセーヌが5着で
消し正解も気持ちよかった👊笑

次はオークス！また予想します🌸

#フローラステークス #アスメシ競馬"""

POST_FLORA_THREADS = """\
🌸【フローラS G2 結果詳細】2026/04/26

今朝投稿した予想の結果をご報告します！

【予想】
◎5番 ラフターラインズ → 🥇1着✅
○13番 エンネ → 🥈2着✅
消し2番 ラベルセーヌ → 5着✅

本命と対抗のワンツーフィニッシュ🎉
AIの能力スコア1位・2位がそのまま決着！

【買い目結果】
✅ 単勝5番（220円）→ 3,000円 → 6,600円
✅ ワイド5-13（310円）→ 2,000円 → 6,200円
✅ 馬連5-13（970円）→ 1,000円 → 9,700円

投資10,000円 → 回収22,500円（回収率225%）🎯

消しにした2人気ラベルセーヌは5着で
消し判断も大正解でした！

次はオークス本番！引き続き応援よろしくお願いします🌸

#フローラステークス #アスメシ競馬 #AI競馬予想"""

# ─────────────────────────────
# 【投稿順③】マイラーズC 結果詳細
# ─────────────────────────────
POST_MILERS_X1 = """\
🏇【マイラーズC G2 結果報告】

今朝予想した通り…

◎9番 アドマイヤズーム 🥇1着✅
▲7番 ベラジオボンド 🥉3着✅
（2着 1番ドラゴンブースト 9人気）

武豊騎手がテンから先手を奪い
後続を完封！完璧なレースでした🔥

本命1着＋3着ダブル的中！
単勝・ワイドしっかり回収しました😤

#マイラーズカップ #アスメシ競馬 #的中"""

POST_MILERS_X2 = """\
💰【マイラーズC 買い目結果】

✅ 単勝 9番（410円）→ 3,000円 → 12,300円
✅ ワイド 9-7（900円）→ 2,000円 → 18,000円

投資10,000円 → 回収30,300円
回収率 303%🔥

武豊×本命1着で単勝だけで12,300円！
ワイドも18,000円しっかり回収😤

#マイラーズカップ #アスメシ競馬"""

POST_MILERS_THREADS = """\
🏇【マイラーズC G2 結果詳細】2026/04/26

今朝投稿した予想の結果をご報告します！

【予想】
◎9番 アドマイヤズーム → 🥇1着✅
▲7番 ベラジオボンド → 🥉3着✅
消し：シックスペンス・エルトンバローズ → 着外✅✅

武豊騎手がテン最速で逃げ切り！
AIの能力1位が完璧な逃げ切りを決めました🔥

【買い目結果】
✅ 単勝9番（410円）→ 3,000円 → 12,300円
✅ ワイド9-7（900円）→ 2,000円 → 18,000円

投資10,000円 → 回収30,300円（回収率303%）🔥

消しにしたシックスペンス（ダート帰り）・
エルトンバローズ（前走大惨敗）も全員着外で
消し判断も全部正解でした👊

来週もAI予想でガチ勝負します💪

#マイラーズカップ #アスメシ競馬 #AI競馬予想"""

# ─────────────────────────────
# 【投稿順④】締めのまとめ
# ─────────────────────────────
POST_MATOME_X = """\
📊【本日のまとめ】2026/04/26

🌸フローラS
◎ラフターラインズ 1着✅ ○エンネ 2着✅
単勝220円・馬連970円・ワイド310円 的中💰
投資1万→回収22,500円（225%）

🏇マイラーズC
◎アドマイヤズーム 1着✅ ▲ベラジオボンド 3着✅
単勝410円・ワイド900円 的中💰
投資1万→回収30,300円（303%）

2レース合計
投資20,000円 → 回収52,800円
回収率 264%🔥 G2 2重賞制覇！

フォロー＆いいね！で来週の予想も
見逃さないでください🙏

#アスメシ競馬 #AI競馬予想 #重賞的中"""

POST_MATOME_THREADS = """\
📊【本日 全予想まとめ】2026/04/26

今日は重賞2レースを予想しました。最終結果をまとめてご報告します！

━━━━━━━━━━━━━━━
🌸 フローラステークス G2
東京競馬場 芝2000m
━━━━━━━━━━━━━━━
◎ラフターラインズ → 1着✅ 本命的中
○エンネ → 2着✅ 対抗的中
単勝220円・馬連970円・ワイド310円 的中💰
投資10,000円 → 回収22,500円（回収率225%）🎯

━━━━━━━━━━━━━━━
🏇 読売マイラーズカップ G2
京都競馬場 芝1600m
━━━━━━━━━━━━━━━
◎アドマイヤズーム → 1着✅ 本命的中
▲ベラジオボンド → 3着✅ 三番手的中
単勝410円・ワイド900円 的中💰
投資10,000円 → 回収30,300円（回収率303%）🎯

━━━━━━━━━━━━━━━
2レース合計 投資20,000円 → 回収52,800円
総回収率 264%🔥

G2 2レース連続で本命1着🎉
今日は本当によく当たりました！

今日の予想を見てくれた皆さん
一緒に勝てていたら嬉しいです😊

来週も競馬予想を投稿しますので
フォローよろしくお願いします🙏

#アスメシ競馬 #AI競馬予想 #重賞的中 #フローラステークス #マイラーズカップ"""


# ══════════════════════════════════════════════════════
# Word生成
# ══════════════════════════════════════════════════════

def main():
    doc = Document()

    # ── 表紙 ──
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("結果報告 SNS投稿案")
    set_font(run, size_pt=16, bold=True, color=RGBColor(0x0F, 0x47, 0x61))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("2026年4月26日（日）フローラS G2 ／ マイラーズC G2")
    set_font(run, size_pt=11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("【投稿順】① 速報  →  ② フローラS詳細  →  ③ マイラーズC詳細  →  ④ まとめ")
    set_font(run, size_pt=9, color=RGBColor(0x88, 0x44, 0x00))

    doc.add_paragraph()

    # ══ ①速報 ══
    add_section_bar(doc, "━━━ ① 速報（レース終了後すぐ） ━━━",
                    RGBColor(0xCC, 0x22, 0x22))
    doc.add_paragraph()

    add_label(doc, "【X投稿】", RGBColor(0x1D, 0xA1, 0xF2))
    add_box(doc, POST_SOKUHOH_X, "FFF8E1",
            note="両レース終了後、なるべく早く投稿")

    add_label(doc, "【Threads】", RGBColor(0x00, 0x00, 0x00))
    add_box(doc, POST_SOKUHOH_THREADS, "F5F5F5",
            note="X投稿と同タイミングで投稿")

    doc.add_page_break()

    # ══ ②フローラS ══
    add_section_bar(doc, "━━━ ② フローラS 結果詳細 ━━━",
                    RGBColor(0xC2, 0x18, 0x5B))
    doc.add_paragraph()

    add_label(doc, "【X投稿①　着順報告】", RGBColor(0x1D, 0xA1, 0xF2))
    add_box(doc, POST_FLORA_X1, "EBF5FB",
            note="速報の30分後くらいに投稿")

    add_label(doc, "【X投稿②　買い目報告】", RGBColor(0x1D, 0xA1, 0xF2))
    add_box(doc, POST_FLORA_X2, "EBF5FB",
            note="①の直後に続けて投稿")

    add_label(doc, "【Threads（1投稿でまとめ）】", RGBColor(0x00, 0x00, 0x00))
    add_box(doc, POST_FLORA_THREADS, "F5F5F5",
            note="XのスレッドよりThreadsは1投稿でまとめる")

    doc.add_page_break()

    # ══ ③マイラーズC ══
    add_section_bar(doc, "━━━ ③ マイラーズC 結果詳細 ━━━",
                    RGBColor(0x0D, 0x47, 0xA1))
    doc.add_paragraph()

    add_label(doc, "【X投稿①　着順報告】", RGBColor(0x1D, 0xA1, 0xF2))
    add_box(doc, POST_MILERS_X1, "EBF5FB",
            note="フローラSの投稿から30〜60分後")

    add_label(doc, "【X投稿②　買い目報告・収支】", RGBColor(0x1D, 0xA1, 0xF2))
    add_box(doc, POST_MILERS_X2, "EBF5FB",
            note="①の直後に続けて投稿")

    add_label(doc, "【Threads（1投稿でまとめ）】", RGBColor(0x00, 0x00, 0x00))
    add_box(doc, POST_MILERS_THREADS, "F5F5F5")

    doc.add_page_break()

    # ══ ④まとめ ══
    add_section_bar(doc, "━━━ ④ 締めのまとめ（最後に投稿） ━━━",
                    RGBColor(0x1B, 0x5E, 0x20))
    doc.add_paragraph()

    add_label(doc, "【X投稿】", RGBColor(0x1D, 0xA1, 0xF2))
    add_box(doc, POST_MATOME_X, "EBF5FB",
            note="全投稿の最後・夜に投稿。フォロー促進CTA付き")

    add_label(doc, "【Threads】", RGBColor(0x00, 0x00, 0x00))
    add_box(doc, POST_MATOME_THREADS, "F5F5F5",
            note="Threadsでの1日の締め投稿")

    doc.save(str(OUTPUT_PATH))
    print(f"保存完了: {OUTPUT_PATH}")


if __name__ == '__main__':
    main()
