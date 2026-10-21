# -*- coding: utf-8 -*-
"""
gen_pick_sns.py — 全レースから期待値の取れそうな数レースを選び、印・短評・買い目のSNS投稿案にする（2026-10-10新規）
==========================================================================================================
ユーザー指示（2026-10-10）: 「平場も含めた全レースから、期待値が取れる買えるレースをSNSで発信したい。数レースに絞って、
本命・対抗・穴の印の理由を短評で」
  ・印・確率・買い目は予想カード research/cards_{日付}.json（predict_day.py の出力）そのまま
  ・選ぶレースと短評は research/pick_content.json に手書き（JRA-VANの数字で確認できることだけ）
      {"title","post_date_label","intro","graded_note","races":[{"key":"東京1","label","why","comments":{馬番:短評},"fav_note"}]}
  ・候補の並べ方は --rank で表示: 🎯勝負レース（市場との比1.3以上の馬連）→ ◎の比（モデル÷市場）の高い順
  ・投稿: X/Threads 共通で「まとめ1本」＋ 1レース2本（印と短評／印の再掲と買い目）。1本500字以内
出力: {日付}/{日付}_注目レース_SNS投稿案.txt（BOM付きUTF-8）と .docx
使い方: python gen_pick_sns.py --dir 20261011 --rank   # 候補の一覧
        python gen_pick_sns.py --dir 20261011          # 投稿案
"""
from __future__ import annotations
import argparse, collections, json, re, sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.stdout.reconfigure(encoding="utf-8")
from kaime_format import circ

BASE = Path.home() / "Desktop" / "競馬予想レポート"
ORDER = "◎○▲△🔥"
SEP = "─────"
TAGS = "#競馬予想 #AI予想 #JRA"
BAN = ("EV", "PDCA", "指数", "sc[", "インフルエンサー", "YouTube", "netkeiba")


def short_name(c):
    return re.sub(r"\s*出馬表.*$", "", c["name"]).replace("３歳以上", "3歳以上").replace("２歳", "2歳").replace("１勝", "1勝").replace("２勝", "2勝").strip()


def rank(cards):
    rows = []
    for c in cards:
        H = {h["uma"]: h for h in c["horses"]}
        hon = next(int(u) for u, m in c["marks"].items() if m == "◎")
        val = [b for b in c["ref_bets"] if b.get("value")]
        rows.append((bool(val), H[hon]["ratio"], c["key"], short_name(c), H[hon]["name"], H[hon]["pop"], [b["combo"] for b in val]))
    rows.sort(key=lambda r: (not r[0], -r[1]))
    for v, r_, k, nm, hn, pop, vv in rows:
        print(f"  {'🎯' if v else '  '} {k:<5} {nm[:16]:<16} ◎{hn}（{pop}番人気・比{r_:.2f}）" + (f" 勝負の馬連 {vv}" if vv else ""))


def bets_lines(c):
    agg = collections.OrderedDict()
    for b in sorted(c["ref_bets"], key=lambda b: ("単勝", "馬連", "馬単", "ワイド", "3連複").index(b["type"])):   # 🎯の馬連を通常の馬連の隣に
        k = (b["type"], tuple(b["combo"]))
        a = agg.setdefault(k, dict(b, stake=0, value=False))
        a["stake"] += b["stake"]
        a["value"] = a["value"] or bool(b.get("value"))
    L = []
    for (t, combo), b in agg.items():
        o = b.get("odds_now")
        est = o is None
        o = o or b["odds_est"]
        L.append(f"{'🎯' if b['value'] else ''}【{t}】{'-'.join(circ(x) for x in combo)} {b['stake']:,}円（的中{b['p_model']*100:.0f}%×{'想定' if est and t != '単勝' else ''}{o:.1f}倍＝{b['p_model']*o:.2f}）")
    return L, sum(b["stake"] for b in agg.values())


def main(A):
    D = BASE / A.dir
    cards = json.loads((D / "research" / f"cards_{A.dir}.json").read_text(encoding="utf-8"))["cards"]
    if A.rank:
        rank(cards)
        return
    P = json.loads((D / "research" / "pick_content.json").read_text(encoding="utf-8"))
    C = {c["key"]: c for c in cards}
    posts = []
    head0 = f"🏇【{P['title']}】{P['post_date_label']}"
    L = [head0, f"━━ 東京・京都 全{len(cards)}レースから{len(P['races'])}つ ━━", ""]
    for r in P["races"]:
        c = C[r["key"]]
        hon = next(int(u) for u, m in c["marks"].items() if m == "◎")
        h = next(x for x in c["horses"] if x["uma"] == hon)
        L.append(f"{'🎯' if any(b.get('value') for b in c['ref_bets']) else '・'}{r['label']}　◎{circ(hon)}{h['name']}")
    L += ["", P["intro"], "", SEP,
          "印は人気ではなく「勝つ確率が人気より高いか（割安さ）」の順。「確率16%（人気13%）」は、こちらの勝つ確率と、オッズが示す勝つ確率です。🎯は人気に対して1.3倍以上割安な馬連が出た勝負レースです（過去3年の実際のオッズで回収率117%）。",
          "オッズは前日時点の値です。当日朝に実際のオッズで計算し直します。", P.get("graded_note", ""), TAGS]
    posts.append(("まとめ", L))
    for r in P["races"]:
        c = C[r["key"]]
        H = {h["uma"]: h for h in c["horses"]}
        marked = sorted(((m, int(u)) for u, m in c["marks"].items()), key=lambda x: (ORDER.index(x[0]), x[1]))
        head = [f"🏇【{r['label']}】{P['post_date_label']}", f"━━ {c['venue']} {c['dist']} / {c['n']}頭 ━━"]
        shobu = any(b.get("value") for b in c["ref_bets"])
        L = head + ([""] + ["🎯勝負レース"] if shobu else []) + [""] + [f"{m}{circ(u)}{H[u]['name']}" for m, u in marked] + ["", r["why"], "", SEP]
        for m, u in marked:
            cm = r["comments"].get(str(u))
            if not cm:
                continue
            L += [f"{m}{circ(u)} {H[u]['pop']}番人気・確率{H[u]['p_model']*100:.0f}%（人気{H[u]['p_market']*100:.0f}%）", f"　{cm}"]
        if r.get("fav_note"):
            L += ["", r["fav_note"]]
        L += ["", TAGS]
        posts.append((f"{r['label']} 印と短評", L))
        bl, tot = bets_lines(c)
        L = head + [""] + [f"{m}{circ(u)}{H[u]['name']}" for m, u in marked] + ["", f"💰 買い目 {len(bl)}点・{tot:,}円"] + bl + ["",
             "◎の単勝と馬連に、ワイドで相手を広く。" + ("🎯の馬連は人気より割安な組み合わせです。" if shobu else ""),
             "想定オッズは単勝から計算した値。当日朝に実際のオッズで計算し直します。", TAGS]
        posts.append((f"{r['label']} 買い目", L))
    # 検査
    probs = []
    out_txt, n = [], 0
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading(f"{P['title']} SNS投稿案 — {P['post_date_label']}", 0)
    doc.add_paragraph("X・Threads 共通。上から順にスレッドで投稿。印・確率・買い目は予想エンジンの値、短評は JRA-VAN の成績・調教から。")
    for t, L in posts:
        body = "\n".join(x for x in L if x is not None)
        body = re.sub(r"\n{3,}", "\n\n", body)
        n += 1                                        # 500字を超えたら検査で知らせる（余白は詰めない＝読みやすさ優先。短評を短くする）
        if len(body) > 500:
            probs.append(f"{t}: {len(body)}字")
        for w in BAN:
            if w in body:
                probs.append(f"{t}: 禁止語 {w}")
        for amt in re.findall(r"([\d,]+)円", body):
            v = int(amt.replace(",", ""))
            if 0 < v < 100000 and v % 100:
                probs.append(f"{t}: 100円単位でない {amt}円")
        out_txt += [f"■ {n}. {t}（{len(body)}字）", body, ""]
        doc.add_heading(f"{n}. {t}　{len(body)}字", 3)
        tb = doc.add_table(rows=1, cols=1); tb.style = "Table Grid"; tb.cell(0, 0).text = body
        doc.add_paragraph()
    stem = D / f"{A.dir}_注目レース_SNS投稿案"
    stem.with_suffix(".txt").write_text("\n".join(out_txt), encoding="utf-8-sig")
    doc.save(stem.with_suffix(".docx"))
    print(f"{n}本 → {stem}.txt / .docx")
    print("検査: " + ("問題なし" if not probs else " ／ ".join(probs)))


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", required=True)
    ap.add_argument("--rank", action="store_true")
    main(ap.parse_args())
