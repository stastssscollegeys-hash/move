# -*- coding: utf-8 -*-
"""
word_order_by_ev.py — Word出力を「期待値の取れるレース（上）／参考（下）」に並べ替える（2026-10-12新規）
=====================================================================================================
ユーザー指示（2026-10-12）: 「Wordの出し方は、期待値の取れるレースは上にまとめて、それ以外は参考として下にまとめる」
  期待値の取れるレース = 予想カードの ref_bets に「勝負馬連（value）」があり、実オッズ（odds_now・無ければ想定）で
                         的中確率×オッズ（期待値）が 1 を超えるもの。検証で効果が確認できているのは
                         「馬連・市場との比≧1.3」だけなので、ワイド・穴の期待値だけが高いレースは入れない。
  対象: {日付}/{日付}_平場_予想とSNS投稿案.docx ・ {日付}/{日付}_予想カード_エンジン版.docx
        （gen_flat_sns.py / predict_day.py の出力を後から並べ替える。元のコードは変えない）
  並び: ◆期待値の取れるレース（発走順）→ ◇参考（発走順）。平場docxは「まとめ投稿」の勝負レース一覧と全レース一覧表も揃える
  何度回しても同じ結果（見出しがあれば何もしない）。元ファイルは {日付}/research/word_before_order/ に退避する
使い方: python word_order_by_ev.py --dir 20261012
"""
from __future__ import annotations
import argparse, copy, json, re, shutil, sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
from docx.text.paragraph import Paragraph

BASE = Path.home() / "Desktop" / "競馬予想レポート"
RT_ENTRY = Path("C:/keiba_db/jv_raw/RT/entry")
RX = re.compile(r"^(京都|東京|中山|阪神|新潟|福島|中京|小倉|札幌|函館)(\d+)R")
RT = re.compile(r"【(京都|東京|中山|阪神|新潟|福島|中京|小倉|札幌|函館)(\d+)R")
H_EV = "◆ 期待値の取れるレース"
H_REF = "◇ 参考"


def ev_keys(cards):
    out = set()
    for c in cards:
        for b in c.get("ref_bets") or []:
            o = b.get("odds_now") or b.get("odds_est")
            if b.get("value") and o and b["p_model"] * o > 1:
                out.add(c["key"])
    return out


def start_times(date):
    """{(場, R): 'HHMM'}（RT の RA の 874〜877 バイト目）"""
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    import target_import as T
    out = {}
    f = RT_ENTRY / f"RA_{date}.txt"
    if f.exists():
        b = f.read_bytes()
        for i in range(0, len(b) - 1271, 1272):
            r = b[i:i + 1272]
            x = T.parse_ra(r)
            if x and x["R"]:
                out[f'{x["venue"]}{x["R"]}'] = r[873:877].decode("ascii", "ignore")
    return out


def set_text(p, text):
    runs = p.runs
    if runs:
        for r in runs[1:]:
            r.text = ""
        runs[0].text = text
    else:
        p.add_run(text)


def reorder(path, ev, times):
    d = Document(path)
    body = d.element.body
    kids = [el for el in body.iterchildren()]
    sect = kids[-1] if kids and kids[-1].tag.endswith("}sectPr") else None
    items = [el for el in kids if el is not sect]
    if any("".join(el.itertext()).strip().startswith(H_EV) for el in items if el.tag.endswith("}p")):
        return "済み（見出しあり）"
    pre, blocks, order, cur, venue_head = [], {}, [], None, None
    for el in items:
        tag = el.tag.split("}")[1]
        t = "".join(el.itertext()).strip()
        if tag == "p":
            m = RX.match(t)
            if m:
                cur = m.group(1) + m.group(2)
            elif t.startswith("■ "):
                venue_head = venue_head if venue_head is not None else el
                cur = None
                continue                                           # 並べ替えると意味が変わる見出しは外す
        elif tag == "tbl":
            m = RT.search(t[:40])
            if m:
                cur = m.group(1) + m.group(2)
        if cur is None:
            pre.append(el)
        else:
            if cur not in blocks:
                blocks[cur] = []
                order.append(cur)
            blocks[cur].append(el)
    if len(blocks) < 2:
        return "対象レースなし"
    key = lambda k: (times.get(k, "9999"), k)
    top = sorted((k for k in order if k in ev), key=key)
    rest = sorted((k for k in order if k not in ev), key=key)
    tmpl = venue_head if venue_head is not None else next((blocks[k][0] for k in order if blocks[k][0].tag.endswith("}p")), None)

    def heading(text):
        el = copy.deepcopy(tmpl)
        p = Paragraph(el, d._body)
        set_text(p, text)
        return el

    for el in items:
        body.remove(el)
    new = list(pre)
    new.append(heading(f"{H_EV}（{len(top)}レース）" if top else f"{H_EV}（該当なし）"))
    for k in top:
        new += blocks[k]
    new.append(heading(f"{H_REF}（{len(rest)}レース）：期待値が1を超える勝負馬連が出ていないレース"))
    for k in rest:
        new += blocks[k]
    for el in new:
        if sect is not None:
            sect.addprevious(el)
        else:
            body.append(el)
    # 全レース一覧の表: 期待値の取れるレースを上に
    for tb in d.tables:
        rows = list(tb.rows)
        if len(rows) < 3 or not RX.match(rows[1].cells[0].text.strip()):
            continue
        head, data = rows[0], rows[1:]
        def rk(r):
            m = RX.match(r.cells[0].text.strip())
            k = m.group(1) + m.group(2) if m else ""
            return (0 if k in ev else 1, times.get(k, "9999"), k)
        for r in data:
            tb._tbl.remove(r._tr)
        for r in sorted(data, key=rk):
            tb._tbl.append(r._tr)
    # まとめ投稿（平場）: 勝負レースの一覧を期待値の取れるレースだけにする
    for tb in d.tables:
        if len(tb.rows) == 1 and len(tb.columns) == 1 and "平場 予想まとめ" in tb.cell(0, 0).text[:40]:
            for p in tb.cell(0, 0).paragraphs:
                lines = p.text.split("\n")
                if not any(x.startswith("🎯 勝負レース") for x in lines):
                    continue
                out, kept, i0 = [], 0, None
                for x in lines:
                    m = RX.match(x[1:]) if x.startswith("・") else None
                    if x.startswith("🎯 勝負レース"):
                        out.append("🎯 期待値の取れるレース（人気より割安な馬連が出て、実オッズでも期待値が1を超えたレース）")
                    elif m and x.startswith("・"):
                        if m.group(1) + m.group(2) in ev:
                            out.append(x); kept += 1
                    elif x.startswith("それ以外の"):
                        if kept == 0:
                            out.append("該当するレースはありません")
                        out.append(f"それ以外の{len(order) - kept}レースは、参考として印と標準の買い目だけを載せています。")
                    else:
                        out.append(x)
                set_text(p, "\n".join(out))
    d.save(path)
    return f"期待値あり{len(top)}・参考{len(rest)}"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", required=True)
    a = ap.parse_args()
    D = BASE / a.dir
    cards = json.loads((D / "research" / f"cards_{a.dir}.json").read_text(encoding="utf-8"))["cards"]
    ev = ev_keys(cards)
    times = start_times(a.dir)
    print("期待値の取れるレース:", sorted(ev, key=lambda k: (times.get(k, "9999"), k)) or "なし")
    bk = D / "research" / "word_before_order"
    bk.mkdir(parents=True, exist_ok=True)
    for fn in (f"{a.dir}_平場_予想とSNS投稿案.docx", f"{a.dir}_予想カード_エンジン版.docx"):
        p = D / fn
        if not p.exists():
            print(fn, "なし"); continue
        if not (bk / fn).exists():
            shutil.copy(p, bk / fn)
        print(fn, "→", reorder(p, ev, times))


if __name__ == "__main__":
    main()
