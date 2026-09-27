# -*- coding: utf-8 -*-
"""
アスメシ競馬予想 SNS投稿案テンプレート
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
使い方:
  1. ── RACE CONFIG ── セクションを今週のレースに書き換える
  2. ── PICKS ── セクションに◎○▲△を記入する
  3. ── ANALYSIS ── セクションに各馬の分析文を記入する
  4. python gen_race_template.py で実行
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
"""
import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from pathlib import Path

# ============================================================
# ── RACE CONFIG ── ここを書き換える
# ============================================================
RACE_DATE   = "20260503"          # YYYYMMDD
RACE_NAME   = "天皇賞・春"        # レース正式名
RACE_GRADE  = "G1"                # G1 / G2 / G3
RACE_VENUE  = "京都競馬場"        # 競馬場名
RACE_COURSE = "芝3200m"           # コース・距離
RACE_COUNT  = "15頭立て"          # 頭数
RACE_TIME   = "15:40"             # 発走時刻

# noteヘッダー画像の出力先（生成後にここへ保存される）
NOTE_IMAGE_OUT = rf"C:\Users\User\Desktop\競馬予想レポート\{RACE_DATE}\note_header_{RACE_DATE}.png"

# Word出力先
OUT = Path(rf"C:\Users\User\Desktop\競馬予想レポート\{RACE_DATE}\{RACE_DATE}_{RACE_NAME}_SNS投稿案.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

# ============================================================
# ── PICKS ── 予想印を記入する
# ============================================================
# 書式: (印, 馬番, 馬名, 騎手, 人気)
PICKS = [
    ("◎",  "3",  "アドマイヤテラ",    "武豊",       "1番人気"),
    ("○",  "7",  "クロワデュノール",   "北村友一",   "2番人気"),
    ("▲",  "4",  "アクアヴァーナル",   "松山弘平",   "4番人気"),
    ("△",  "6",  "エヒト",            "川田将雅",   ""),
    ("△", "14",  "ホーエリート",       "戸崎圭太",   "9番人気"),
]

# AI能力スコア（上位3頭）
SCORES = {
    "クロワデュノール":  "32.1%",
    "アドマイヤテラ":    "28.4%",
    "アクアヴァーナル":  "16.8%",
}

# ============================================================
# ── BETTING ── 買い目と資金配分を記入する（合計は BUDGET に合わせる）
# ============================================================
BUDGET = 10000  # 合計予算（円）

# 書式: (馬券種, 組み合わせ, 金額, 備考)
BETS = [
    ("単勝",  "3番（アドマイヤテラ）",         3000, "軸馬に最大集中"),
    ("馬連",  "3-7",                          1500, "◎○本線"),
    ("馬連",  "3-4",                          1000, "◎▲"),
    ("ワイド", "3-7",                          1000, "安全網"),
    ("ワイド", "3-4",                           700, "◎▲ワイド"),
    ("ワイド", "3-6",                           500, "川田エヒト穴狙い"),
    ("3連複", "3-7-4（本線）",                 1500, "◎○▲本線"),
    ("3連複", "3-7-6（川田穴馬パターン）",       800, "高配当狙い"),
]

# 的中シナリオ別の期待回収
SCENARIOS = [
    ("◎→○本線決着",   "約20,000〜35,000円"),
    ("◎→▲決着",       "約18,000〜28,000円"),
    ("◎→穴馬エヒト",   "約25,000〜50,000円"),
    ("◎のみ3着以内",   "約8,000〜12,000円"),
]

# ============================================================
# ── ANALYSIS ── 各馬の分析文を記入する
# ============================================================
ANALYSIS_HONMEI = """\
✅ AI能力スコア 28.4%（出走全馬中2位・ただし長距離適性で実質1位）
✅ 前走・阪神大賞典 3000m 1着。王道ローテ
✅ 武豊騎手：京都長距離の信頼度は通算最上位
✅ 2枠3番の内枠でポジション取りにロスなし
✅ スロー想定の3200mで末脚が活きる先行スタイル

阪神大賞典→天皇賞春の王道ローテ＋武豊×京都の最強コンビ。"""

ANALYSIS_TAIKOU = """\
AI能力スコア32.1%（全馬1位）
大阪杯G1制覇の実力馬。4枠7番好位置から上位争い確実。
唯一の懸念は3200mが初距離。実力でカバーできるか注目。"""

ANALYSIS_SANKABAN = """\
AI能力スコア16.8%（全馬3位）
前走阪神大賞典2着。◎との差は首差で実力は同等。
PDCA改善：前走8着のヘデントールより確実性UP。距離延長で浮上。"""

ANALYSIS_穴 = """\
エヒト（川田将雅）：川田乗り替わりで信頼度UP。長距離実績あり。穴候補。
ホーエリート（戸崎圭太）：ダイヤモンドS経験のステイヤー型。距離延長を味方に。"""

ANALYSIS_BET = "昨日ユニコーンS◎○▲ワンツースリー的中の勢いそのままに！"

# ============================================================
# ── NOTE記事 ── 各セクションの本文を記入する
# ============================================================
NOTE_INTRO = (
    f"こんにちは！アスメシ競馬予想です。\n"
    f"今日は{RACE_NAME}（{RACE_GRADE}）をAIで徹底分析しました！\n"
    f"PDCAで改善し続ける手法で今週も攻めます。"
)

NOTE_COURSE = (
    f"{RACE_VENUE}{RACE_COURSE}の特徴を踏まえ、"
    f"ペースシナリオと脚質有利不利を分析しました。"
    f"スロー想定でこそ末脚・先行有利の展開が予測されます。"
)

NOTE_CLOSING = (
    "今日もアスメシ競馬予想の予想を応援よろしくお願いします。\n"
    "予想が当たったらぜひコメントで教えてください！"
)

# ============================================================
# ── ここから下は書き換え不要 ── 自動生成ロジック
# ============================================================
doc = Document()
for section in doc.sections:
    section.top_margin    = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)

FONT = '游ゴシック'

def set_font(run, size=10, bold=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

def h1(text, color=(15, 71, 97)):
    p = doc.add_heading(text, level=1)
    for run in p.runs:
        run.font.color.rgb = RGBColor(*color)
        set_font(run, 14)

def h2(text):
    p = doc.add_heading(text, level=2)
    for run in p.runs:
        set_font(run, 12)

def h3(text):
    p = doc.add_heading(text, level=3)
    for run in p.runs:
        set_font(run, 11)

def box(lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent  = Cm(0.5)
    p.paragraph_format.right_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(8)
    run = p.add_run("\n".join(lines))
    set_font(run, 10)

def body(text, size=10):
    p = doc.add_paragraph(text)
    for run in p.runs:
        set_font(run, size)
    return p

# ──── ヘッダー ────
t = doc.add_heading(f"SNS投稿案  {RACE_DATE[:4]}/{RACE_DATE[4:6]}/{RACE_DATE[6:]}（日）", level=0)
for run in t.runs:
    set_font(run, 18)
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph(f"{RACE_NAME} {RACE_GRADE}  {RACE_VENUE} {RACE_COURSE}  {RACE_COUNT}")
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in sub.runs:
    set_font(run, 11)
doc.add_paragraph()

# ──── X投稿文（長文版） ────
h1(f"{RACE_NAME}（{RACE_GRADE}）  {RACE_VENUE}  {RACE_COURSE}")
h2("X投稿文（認証済み・長文版）")

picks_lines = [f"{印}  {馬番.rjust(2)}番 {馬名}（{騎手}{('  /  ' + 人気) if 人気 else ''}）"
               for 印, 馬番, 馬名, 騎手, 人気 in PICKS]
honmei = next((p for p in PICKS if p[0] == "◎"), None)
taikou = next((p for p in PICKS if p[0] == "○"), None)
sankaban = next((p for p in PICKS if p[0] == "▲"), None)

bet_lines = [f"{種:4}  {組み合わせ:<22} {金額:,}円  ({備考})" for 種, 組み合わせ, 金額, 備考 in BETS]
scenario_lines = [f"{s}：{r}" for s, r in SCENARIOS]

x_lines = [
    "【コピー用】",
    f"🏇【{RACE_NAME} {RACE_GRADE}】{RACE_DATE[:4]}/{RACE_DATE[4:6]}/{RACE_DATE[6:]}",
    f"━━ {RACE_VENUE} {RACE_COURSE} / {RACE_COUNT} ━━",
    "",
    *picks_lines,
    "",
    "─────────────────────",
    f"◎ 本命：{honmei[2]}を選んだ理由",
    "─────────────────────",
    "",
    *ANALYSIS_HONMEI.split("\n"),
    "",
    "─────────────────────",
    f"○ 対抗：{taikou[2]}（{taikou[4]}）",
    "─────────────────────",
    "",
    *ANALYSIS_TAIKOU.split("\n"),
    "",
    "─────────────────────",
    f"▲ 3番手：{sankaban[2]}（{sankaban[4]}）",
    "─────────────────────",
    "",
    *ANALYSIS_SANKABAN.split("\n"),
    "",
    "─────────────────────",
    "★ 穴馬",
    "─────────────────────",
    "",
    *ANALYSIS_穴.split("\n"),
    "",
    "─────────────────────",
    f"買い目（合計{BUDGET:,}円）",
    "─────────────────────",
    "",
    *bet_lines,
    "",
    "【的中シナリオ別 期待回収】",
    *scenario_lines,
    "",
    ANALYSIS_BET,
]
box(x_lines)

# ──── Threads ────
h2("Threads投稿（スレッド形式・5投稿）")

h3("投稿1/5  冒頭・レース紹介")
box([
    "【コピー用】",
    f"{RACE_NAME}（{RACE_GRADE}）予想 {RACE_DATE[:4]}",
    f"{RACE_VENUE} {RACE_COURSE} / {RACE_COUNT}",
    "",
    f"AIと期待値計算でPDCA改善を続ける手法で今週も挑みます！",
    "本命から穴まで全部お伝えします",
])

h3("投稿2/5  予想印発表")
box([
    "【コピー用】",
    "【予想印】",
    *picks_lines,
    "",
    f"本命は{honmei[2]}。穴は{'・'.join(p[2] for p in PICKS if p[0] == '△')}！",
])

h3("投稿3/5  本命詳細分析")
box([
    "【コピー用】",
    f"【本命】{honmei[1]}番 {honmei[2]}（{honmei[3]}  {honmei[4]}）",
    "",
    *ANALYSIS_HONMEI.split("\n"),
])

h3("投稿4/5  対抗・穴馬")
box([
    "【コピー用】",
    "【対抗・穴馬のポイント】",
    "",
    f"○{taikou[1]}番 {taikou[2]}（{taikou[4]}）",
    *[f"→ {l}" for l in ANALYSIS_TAIKOU.split("\n") if l],
    "",
    f"▲{sankaban[1]}番 {sankaban[2]}（{sankaban[4]}）",
    *[f"→ {l}" for l in ANALYSIS_SANKABAN.split("\n") if l],
    "",
    *ANALYSIS_穴.split("\n"),
])

h3("投稿5/5  買い目・締め")
box([
    "【コピー用】",
    f"【買い目まとめ（合計{BUDGET:,}円）】",
    *bet_lines,
    "",
    "【的中シナリオ別 期待回収】",
    *scenario_lines,
    "",
    ANALYSIS_BET,
    "",
    "予想が参考になったらフォロー&いいねお願いします",
])

# ──── Note記事 ────
h2("Note記事（充実版）")

scores_text = "\n".join(
    f"{'◎' if i==0 else '○' if i==1 else '▲'}  {馬番.rjust(2)}番 {馬名}  {人気}  AI能力スコア {SCORES.get(馬名, '-')}"
    for i, (印, 馬番, 馬名, 騎手, 人気) in enumerate(PICKS[:3])
)
delta_text = "\n".join(
    f"△  {馬番.rjust(2)}番 {馬名}"
    for 印, 馬番, 馬名, 騎手, 人気 in PICKS if 印 == "△"
)

note_text = (
    f"# {RACE_NAME}{RACE_DATE[:4]} AI予想完全版\n\n"
    f"{NOTE_INTRO}\n\n"
    "---\n\n"
    "## レース概要\n\n"
    f"開催：{RACE_VENUE}\n"
    f"距離：{RACE_COURSE}\n"
    f"グレード：{RACE_GRADE}  頭数：{RACE_COUNT}  発走：{RACE_TIME}\n\n"
    f"### コース特徴\n{NOTE_COURSE}\n\n"
    "---\n\n"
    "## 予想印一覧\n\n"
    f"{scores_text}\n"
    f"{delta_text}\n\n"
    "---\n\n"
    f"## 本命：{honmei[1]}番 {honmei[2]}（{honmei[3]}  {honmei[4]}）\n\n"
    f"{ANALYSIS_HONMEI}\n\n"
    "---\n\n"
    f"## 対抗：{taikou[1]}番 {taikou[2]}（{taikou[3]}  {taikou[4]}）\n\n"
    f"{ANALYSIS_TAIKOU}\n\n"
    "---\n\n"
    f"## ▲3番手：{sankaban[1]}番 {sankaban[2]}（{sankaban[3]}  {sankaban[4]}）\n\n"
    f"{ANALYSIS_SANKABAN}\n\n"
    "---\n\n"
    "## 穴馬分析\n\n"
    f"{ANALYSIS_穴}\n\n"
    "---\n\n"
    f"## 買い目（合計{BUDGET:,}円）\n\n"
    + "\n".join(bet_lines) +
    "\n\n### 的中シナリオ別 期待回収\n\n"
    + "\n".join(f"- {s}：{r}" for s, r in SCENARIOS) +
    "\n\n---\n\n"
    "## まとめ\n\n"
    f"{NOTE_CLOSING}\n"
)
body(note_text, 10)

# ──── noteヘッダー画像生成コマンド ────
h2("noteヘッダー画像 生成コマンド")
box([
    "以下のコマンドをターミナルで実行してください：",
    "",
    "cd C:/Users/User/dev/move/.claude/skills/nanobanana-pro",
    "",
    "PYTHONUNBUFFERED=1 python -u scripts/run.py image_generator.py \\",
    '  --attach-image "C:/Users/User/Desktop/競馬予想レポート/20260412/note_header_ouka2026.png" \\',
    '  --attach-image "C:/Users/User/Downloads/アスメシキャラクターシート文字なし.png" \\',
    f'  --output "{NOTE_IMAGE_OUT}" \\',
    "  --show-browser --skip-thinking-mode --timeout 600 \\",
    f'  --prompt "添付1枚目と同スタイルで{RACE_NAME}({RACE_GRADE})のnoteサムネイル（横長16:9）。',
    f'2枚目キャラ右下配置。テキスト：左上白・金縁「{RACE_NAME}({RACE_GRADE})」「{RACE_DATE[:4]} AI予想」、',
    f'中段左「アスメシ競馬予想」、下部黒バー「【データ&AI活用による勝ち馬診断】{RACE_VENUE} {RACE_COURSE}」。',
    "鮮やかなアニメ調デジタルイラスト。\"",
])

doc.save(str(OUT))
print(f"保存完了: {OUT}")
print(f"noteヘッダー画像の出力先: {NOTE_IMAGE_OUT}")
