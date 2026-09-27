# -*- coding: utf-8 -*-
"""
eval_results.py — 公開予想（平場ルール適用版ダイジェスト）の答え合わせPDCA（汎用・日付指定）

検証対象:
  A. 印の精度: ◎○▲△🔥穴❌ の着順分布（1着率・連対率・3着内率・平均人気/オッズ）
  B. 公開買い目（flat_race_rules: 4類型多層＋v5.5ガード＋若馬上限）の実払戻ベース回収率
     — netkeiba結果ページの実着順×払戻で精算（推定で報告しない）
  C. 参考: v5エンジン素の選択（choose_pattern_v55）の回収率、見送りレースの「買っていたら」
  D. 類型別・自信度別・会場別の集計 → 週次設定レビュー用JSON
出力: Desktop/競馬予想レポート/{date}/{date}_答え合わせPDCA.docx / .json
使い方: python eval_results.py --date 20260822
"""
import sys, re, json, time, argparse
sys.stdout.reconfigure(encoding='utf-8')
from pathlib import Path
from collections import defaultdict
import requests
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn

SRC = Path(__file__).resolve().parent
sys.path.insert(0, str(SRC)); sys.path.insert(0, str(SRC / 'scripts')); sys.path.insert(0, str(SRC / 'scraper'))
from gen_digest_sns import load, is_grade, CONF_MIN, WEEKDAY
from v55_guard import apply_young_cap, choose_pattern_v55
from gen_kaime_v5 import BUDGET_BY_CONF
from kaime_format import circ
from convert_netkeiba_to_features import rcode_to_netkeiba_id
from smartrc_api import SmartRCAPI

SESSION = requests.Session()
SESSION.headers.update({"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/120.0 Safari/537.36", "Accept-Language": "ja"})
VEN = {"01": "札幌", "02": "函館", "03": "福島", "04": "新潟", "05": "東京", "06": "中山", "07": "中京", "08": "京都", "09": "阪神", "10": "小倉"}
FONT = '游ゴシック'


def fetch_result(race_id):
    """netkeiba結果ページ → (着順リスト[(着順, 馬番, 人気, 単勝オッズ)], payouts{券種: [(組合せ, 払戻100円)]})"""
    url = f"https://race.netkeiba.com/race/result.html?race_id={race_id}"
    time.sleep(1.2)
    r = SESSION.get(url, timeout=20); r.encoding = r.apparent_encoding or "utf-8"
    soup = BeautifulSoup(r.text, "html.parser")
    order = []
    for tr in soup.select("table.RaceTable01 tr.HorseList"):
        tds = tr.find_all("td")
        if len(tds) < 11: continue
        try:
            fin = int(tds[0].get_text(strip=True)); num = int(tds[2].get_text(strip=True))
        except ValueError:
            continue
        pop = tds[9].get_text(strip=True); odds = tds[10].get_text(strip=True)
        order.append((fin, num, int(pop) if pop.isdigit() else None, float(odds) if re.match(r'^[\d.]+$', odds) else None))
    order.sort()
    payouts = defaultdict(list)
    for table in soup.select("table.Payout_Detail_Table"):
        for tr in table.select("tr"):
            th = tr.find("th")
            if not th: continue
            btype = th.get_text(strip=True)
            rtd, ptd = tr.select_one("td.Result"), tr.select_one("td.Payout")
            if not rtd or not ptd: continue
            combos = [el.get_text("-", strip=True) for el in (rtd.find_all(["ul", "div"], recursive=False) or [rtd]) if el.get_text(strip=True)]
            if not combos: combos = [rtd.get_text("-", strip=True)]
            pays = [int(p.replace(",", "")) for p in re.findall(r"([\d,]+)円", ptd.get_text())]
            if len(combos) == 1 and len(pays) > 1:
                nums = re.findall(r"\d+", combos[0])
                if btype == '複勝': combos = nums[:len(pays)]
                elif btype == 'ワイド' and len(nums) >= 2 * len(pays): combos = [f"{nums[2*i]}-{nums[2*i+1]}" for i in range(len(pays))]
            for c, p in zip(combos, pays): payouts[btype].append((c, p))
    return order, payouts


def settle(bets, top3, payouts):
    """bets=[(券種, idxs→馬番list, 金額)] → (投資, 回収, 行)"""
    inv = ret = 0; lines = []
    norm = lambda ns: "-".join(str(n) for n in sorted(int(x) for x in ns))
    for btype, nums, amount in bets:
        nums = [int(n) for n in nums]; inv += amount; hit = False; pay = 0
        if btype == '単勝':
            hit = bool(top3) and nums[0] == top3[0]
            if hit: pay = next((p for c, p in payouts.get('単勝', []) if re.sub(r'\D', '', c) == str(nums[0])), 0)
        elif btype == '馬連':
            hit = len(top3) >= 2 and set(nums) == set(top3[:2])
            if hit: pay = next((p for c, p in payouts.get('馬連', [])), 0)
        elif btype == '馬単':
            hit = len(top3) >= 2 and nums == top3[:2]
            if hit: pay = next((p for c, p in payouts.get('馬単', [])), 0)
        elif btype == 'ワイド':
            hit = set(nums) <= set(top3)
            if hit:
                for c, p in payouts.get('ワイド', []):
                    cn = re.findall(r'\d+', c)
                    if len(cn) >= 2 and norm(cn[:2]) == norm(nums): pay = p; break
        elif btype == '3連複':
            hit = len(top3) >= 3 and set(nums) == set(top3[:3])
            if hit: pay = next((p for c, p in payouts.get('3連複', [])), 0)
        elif btype == '3連単':
            hit = len(top3) >= 3 and nums == top3[:3]
            if hit: pay = next((p for c, p in payouts.get('3連単', [])), 0)
        this = int(pay * amount / 100) if hit and pay else 0; ret += this
        sep = '→' if btype in ('馬単', '3連単') else '-'
        lines.append(f"{'✅' if hit else '　'} {btype} {sep.join(circ(n) for n in nums)} {amount:,}円" + (f" → 払戻{this:,}円（{pay}円）" if this else ("（的中・払戻取得失敗）" if hit else "")))
    return inv, ret, lines


def set_font(run, size=10, bold=False, color=None):
    run.font.name = FONT; run.font.size = Pt(size); run.font.bold = bold
    if color: run.font.color.rgb = RGBColor(*color)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

def H(doc, t, lv=1):
    p = doc.add_heading(t, level=lv)
    for r in p.runs: set_font(r, 14 if lv == 1 else 11.5)

def P(doc, lines, size=10, color=None):
    p = doc.add_paragraph(); r = p.add_run("\n".join(lines) if isinstance(lines, list) else lines); set_font(r, size, color=color)

def T(doc, rows, size=8.5):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = 'Table Grid'
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.text = ''; run = c.paragraphs[0].add_run(str(v)); set_font(run, size, bold=(i == 0))
    doc.add_paragraph()


def main():
    ap = argparse.ArgumentParser(); ap.add_argument('--date', required=True); a = ap.parse_args()
    day = a.date; wd = WEEKDAY.get(day, '')
    base = Path.home() / 'Desktop' / '競馬予想レポート' / day
    races, conf, flat = load(day, {})
    keys = [k for k in races if conf[k][0] >= CONF_MIN and not is_grade(races[k][0]['レース名'])]
    keep, young_dropped = apply_young_cap(keys, races, conf)
    keys.sort(key=lambda k: (-conf[k][0], k[1], k[2]))

    # race_id マップ
    api = SmartRCAPI(); rid = {}
    try:
        for r in api.fetch_races(day):
            nid = rcode_to_netkeiba_id(r['rcode']); rid[(VEN.get(nid[4:6], '?'), int(r.get('rno', 0)))] = nid
    finally:
        api.close()

    # 全レース結果（印精度は全36レース対象）
    results = {}
    for k in sorted(races):
        nid = rid.get((k[1], k[2]))
        if not nid: print(f"[WARN] race_id不明 {k}"); continue
        order, pay = fetch_result(nid)
        if not order: print(f"[WARN] 結果未取得 {k}"); continue
        results[k] = (order, pay)
    print(f"結果取得: {len(results)}/{len(races)}レース")

    # ── A. 印精度 ──
    stat = defaultdict(lambda: {'n': 0, 'w': 0, 'p2': 0, 'p3': 0, 'pop': [], 'odds': []})
    for k, recs in races.items():
        if k not in results: continue
        order, _ = results[k]; fin = {num: f for f, num, _, _ in order}; popm = {num: (p, o) for _, num, p, o in order}
        marks, roles, info, kb = flat[k]
        for num, m in marks.items():
            f = fin.get(num)
            if f is None: continue
            s = stat[m]; s['n'] += 1; s['w'] += (f == 1); s['p2'] += (f <= 2); s['p3'] += (f <= 3)
            p, o = popm.get(num, (None, None))
            if p: s['pop'].append(p)
            if o: s['odds'].append(o)
    # ◎が3着内の率（週次制御の判定値）
    hon = stat['◎']; hon_p3 = hon['p3'] / hon['n'] if hon['n'] else 0

    # ── B. 公開買い目の精算 ──
    race_rows = []; tot_inv = tot_ret = 0; by_arch = defaultdict(lambda: [0, 0, 0, 0]); by_conf = defaultdict(lambda: [0, 0, 0, 0]); by_venue = defaultdict(lambda: [0, 0, 0, 0])
    skipped_would = []
    eng_inv = eng_ret = 0
    for k in keys:
        recs = races[k]; c = conf[k][0]; marks, roles, info, kb = flat[k]
        if k not in results: continue
        order, pay = results[k]; top3 = [num for f, num, _, _ in order[:3]]
        fin = {num: f for f, num, _, _ in order}
        hon_fin = fin.get(recs[0]['馬番'])
        arch = kb['arch_name']
        if k not in keep or kb['skip']:
            why = young_dropped.get(k, '期待値基準未達')
            # 見送りの「買っていたら」: 多層の最良案 or なし
            would = None
            if kb['res']:
                bets = [(b, [recs[i]['馬番'] for i in idxs], amt) for b, idxs, amt, _ in kb['res']['bets']]
                inv, ret, _ = settle(bets, top3, pay); would = (inv, ret)
            skipped_would.append((k, recs[0]['レース名'], c, why, hon_fin, top3, would))
            continue
        bets = [(b, [recs[i]['馬番'] for i in idxs], amt) for b, idxs, amt, _ in kb['res']['bets']]
        inv, ret, lines = settle(bets, top3, pay)
        tot_inv += inv; tot_ret += ret
        for d_, key in ((by_arch, arch.split('・')[0].split('→')[0]), (by_conf, c), (by_venue, k[1])):
            d_[key][0] += 1; d_[key][1] += inv; d_[key][2] += ret; d_[key][3] += (ret > 0)
        # 参考: エンジン素の選択
        name, desc, res, ranking, p, q, vinfo = choose_pattern_v55(recs, BUDGET_BY_CONF[c])
        e_inv = e_ret = 0; e_name = name
        if res:
            ebets = [(b, [recs[i]['馬番'] for i in idxs], amt) for b, idxs, amt, _ in res['bets']]
            e_inv, e_ret, _ = settle(ebets, top3, pay)
        eng_inv += e_inv; eng_ret += e_ret
        race_rows.append({'key': k, 'title': recs[0]['レース名'], 'conf': c, 'arch': arch, 'inv': inv, 'ret': ret, 'lines': lines,
                          'top3': top3, 'hon_fin': hon_fin, 'order': order[:5], 'engine': (e_name, e_inv, e_ret),
                          'marks': marks, 'E': kb['res']['E_rate'], 'P300': kb['res']['P_target']})

    # ── docx ──
    doc = Document()
    for s in doc.sections: s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(1.8)
    t = doc.add_heading(f"{day[:4]}/{day[4:6]}/{day[6:]}（{wd}） 答え合わせPDCA — 公開予想（平場ルール適用版）", level=0)
    for r in t.runs: set_font(r, 15)
    roi = tot_ret / tot_inv * 100 if tot_inv else 0
    P(doc, [f"参加{len(race_rows)}レース 投資{tot_inv:,}円 → 回収{tot_ret:,}円（回収率{roi:.0f}%・的中{sum(1 for x in race_rows if x['ret'] > 0)}レース）",
            f"参考: v5エンジン素の選択なら 投資{eng_inv:,}円 → 回収{eng_ret:,}円（{(eng_ret/eng_inv*100 if eng_inv else 0):.0f}%）",
            f"◎3着内率（全{hon['n']}レース）= {hon_p3*100:.0f}% → 週次制御の判定: {'ON継続（<45%）' if hon_p3 < 0.45 else ('解除可（≧50%）' if hon_p3 >= 0.50 else '据え置き（45〜50%）')}",
            f"実着順・払戻はnetkeiba結果ページから取得（{len(results)}/{len(races)}レース）"], size=10.5)

    H(doc, "A. 印の精度（全レース）")
    rows = [["印", "頭数", "1着", "連対", "3着内", "1着率", "3着内率", "平均人気", "平均単勝"]]
    for m in ['◎', '○', '▲', '△', '🔥穴', '❌']:
        s = stat[m]
        if not s['n']: continue
        rows.append([m, s['n'], s['w'], s['p2'], s['p3'], f"{s['w']/s['n']*100:.0f}%", f"{s['p3']/s['n']*100:.0f}%",
                     f"{sum(s['pop'])/len(s['pop']):.1f}" if s['pop'] else '-', f"{sum(s['odds'])/len(s['odds']):.1f}" if s['odds'] else '-'])
    T(doc, rows)

    H(doc, "B. 公開買い目の答え合わせ（参加レース）")
    for x in sorted(race_rows, key=lambda x: -x['conf']):
        k = x['key']; r100 = x['ret'] / x['inv'] * 100 if x['inv'] else 0
        ordtxt = " ".join(f"{f}着{circ(num)}({p}人気)" for f, num, p, o in x['order'][:3])
        H(doc, f"{k[1]}{k[2]}R {x['title']} 自信度{x['conf']} ｜ {x['arch']} ｜ {x['inv']:,}円→{x['ret']:,}円（{r100:.0f}%）{' ✅' if x['ret'] else ''}", 2)
        P(doc, [f"結果: {ordtxt} ／ ◎{circ(races[k][0]['馬番'])}{races[k][0]['馬名']}は{x['hon_fin']}着",
                f"印: " + "　".join(f"{m}{circ(n)}" for n, m in x['marks'].items()),
                f"事前評価: 期待回収率{x['E']*100:.0f}% / 300%超{x['P300']*100:.0f}%",
                f"参考（v5エンジン素の選択 {x['engine'][0] or '見送り'}）: {x['engine'][1]:,}円→{x['engine'][2]:,}円"] + x['lines'], size=9.5)

    H(doc, "C. 見送りレースの結果（買っていたら）")
    rows = [["会場R", "レース", "自信度", "見送り理由", "◎着順", "1-2-3着", "多層案の仮精算"]]
    for k, title, c, why, hf, top3, would in skipped_would:
        rows.append([f"{k[1]}{k[2]}R", title[:14], c, why[:22], hf, "-".join(circ(n) for n in top3), (f"{would[0]:,}→{would[1]:,}円" if would else "案なし")])
    T(doc, rows)

    H(doc, "D. 集計（類型別／自信度別／会場別）")
    for label, d_ in (("類型", by_arch), ("自信度", by_conf), ("会場", by_venue)):
        rows = [[label, "レース", "投資", "回収", "回収率", "的中R"]]
        for key, (n, inv, ret, hit) in sorted(d_.items(), key=lambda x: str(x[0])):
            rows.append([key, n, f"{inv:,}", f"{ret:,}", f"{(ret/inv*100 if inv else 0):.0f}%", hit])
        T(doc, rows)

    out = base / f"{day}_答え合わせPDCA.docx"; doc.save(str(out))
    js = {'date': day, 'invest': tot_inv, 'return': tot_ret, 'roi': roi, 'hon_p3': hon_p3,
          'marks': {m: {kk: (v if kk not in ('pop', 'odds') else (sum(v)/len(v) if v else None)) for kk, v in s.items()} for m, s in stat.items()},
          'races': [{'key': list(x['key']), 'title': x['title'], 'conf': x['conf'], 'arch': x['arch'], 'inv': x['inv'], 'ret': x['ret'], 'top3': x['top3'], 'hon_fin': x['hon_fin'], 'engine': x['engine']} for x in race_rows],
          'skipped': [{'key': list(k), 'title': t_, 'conf': c, 'why': w, 'hon_fin': hf, 'top3': t3, 'would': wd_} for k, t_, c, w, hf, t3, wd_ in skipped_would],
          'by_arch': dict(by_arch), 'by_conf': {str(k): v for k, v in by_conf.items()}, 'by_venue': dict(by_venue), 'engine': [eng_inv, eng_ret]}
    json.dump(js, open(base / f"{day}_答え合わせPDCA.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"保存: {out}")
    print(f"参加{len(race_rows)}R 投資{tot_inv:,} 回収{tot_ret:,} ROI{roi:.0f}% | エンジン素 {eng_inv:,}→{eng_ret:,} | ◎3着内率{hon_p3*100:.0f}%")
    for x in race_rows: print(f"  {x['key'][1]}{x['key'][2]}R {x['title'][:12]} {x['arch'][:16]} {x['inv']:,}→{x['ret']:,} ◎{x['hon_fin']}着 top3={x['top3']}")
    for m in ['◎', '○', '▲', '△', '🔥穴', '❌']:
        s = stat[m]
        if s['n']: print(f"  {m}: n={s['n']} 1着{s['w']} 3内{s['p3']} ({s['p3']/s['n']*100:.0f}%)")


if __name__ == '__main__':
    main()
